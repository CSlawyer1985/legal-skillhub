#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from common import ASSETS, SKILL_VERSION, WORD_DOCS_DEFAULT, FlowError, clean_filename, deterministic_zip, safe_write_bytes

STYLE = json.loads((ASSETS / "word-style.json").read_text(encoding="utf-8"))
BOUNDARY = "本文件为AI数据合规工作候选材料。客户事实在取得证据并确认前均为待核实；评估结论须经律师复核，不得据此直接认定合规、违法或替代专项法律意见。"

# 标准尽调内容默认值（profile 未配置时回退；配置了则使用 profile 值）
DEFAULT_CONTENT = {
    "legal_bases": ["个人同意", "合同履行必要", "法定义务", "合法利益", "公开信息"],
    "data_lifecycle": ["收集", "存储", "使用", "共享", "跨境传输", "删除"],
    "interview_questions": {
        "管理层": ["项目立项背景与业务目标是什么？", "合规评估的范围和边界如何界定？", "项目投入预算与合规资源如何安排？", "上线时间表与合规节点如何衔接？"],
        "业务负责人": ["业务流程中AI与人工如何分工？", "高风险场景的转人工规则如何设计？", "客户投诉处理流程如何闭环？", "业务指标与合规要求冲突时如何取舍？"],
        "数据安全": ["数据如何分类分级？", "访问控制策略如何落地？", "加密与脱敏措施覆盖哪些环节？", "数据泄露应急响应预案是否就绪？"],
        "法务合规": ["算法备案进展如何？", "用户协议与隐私政策是否更新？", "生成内容标识义务如何履行？", "供应商合同中的合规条款是否齐备？"],
        "供应商负责人": ["供应商遴选标准如何设定？", "合同中数据安全与责任条款是否明确？", "SLA与违约责任如何约定？", "供应商退出时数据如何删除或返还？"],
        "技术架构师": ["系统架构中数据流向如何设计？", "API鉴权与调用日志如何管理？", "提示词注入防护如何实现？", "模型输出过滤与敏感信息屏蔽如何落地？"],
        "产品经理": ["产品功能边界如何定义？", "用户交互中的告知与授权如何呈现？", "转人工策略如何触发？", "内容审核机制如何运作？"],
        "客服运营": ["客服话术与知识库如何维护更新？", "人工客服操作日志如何记录？", "投诉升级路径如何设计？", "员工培训与合规意识如何保障？"],
        "算法工程师": ["模型训练数据来源与授权如何管理？", "模型评测与公平性如何验证？", "日志留存与审计如何实现？", "模型版本更新如何管控？"],
    },
    "data_assets": [
        {"name": "用户对话记录", "category": "个人信息", "source": "平台用户交互", "purpose": "智能客服服务", "legal_basis": "个人同意/合同履行必要", "storage": "生产环境数据库", "retention": "按平台政策", "sharing": "客服Agent处理", "cross_border": "否"},
        {"name": "订单信息", "category": "敏感个人信息", "source": "电商交易系统", "purpose": "订单查询与售后", "legal_basis": "合同履行必要", "storage": "生产环境数据库", "retention": "按平台政策", "sharing": "客服Agent处理", "cross_border": "否"},
    ],
    "suppliers": [
        {"name": "外部模型服务商", "service": "生成式AI模型API", "role": "受托处理者", "cross_border": "否", "security": "加密传输、访问控制、调用日志", "exit": "合同终止后删除或返还数据"},
    ],
    "risk_matrix": [
        {"theme": "个人信息泄露", "description": "处理过程中个人信息被未授权访问或披露", "law": "《个人信息保护法》第9、51条", "likelihood": "中", "impact": "高", "level": "高", "controls": "加密、访问控制、脱敏", "residual": "中", "remediation": "定期渗透测试与审计"},
        {"theme": "权限越界", "description": "AI或人员访问超出授权范围的数据", "law": "《个人信息保护法》第6条", "likelihood": "中", "impact": "中", "level": "中", "controls": "分层权限、登录态校验", "residual": "中", "remediation": "权限定期复核"},
    ],
    "gate_conditions": [
        {"condition": "P0资料完整", "evidence": "P0资料清单全部已提供", "method": "核对资料清单"},
        {"condition": "事实已经确认", "evidence": "事实确认稿签署", "method": "核对确认主体与日期"},
        {"condition": "数据与系统边界明确", "evidence": "数据地图与系统地图", "method": "文档查验"},
        {"condition": "供应商责任明确", "evidence": "供应商合同与责任条款", "method": "合同查验"},
        {"condition": "重大风险已有控制", "evidence": "风险控制措施清单", "method": "控制措施核验"},
        {"condition": "律师复核字段明确", "evidence": "律师复核意见", "method": "复核记录查验"},
    ],
}


def content_for(profile: dict) -> dict:
    """合并 profile 内容配置与默认值：profile 有则用，无则回退默认。"""
    content = {}
    for key, default in DEFAULT_CONTENT.items():
        value = profile.get(key)
        content[key] = value if value is not None else default
    return content


def set_run_font(run, name: str, size: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent=140) -> None:
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            cell.width = Mm(width / 1440 * 25.4)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def configure_styles(doc: Document) -> None:
    fonts = STYLE["fonts"]
    sizes = STYLE["sizes_pt"]
    normal = doc.styles["Normal"]
    normal.font.name = fonts["body"]
    normal.font.size = Pt(sizes["body"])
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["body"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for style_name, font_name, size, before, after in (
        ("Title", fonts["title"], sizes["title"], 0, 14),
        ("Subtitle", fonts["secondary"], 11, 0, 12),
        ("Heading 1", fonts["heading"], sizes["h1"], 16, 8),
        ("Heading 2", fonts["heading"], sizes["h2"], 12, 6),
        ("Heading 3", fonts["heading"], sizes["h3"], 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True
        if style_name == "Title":
            p_pr = style._element.get_or_add_pPr()
            p_bdr = p_pr.find(qn("w:pBdr"))
            if p_bdr is not None:
                p_pr.remove(p_bdr)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, STYLE["fonts"]["title"], 9)


def configure_page(doc: Document, project_name: str, short_title: str, candidate: bool) -> None:
    page = STYLE["page"]
    for section in doc.sections:
        section.page_width = Mm(page["width_mm"])
        section.page_height = Mm(page["height_mm"])
        section.top_margin = Mm(page["top_mm"])
        section.bottom_margin = Mm(page["bottom_mm"])
        section.left_margin = Mm(page["left_mm"])
        section.right_margin = Mm(page["right_mm"])
        section.header_distance = Mm(page["header_mm"])
        section.footer_distance = Mm(page["footer_mm"])
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_after = Pt(0)
        hr = hp.add_run(f"{project_name}｜{short_title}")
        set_run_font(hr, STYLE["fonts"]["title"], 9)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prefix = "候选版／需律师确认  " if candidate else "客户事实待核实  "
        fr = fp.add_run(prefix + "第")
        set_run_font(fr, STYLE["fonts"]["title"], 9)
        add_page_field(fp)
        tail = fp.add_run("页")
        set_run_font(tail, STYLE["fonts"]["title"], 9)


def add_notice(doc: Document, candidate: bool) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    cell.text = ""
    set_cell_shading(cell, STYLE["colors"]["callout"])
    p = cell.paragraphs[0]
    r = p.add_run(("候选版／需律师确认。" if candidate else "") + BOUNDARY)
    set_run_font(r, STYLE["fonts"]["secondary"], 10.5, bold=True)
    set_table_geometry(table, [8700])
    prevent_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def new_document(project: dict, title: str, short_title: str, doc_no: str, candidate: bool) -> Document:
    doc = Document()
    configure_styles(doc)
    configure_page(doc, project["project_name"], short_title, candidate)
    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    set_run_font(title_run, STYLE["fonts"]["title"], STYLE["sizes_pt"]["title"], bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    label = "候选版／需律师确认" if candidate else "客户启动工作文件"
    sr = subtitle.add_run(f"文件编号：{project['project_id']}-{doc_no}    版本：{SKILL_VERSION}")
    set_run_font(sr, STYLE["fonts"]["secondary"], 10.5)
    property_line = doc.add_paragraph(style="Subtitle")
    property_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    property_line.paragraph_format.space_after = Pt(10)
    pr = property_line.add_run(f"文件属性：{label}")
    set_run_font(pr, STYLE["fonts"]["secondary"], 10.5)
    add_notice(doc, candidate)
    return doc


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_paragraph(doc: Document, text: str, bold_label: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_label and text.startswith(bold_label):
        first = p.add_run(bold_label)
        set_run_font(first, STYLE["fonts"]["body"], 12, bold=True)
        rest = p.add_run(text[len(bold_label):])
        set_run_font(rest, STYLE["fonts"]["body"], 12)
    else:
        run = p.add_run(text)
        set_run_font(run, STYLE["fonts"]["body"], 12)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, STYLE["colors"]["table_header"])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, STYLE["fonts"]["heading"], 10.5, bold=True)
    set_repeat_header(table.rows[0])
    prevent_split(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        prevent_split(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, STYLE["fonts"]["body"], 10.5)
    set_table_geometry(table, widths)


def save_docx(doc: Document, path: Path) -> None:
    if path.exists():
        raise FlowError(f"同名Word成果已存在：{path}", 4)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp-{os.getpid()}.docx")
    doc.save(temporary)
    safe_write_bytes(path, temporary.read_bytes())
    temporary.unlink(missing_ok=True)


def build_startup_pack(project: dict, profile: dict, output_dir: Path) -> list[Path]:
    project_name = project["project_name"]
    word_docs = profile.get("word_docs") or {}
    docs = word_docs.get("startup") or WORD_DOCS_DEFAULT["startup"]
    content = content_for(profile)
    output_dir.mkdir(parents=True, exist_ok=True)
    built: list[Path] = []
    for entry in docs:
        doc_no = str(entry["doc_no"])
        title = entry["title"]
        short_title = entry.get("short_title", "")
        purpose = entry.get("purpose", "")
        full_title = f"{project_name}\n{title}"
        doc = new_document(project, full_title, short_title, f"START-{doc_no}", False)
        add_heading(doc, "一、文件用途")
        add_paragraph(doc, purpose)
        add_heading(doc, "二、项目基本信息")
        add_table(doc, ["项目", "内容"], [
            ["行业配置", profile["display_name"]],
            ["目标用例", project["use_case"]],
            ["适用地区", project["jurisdiction"]],
            ["当前边界", "客户事实待核实；本文件不构成正式法律意见"],
        ], [1900, 6800])
        if doc_no == "01":
            add_heading(doc, "三、P0资料清单")
            add_table(doc, ["序号", "资料名称", "用途说明", "提交状态"], [[str(i), item, "确认项目边界所需最低资料", "待提供"] for i, item in enumerate(profile["p0_items"], 1)], [700, 3300, 3300, 1400])
            add_heading(doc, "四、填写指引")
            add_paragraph(doc, "请逐项标注提交状态（已提供/待提供）；已提供的资料请填写证据编号或文件路径，供合规团队核验。")
        elif doc_no == "02":
            add_heading(doc, "三、访谈对象与核验问题")
            add_paragraph(doc, "按角色开展访谈，逐项记录回答要点与证据。问题可由项目组按实际增删。")
            questions = content["interview_questions"]
            rows = []
            for index, role in enumerate(profile["roles"], 1):
                role_questions = questions.get(role, questions.get("管理层", []))
                for q_index, question in enumerate(role_questions, 1):
                    rows.append([f"{index}.{q_index}", role if q_index == 1 else "", question, ""])
            add_table(doc, ["序号", "角色", "核验问题", "回答要点/证据"], rows, [800, 1800, 3900, 2200])
        elif doc_no == "03":
            add_heading(doc, "三、数据资产清单")
            add_paragraph(doc, "请逐项填写实际处理的数据资产；敏感个人信息请特别标注。")
            assets = content["data_assets"]
            rows = []
            for index, asset in enumerate(assets, 1):
                rows.append([
                    str(index),
                    asset.get("name", ""),
                    asset.get("category", ""),
                    asset.get("source", ""),
                    asset.get("purpose", ""),
                    asset.get("legal_basis", ""),
                    asset.get("storage", ""),
                    asset.get("retention", ""),
                    asset.get("sharing", ""),
                    asset.get("cross_border", ""),
                ])
            add_table(doc, ["序号", "数据资产", "分类", "来源", "处理目的", "合法性基础", "存储位置", "保留期限", "共享对象", "跨境"], rows, [600, 1200, 850, 900, 1000, 1100, 850, 800, 800, 600])
            add_heading(doc, "四、数据生命周期说明")
            add_paragraph(doc, "生命周期各环节处理情况：收集 / 存储 / 使用 / 共享 / 跨境传输 / 删除。" + "；".join(content["data_lifecycle"]))
            add_heading(doc, "五、合法性基础选项")
            add_table(doc, ["序号", "合法性基础", "说明"], [[str(i), item, "请勾选适用情形并说明依据"] for i, item in enumerate(content["legal_bases"], 1)], [800, 3000, 4900])
        elif doc_no == "04":
            add_heading(doc, "三、供应商及外部模型清单")
            add_paragraph(doc, "请逐项填写实际使用的供应商与外部模型；跨境调用请特别标注。")
            suppliers = content["suppliers"]
            rows = []
            for index, supplier in enumerate(suppliers, 1):
                rows.append([
                    str(index),
                    supplier.get("name", ""),
                    supplier.get("service", ""),
                    supplier.get("role", ""),
                    supplier.get("cross_border", ""),
                    supplier.get("security", ""),
                    supplier.get("exit", ""),
                ])
            add_table(doc, ["序号", "供应商", "服务内容", "处理角色", "跨境", "安全措施", "退出/删除机制"], rows, [600, 1300, 1300, 1100, 700, 2000, 1700])
        elif doc_no == "05":
            add_heading(doc, "三、风险识别与评估矩阵")
            add_paragraph(doc, "逐项评估可能性、影响与等级，并记录现有控制措施、剩余风险与整改建议。")
            matrix = content["risk_matrix"]
            rows = []
            for index, risk in enumerate(matrix, 1):
                rows.append([
                    str(index),
                    risk.get("theme", ""),
                    risk.get("description", ""),
                    risk.get("law", ""),
                    risk.get("likelihood", ""),
                    risk.get("impact", ""),
                    risk.get("level", ""),
                    risk.get("controls", ""),
                    risk.get("residual", ""),
                    risk.get("remediation", ""),
                ])
            add_table(doc, ["序号", "风险主题", "风险描述", "关联法源", "可能性", "影响", "等级", "现有控制", "剩余风险", "整改建议"], rows, [600, 1100, 1400, 1100, 700, 600, 600, 1000, 800, 900])
        elif doc_no == "06":
            add_heading(doc, "三、上线准入条件核验")
            add_paragraph(doc, "全部条件满足并取得书面决定前，不得使用真实数据开展试点。")
            gates = content["gate_conditions"]
            rows = []
            for index, gate in enumerate(gates, 1):
                rows.append([str(index), gate.get("condition", ""), gate.get("evidence", ""), gate.get("method", ""), "未核验", "", ""])
            add_table(doc, ["序号", "准入条件", "证据要求", "核验方法", "状态", "核验人", "核验日期"], rows, [600, 1900, 2300, 1500, 800, 800, 900])
        else:
            add_heading(doc, "三、填写与核验记录")
            add_table(doc, ["序号", "事项", "客户填写/证据编号", "核验状态"], [[str(i), theme, "", "待核实"] for i, theme in enumerate(profile["risk_themes"][:6], 1)], [700, 2900, 3300, 1800])
        path = output_dir / f"{doc_no}-{title}.docx"
        save_docx(doc, path)
        built.append(path)
    deterministic_zip(output_dir.parent / "客户启动Word公文包.zip", [(path, path.name) for path in built])
    return built


def build_candidate_pack(project: dict, profile: dict, records: dict[str, list[dict]], output_dir: Path) -> list[Path]:
    project_name = project["project_name"]
    word_docs = profile.get("word_docs") or {}
    documents = word_docs.get("candidate") or WORD_DOCS_DEFAULT["candidate"]
    output_dir.mkdir(parents=True, exist_ok=True)
    built: list[Path] = []
    counts = {name: len(rows) for name, rows in records.items()}
    for entry in documents:
        doc_no = str(entry["doc_no"])
        title = entry["title"]
        short_title = entry.get("short_title", "")
        doc = new_document(project, f"{project_name}\n{title}", short_title, f"CAND-{doc_no}", True)
        add_heading(doc, "一、使用边界")
        add_paragraph(doc, BOUNDARY)
        add_heading(doc, "二、项目摘要")
        add_table(doc, ["项目", "内容"], [
            ["行业配置", profile["display_name"]],
            ["目标用例", project["use_case"]],
            ["事实状态", "已通过本轮资料门禁；仍以书面证据和律师复核为准"],
            ["交付属性", "候选版／需律师确认"],
        ], [1900, 6800])
        add_heading(doc, "三、证据与评估概况")
        add_table(doc, ["记录类型", "数量", "复核要求"], [
            ["来源记录", str(counts.get("source_records", 0)), "法源版本及适用性需复核"],
            ["主张记录", str(counts.get("claim_records", 0)), "法律效力和条件需复核"],
            ["事实记录", str(counts.get("fact_records", 0)), "以客户书面确认和证据为准"],
            ["评估记录", str(counts.get("assessment_records", 0)), "按真实事实筛选，不统一评分"],
            ["问题记录", str(counts.get("issue_records", 0)), "整改、复验和关闭证据待跟踪"],
        ], [2200, 1200, 5300])
        if doc_no == "02":
            add_heading(doc, "四、客户事实确认明细")
            fact_rows = records.get("fact_records") or []
            rows = []
            for fact in fact_rows:
                rows.append([
                    fact.get("fact_id", ""),
                    fact.get("fact", ""),
                    fact.get("evidence_id", ""),
                    fact.get("fact_status", ""),
                    fact.get("confirmed_by", ""),
                    fact.get("confirmed_at", ""),
                    fact.get("confirmation_method", ""),
                ])
            if not rows:
                rows = [["", "暂无事实记录", "", "", "", "", ""]]
            add_table(doc, ["编号", "事实", "证据编号", "状态", "确认主体", "确认日期", "确认方式"], rows, [900, 2400, 1300, 1200, 1200, 800, 900])
        add_heading(doc, "四、本文件关注事项" if doc_no != "02" else "五、本文件关注事项")
        themes = profile["risk_themes"][:8]
        add_table(doc, ["序号", "事项", "候选状态"], [[str(i), item, "需律师确认"] for i, item in enumerate(themes, 1)], [800, 5700, 2200])
        path = output_dir / f"{doc_no}-{title}.docx"
        save_docx(doc, path)
        built.append(path)
    deterministic_zip(output_dir.parent / "AI数据合规候选交付Word包.zip", [(path, path.name) for path in built])
    return built


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        parts = []
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                parts.append(archive.read(name).decode("utf-8", errors="ignore"))
        return "\n".join(parts)


# 访谈问题库：角色 → [(问题, 追问要点, 证据锚点)]（AI 数据合规进企访谈，按手册评估方法组织）
INTERVIEW_QUESTIONS = {
    "管理层": [
        ("本生成式AI系统的业务定位、上线时间与当前使用规模？",
         "追问：立项与决策人是谁？有无立项评审/试运行记录？",
         "立项文档、业务介绍、试运行报告"),
        ("数据安全管理责任是否落实到具体负责人与部门？",
         "追问：有无正式任命文件？是否定期向管理层述职？",
         "数据安全负责人任命文件、组织架构图"),
        ("是否完成生成式AI备案或算法安全评估？",
         "追问：备案编号与评估结论；未备案的原因与计划？",
         "备案凭证、安全评估报告"),
        ("数据安全与合规投入情况（预算、人员、外包）？",
         "追问：安全投入占IT预算比例；外包安全服务商名单？",
         "预算文件、安全服务合同"),
        ("数据安全事件的上报路径与处置流程？",
         "追问：是否发生过事件？有无预案与演练？",
         "应急预案、演练记录、事件复盘报告"),
        ("隐私政策/用户协议对外承诺与实际操作是否一致？",
         "追问：已知不一致点有哪些？整改计划与时限？",
         "隐私政策、用户协议、整改记录"),
        ("是否向第三方提供数据（含模型供应商、云服务商）？",
         "追问：供应商名单与数据范围？有无合同与授权？",
         "供应商清单、委托处理合同（DPA）"),
        ("是否存在数据出境安排？",
         "追问：出境方式（传输/访问）？是否完成安全评估或备案？",
         "出境安全评估报告、标准合同备案凭证"),
        ("员工接触数据的权限如何管理？",
         "追问：权限申请/审批流程？有无定期审计？",
         "权限管理制度、权限审计记录"),
        ("业务快速迭代与合规审查如何协调？",
         "追问：新功能上线前有无合规评审节点？",
         "上线流程文档、评审记录"),
        ("是否收到过监管询问、用户投诉或整改要求？",
         "追问：处理结果与闭环记录？",
         "监管往来函件、投诉台账、整改报告"),
        ("对外宣传/招投标材料中关于AI能力的表述是否与实际情况一致？",
         "追问：宣传与系统真实能力差异点？",
         "宣传材料、产品能力说明"),
    ],
    "数据安全负责人": [
        ("数据分类分级制度是否建立并落地执行？",
         "追问：分级标准？哪些数据被定为敏感/重要数据？",
         "分类分级制度、分级结果台账"),
        ("是否建立数据资产台账并定期更新？",
         "追问：台账覆盖范围（库/表/文件）？更新频率？",
         "数据资产清单"),
        ("是否开展过个人信息保护影响评估（PIA）？",
         "追问：评估范围与结论？整改项是否闭环？",
         "PIA报告、整改记录"),
        ("数据收集是否遵循最小必要原则？",
         "追问：收集字段清单与业务目的的对应关系？",
         "收集字段清单、授权记录"),
        ("去标识化/匿名化技术应用于哪些环节？",
         "追问：采用何种技术？再识别风险评估？",
         "技术方案、算法说明"),
        ("数据留存期限如何设定与执行？",
         "追问：到期自动删除机制？删除执行记录？",
         "留存策略文档、删除执行记录"),
        ("访问控制如何实现最小权限？",
         "追问：权限矩阵？特权账号管理？离职回收流程？",
         "权限矩阵、系统配置、账号管理记录"),
        ("传输与存储加密措施及密钥管理？",
         "追问：加密算法与范围？密钥轮换机制？",
         "加密方案、密钥管理制度"),
        ("备份与恢复机制及演练情况？",
         "追问：备份频率与恢复点目标（RPO）？",
         "备份策略、恢复演练记录"),
        ("第三方数据共享的审批流程？",
         "追问：共享清单？审批留痕？",
         "共享审批记录、共享数据清单"),
        ("日志记录范围与审计安排？",
         "追问：覆盖哪些环节（登录/操作/数据导出）？审计频率？",
         "日志规范、审计记录"),
        ("用户权利请求（查询/更正/删除）处理流程与时限？",
         "追问：响应时限？拒绝理由记录？",
         "处理流程文档、请求响应记录"),
        ("员工数据安全培训频率与覆盖范围？",
         "追问：培训内容与考核？新员工入职培训？",
         "培训计划、培训签到与考核记录"),
        ("数据安全事件应急响应与演练情况？",
         "追问：预案版本与更新时间？最近一次演练？",
         "应急预案、演练与复盘报告"),
    ],
    "技术负责人": [
        ("模型来源（自研/开源/商业API）与训练数据构成？",
         "追问：基座模型与版本？微调数据来源？",
         "模型说明文档、训练数据清单"),
        ("训练数据来源的合法性与授权链？",
         "追问：爬取/购买/自有数据？版权与个人信息评估？",
         "数据来源协议、合法性评估记录"),
        ("是否完成生成式AI备案与安全评估？",
         "追问：备案编号？评估报告关键结论？",
         "备案凭证、安全评估报告"),
        ("输入输出内容安全措施（敏感词/内容标识/水印）？",
         "追问：拦截策略？标识方式？测试覆盖？",
         "内容安全方案、测试记录"),
        ("推理服务的输入输出是否留存？留存范围与期限？",
         "追问：留存目的与授权依据？",
         "留存策略文档"),
        ("系统架构与数据流是否与合规文档一致？",
         "追问：现场核对数据流图与实际系统？",
         "架构文档、数据流图"),
        ("依赖哪些外部API/云服务？",
         "追问：依赖清单？合同与数据流向？",
         "依赖清单、供应商合同"),
        ("漏洞管理与安全测试安排？",
         "追问：扫描频率？渗透测试最近一次？",
         "漏洞管理流程、渗透测试报告"),
        ("多租户/多客户数据隔离措施？",
         "追问：隔离粒度（库/表/行）？",
         "数据隔离方案"),
        ("未成年人/特殊群体保护的技术措施？",
         "追问：识别机制？限制策略？",
         "识别与保护方案"),
        ("模型版本管理与回滚机制？",
         "追问：版本记录？线上回滚流程？",
         "版本管理规范"),
        ("输出真实性/幻觉风险的控制措施？",
         "追问：质检机制？人工兜底方案？",
         "质检机制、人工兜底流程"),
        ("日志与审计系统覆盖哪些环节？",
         "追问：日志留存时长？是否可追溯至个人？",
         "日志架构文档"),
        ("运维变更流程与审批？",
         "追问：变更窗口？回退方案？",
         "变更管理记录"),
    ],
    "法务合规负责人": [
        ("隐私政策/用户协议是否与产品功能一致并经评审？",
         "追问：最近一次评审时间？功能变更后是否同步更新？",
         "隐私政策、用户协议、评审记录"),
        ("个人信息收集的知情同意链路（弹窗/勾选/撤回）？",
         "追问：同意记录是否留存？撤回机制是否可用？",
         "交互截图、同意记录样本"),
        ("供应商合同是否含数据委托处理条款（DPA）？",
         "追问：覆盖哪些供应商？责任边界是否清晰？",
         "DPA、委托处理合同"),
        ("员工个人信息处理（入职/监控/离职）是否合规？",
         "追问：监控告知？离职数据处置？",
         "员工数据处理相关文件"),
        ("第三方数据来源合同是否含合法性保证条款？",
         "追问：保证范围？违约追责机制？",
         "数据采购合同"),
        ("监管报送/备案材料清单与更新机制？",
         "追问：最近一次报送？负责人？",
         "报送记录"),
        ("个人信息投诉处理机制与时限？",
         "追问：处理流程？法定时限遵守情况？",
         "投诉处理制度、投诉台账"),
        ("数据跨境的法律评估与批准文件？",
         "追问：评估结论？批准层级？",
         "跨境评估文件、批准记录"),
        ("训练数据知识产权（版权）评估？",
         "追问：输出内容版权归属约定？",
         "知识产权评估意见"),
        ("合规审查流程如何运作（谁审/留痕）？",
         "追问：审查清单？审查意见留存？",
         "审查流程文档、审查记录"),
        ("与业务/技术团队的合规协作机制？",
         "追问：新项目早期介入机制？",
         "协作流程文档"),
        ("合同中关于AI能力的承诺与实际履约情况？",
         "追问：有无承诺无法兑现的情形？",
         "合同、履约记录"),
    ],
    "业务运营": [
        ("客服场景的典型用户路径与AI接触点？",
         "追问：AI答复与人工客服的切换节点？",
         "业务流程文档"),
        ("实际收集的数据类型与业务目的的匹配性？",
         "追问：是否有超出目的收集？",
         "字段清单、目的说明"),
        ("自动化决策/智能答复的人工介入机制？",
         "追问：介入触发条件？介入记录？",
         "人工介入流程、介入记录"),
        ("用户查询/删除/更正请求的处理通道？",
         "追问：线上通道？处理时效？",
         "处理通道说明、处理记录"),
        ("未成年人用户的识别与保护措施？",
         "追问：识别方式？特殊策略？",
         "识别机制说明"),
        ("运营数据是否用于二次用途（营销/分析/再训练）？",
         "追问：用途清单？授权依据？",
         "数据用途说明"),
        ("与外包/第三方客服的数据共享？",
         "追问：外包商名单？脱敏措施？",
         "外包合同、脱敏方案"),
        ("新功能上线的合规评审节点？",
         "追问：评审流程？评审记录？",
         "上线流程、评审记录"),
        ("营销/推荐等个性化功能的授权情况？",
         "追问：单独同意？拒绝选项？",
         "授权记录、功能开关"),
        ("用户投诉的统计与趋势？",
         "追问：投诉集中的问题类型？",
         "投诉台账、统计报告"),
    ],
}

# 本场景必查项关键词（命中任一风险主题即标记必查；用于 01/02 表）
REQUIRED_KEYWORDS = ["个人信息", "数据安全", "内容安全", "算法", "跨境", "未成年人", "权限", "日志", "备份", "删除", "供应商", "授权", "训练数据", "输出"]

# 法律依据映射（2026-08-20 经元典核验的现行有效版本；网安法为 2025 修正版）
# 01 信息调研表：来源表 → 法条（短格式）
SURVEY_LAW_MAP = {
    "表1-1": "暂行办法§7；个保法§52",
    "表1-2": "数安法§27",
    "表1-3": "数安法§21",
    "表1-4": "数安法§21",
    "表1-5": "个保法§5、6、13",
    "表1-6": "个保法§5、6、13、14",
    "表1-7": "个保法§51；数安法§27",
    "表1-8": "网安法§23",
    "表1-9": "个保法§6",
    "表1-10": "个保法§24；暂行办法§12",
    "表1-11": "数安法§27；暂行办法§8",
    "表1-12": "个保法§23",
    "表1-13": "个保法§23",
    "表1-14": "个保法§47",
    "表1-15": "个保法§38；数安法§31",
    "表1-16": "个保法§21",
    "表1-17": "数安法§29",
    "表1-18": "数安法§29",
    "表1-19": "数安法§27",
    "表1-20": "数安法§27",
    "表1-21": "数安法§29；网安法§42",
    "表1-22": "网安法§23；个保法§51",
}

# 02 风险识别表：安全子类关键词 → 法条（短格式；按序首条命中）
LAW_BASIS_RULES = [
    ("收集合法|正当性|合法、诚信|正当、必要", "个保法§5、6、13"),
    ("同意", "个保法§14"),
    ("告知", "个保法§17"),
    ("委托处理|共同处理|向他人提供|第三方|合作方", "个保法§21、23"),
    ("自动化决策", "个保法§24"),
    ("生物特征", "个保法§28、29"),
    ("保存|删除", "个保法§47"),
    ("查阅|更正|可携带|权利", "个保法§45、46"),
    ("影响评估|保护负责人", "个保法§55、52"),
    ("分类分级", "数安法§21"),
    ("资产管理|^数据$", "数安法§21"),
    ("制度体系|制度落实|组织架构|岗位|培训|保密协议|人员录用|转岗离岗", "数安法§27"),
    ("监测预警|应急|威胁和事件|备份恢复", "数安法§29"),
    ("训练数据|预处理|标注", "暂行办法§7、8"),
    ("内容安全|输出", "暂行办法§12、14"),
    ("举报投诉", "个保法§50"),
    ("数据收集设备及环境安全", "数安法§27；网安法§23"),
    ("合作协议约束", "个保法§21、23；数安法§27"),
    ("个人信息保护措施", "个保法§51"),
    ("脱敏", "个保法§51"),
    ("传输链路|网络安全防护|接口|访问控制|身份鉴别|日志留存|审计|云数据|外包人员|第三方接入|存储介质|防泄漏|安全控制|开发运维|授权管理|对外接口|通用规则", "网安法§23"),
]

LAW_ABBREV_NOTE = "法律依据为现行有效法规的条款引用：个保法=《中华人民共和国个人信息保护法》；数安法=《中华人民共和国数据安全法》；网安法=《中华人民共和国网络安全法》（2025修正）；暂行办法=《生成式人工智能服务管理暂行办法》。条款内容经元典平台核验（2026-08-20）。"


def _required_keywords(model: dict | None) -> set[str]:
    """从模型风险清单提炼本场景必查关键词。"""
    if not model:
        return set()
    text = " ".join(
        f"{r.get('theme', '')} {r.get('description', '')}" for r in model.get("risks", [])
    )
    return {kw for kw in REQUIRED_KEYWORDS if kw in text}


# 适用性裁剪规则：intake 答案确认"不涉及"时，相关调研项/评估项标不适用（不物理删除，保留复核痕迹）
APPLICABILITY_RULES = [
    {
        "q_pattern": "BB-04",
        "hit": ["跨境", "港澳", "境外", "出海", "国外"],
        "target01": "表1-15",
        "target02": ["出境", "跨境"],
        "reason": "已确认无跨境业务，数据出境相关项不适用",
    },
    {
        "q_pattern": "DS-03|CD-01",
        "hit": ["生物识别", "人脸", "指纹", "虹膜"],
        "target01": None,
        "target02": ["生物特征"],
        "reason": "已确认不处理生物识别信息，相关项不适用",
    },
    {
        "q_pattern": "BB-03|CH-01|CD-01",
        "hit": ["未成年", "儿童", "学生"],
        "target01": None,
        "target02": ["未成年"],
        "reason": "已确认不面向未成年人，相关项不适用",
    },
]


def _intake_answer_text(intake: dict | None) -> dict[str, str]:
    """question_id → 答案文本（含 asked 标记）。"""
    if not intake:
        return {}
    answers = intake.get("answers", {})
    out = {}
    for qid, entry in answers.items():
        if isinstance(entry, dict) and entry.get("asked") and entry.get("answer"):
            out[qid] = str(entry["answer"])
    return out


def _applicability_map(intake: dict | None) -> dict[str, str]:
    """返回 {target01: 理由} 与 {target02关键词: 理由} 合并字典；无信号/答案缺失时保守不裁剪。"""
    result: dict[str, str] = {}
    answers = _intake_answer_text(intake)
    if not answers:
        return result
    for rule in APPLICABILITY_RULES:
        patterns = rule["q_pattern"].split("|")
        texts = [v for k, v in answers.items() if any(p in k for p in patterns)]
        if not texts:
            continue  # 相关信号问题未答 → 不裁剪
        joined = " ".join(texts)
        if any(kw in joined for kw in rule["hit"]):
            continue  # 答案确认涉及 → 适用
        # 答案明确但不涉及 → 标记不适用
        if rule.get("target01"):
            result[f"01::{rule['target01']}"] = rule["reason"]
        for kw in rule.get("target02", []):
            result[f"02::{kw}"] = rule["reason"]
    return result


def _load_domain_extensions() -> list[dict]:
    """读取领域扩充规则库：优先全局 skill（ai-data-compliance-field-rules），缺失回退本地 assets/。"""
    candidates = [
        Path.home() / ".myagents/projects/mino/.claude/skills/ai-data-compliance-field-rules/assets/field-extensions.json",
        Path(__file__).resolve().parents[1] / "assets/field-extensions.json",
    ]
    for path in candidates:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data.get("extensions", [])
        except (OSError, json.JSONDecodeError):
            continue
    return []


def _match_domain_extensions(model: dict | None, intake: dict | None, extensions: list[dict]) -> list[dict]:
    """按领域关键词匹配：model 元信息 + intake 全部答案文本。"""
    texts = []
    if model:
        meta = model.get("meta", {})
        texts.append(str(meta.get("industry", "")))
        texts.append(str(meta.get("title", "")))
    if intake:
        texts.extend(_intake_answer_text(intake).values())
    joined = " ".join(texts)
    return [e for e in extensions if any(kw in joined for kw in e.get("match", []))]


def _survey_law(table_label: str) -> str:
    """按来源表前缀匹配法律依据。"""
    for prefix, basis in SURVEY_LAW_MAP.items():
        if table_label.startswith(prefix):
            return basis
    return ""


def _subclass_law(subclass: str) -> str:
    """按安全子类关键词匹配法律依据（首条命中）。"""
    import re as re_module
    for pattern, basis in LAW_BASIS_RULES:
        if re_module.search(pattern, subclass):
            return basis
    return ""


# 现场核查动作明细（结构化版，用于 06 现场核查表）
FIELD_SITE_CHECKS_DETAIL = [
    ("系统演示：权限管理", "演示访问控制台：账号创建/审批/回收流程，特权账号清单", "访谈+演示", "确认权限矩阵与制度一致"),
    ("系统演示：日志查看", "演示日志检索：登录/操作/数据导出记录，留存期限", "访谈+演示", "确认日志范围与留存策略"),
    ("系统演示：数据删除", "演示删除流程：用户请求删除/到期删除，含执行记录", "访谈+演示", "确认删除机制可用且留痕"),
    ("日志抽查", "抽取近 30 日日志，核对访问记录与留存范围", "抽样", "抽查结果记录"),
    ("机房/托管环境", "物理访问控制、监控覆盖、进出登记", "现场查看", "如适用；云上环境核对合同"),
    ("安全测试核验", "渗透测试/漏洞扫描报告与修复记录", "文档核验", "核对报告真实性"),
    ("内容安全功能测试", "输入敏感词/违规请求，验证拦截与标识", "现场测试", "确认内容标识方案生效"),
]

# 数据资产盘点清单示例行（客服场景示意，生成后由律师按场景修订；不含状态列，由调用方追加）
DATA_ASSET_SAMPLE_ROWS = [
    ["示例1", "用户对话记录（客服对话/输入输出）", "个人/敏感", "客服场景收集/对话服务", "对话系统/6个月", "无", "用户可删除"],
    ["示例2", "订单与交易信息（订单号/金额/收货）", "个人", "交易处理/订单履约", "交易系统/法定期限", "物流公司", "用户可删除"],
    ["示例3", "人工客服操作日志（操作人/时间/内容）", "内部", "审计追溯/日志系统", "日志系统/3年", "无", "到期删除"],
    ["示例4", "模型训练数据集（脱敏对话/标注）", "内部", "模型训练/训练平台", "训练平台/按项目", "训练服务商", "项目结束删除"],
    ["示例5", "推理服务日志（请求/响应/标识）", "内部", "质量审计/日志系统", "日志系统/90天", "无", "到期删除"],
]


FIELD_CHECKLIST = {
    "一、组织与治理": [
        ("企业基本信息表", "企业全称/统一社会信用代码/法定代表人/数据安全负责人", "原件/盖章件", "访谈索取"),
        ("组织架构与任命文件", "数据安全负责人与部门、职责划分", "盖章件", "访谈索取"),
        ("制度汇编", "数据安全管理制度/分类分级制度/应急预案", "版本化文本", "制度库导出"),
    ],
    "二、数据处理活动": [
        ("数据资产清单（字段级）", "字段名/类型/敏感级别/来源/用途/留存期", "台账导出", "系统导出"),
        ("全生命周期处理活动表", "收集/存储/使用/共享/跨境/删除各环节", "台账/说明", "访谈+导出"),
        ("个人信息授权记录样本", "同意时间/版本/渠道（脱敏）", "脱敏样本", "系统导出"),
    ],
    "三、系统与安全技术": [
        ("系统架构图", "网络拓扑/系统边界/数据流", "版本化图表", "技术访谈"),
        ("安全防护措施清单", "加密/访问控制/备份/防泄漏", "配置清单", "技术访谈"),
        ("日志与审计规范", "日志范围/留存期限/审计安排", "版本化文档", "技术访谈"),
    ],
    "四、模型与算法": [
        ("模型说明", "来源/基座版本/训练数据构成", "版本化文档", "技术访谈"),
        ("备案与安全评估材料", "备案编号/评估报告/标识方案", "官方回执+报告", "法务访谈"),
        ("内容安全方案", "敏感词/水印/内容标识/人工兜底", "方案文档", "技术访谈"),
    ],
    "五、个人信息保护": [
        ("隐私政策与用户协议", "现行版本与最近更新时间", "官网版本存档", "法务访谈"),
        ("PIA报告", "评估范围/结论/整改闭环", "报告原件", "数据安全访谈"),
        ("用户权利请求处理记录", "查询/更正/删除的响应记录（脱敏）", "台账", "业务访谈"),
    ],
    "六、第三方与供应商": [
        ("供应商清单", "模型/云服务/外包客服全量清单", "清单台账", "访谈索取"),
        ("DPA/委托处理合同", "含数据范围与责任边界条款", "合同复印件", "法务访谈"),
        ("外部API与云服务合同", "数据流向与服务条款", "合同复印件", "法务访谈"),
    ],
    "七、日志与退出": [
        ("日志留存策略与样本", "留存范围/期限/访问权限", "策略文档+样本", "系统导出"),
        ("数据删除/退出机制文档", "到期删除/用户删除的执行流程", "流程文档", "业务访谈"),
        ("删除执行记录", "近期删除任务的执行日志（脱敏）", "执行记录", "系统导出"),
    ],
    "八、合规记录": [
        ("历史投诉与整改记录", "投诉类型/数量/整改闭环", "台账", "法务访谈"),
        ("监管往来函件", "询问/检查/整改要求及回复", "函件存档", "法务访谈"),
        ("员工培训记录", "数据安全培训计划与签到", "记录台账", "数据安全访谈"),
        ("应急预案与演练记录", "预案版本/演练/复盘", "文档+记录", "数据安全访谈"),
    ],
}

FIELD_SITE_CHECKS = [
    ("系统演示", "权限管理、日志查看、数据删除流程实际操作", "演示记录"),
    ("日志抽查", "抽取近期日志核对留存范围与访问记录", "抽查记录"),
    ("机房/托管环境", "物理访问控制、监控、进出登记（如适用）", "现场照片+登记"),
    ("安全测试核验", "渗透测试/漏扫报告与修复记录核对", "报告核对"),
]


def build_field_pack(project: dict, handbook: dict, output_dir: Path, model: dict | None = None, intake: dict | None = None) -> list[Path]:
    """第四步：以《生成式AI行业网络数据安全风险评估工作手册》为指引，生成进企调研前全套文件。

    handbook 结构：
      survey_items: list[dict]（survey_id/table_id/section/table_label/original_text/pdf_page）
      evaluation_indicators: list[dict]（指标ID/来源表ID/安全子类/原文评估文本/评估方法/判定依据/预期证据/来源模块/PDF页码）
    model（可选）：industry-model.json，用于向访谈提纲注入本场景重点核查问题（risks/verificationItems）。
    intake（可选）：00_project/intake.json，用于适用性裁剪（不涉及项标不适用）与领域扩充匹配。
    """
    project_name = project["project_name"]
    output_dir.mkdir(parents=True, exist_ok=True)
    built: list[Path] = []

    def new_field_doc(title: str, doc_no: str) -> Document:
        return new_document(project, f"{project_name}\n{title}", title, f"FIELD-{doc_no}", True)

    # 适用性裁剪与领域扩充（一次计算，01/02 共用）
    applicability = _applicability_map(intake)
    extensions = _match_domain_extensions(model, intake, _load_domain_extensions())
    domain_names = "、".join(e["domain"] for e in extensions) or ""

    # 1. 信息调研表（手册附录1，survey_items 107 项）
    doc = new_field_doc("信息调研表（手册附录1）", "01")
    add_heading(doc, "一、用途与边界")
    add_paragraph(doc, "本表依据《生成式人工智能行业网络数据安全风险评估工作手册》附录1编制，用于进入企业实地调研前的信息收集。全部调研项须以客户事实和证据判断为准。")
    add_heading(doc, "二、调研项清单")
    required = _required_keywords(model)
    survey = handbook.get("survey_items", [])
    rows = []
    for index, item in enumerate(survey, 1):
        table_label = item.get("table_label", "")
        text = f"{table_label} {item.get('original_text', '')}"
        mark = "★必查" if any(kw in text for kw in required) else ""
        status = applicability.get(f"01::{table_label[:6]}", "待调研")
        rows.append([
            item.get("survey_id", ""),
            table_label,
            item.get("original_text", ""),
            _survey_law(table_label),
            mark,
            status,
        ])
    if not rows:
        rows = [["", "暂无调研项", "", "", "", "待调研"]]
    add_table(doc, ["编号", "来源表", "调研内容", "法律依据", "本场景必查", "状态"], rows, [600, 1500, 3800, 1500, 700, 700])
    # 领域专项调研项
    if extensions:
        add_heading(doc, f"三、领域专项调研项（{domain_names}）")
        add_paragraph(doc, f"以下调研项由「{domain_names}」领域规则库自动追加，对应领域专项法规，进企调研时与通用项一并完成。")
        rows = []
        for index, item in enumerate(extensions[0].get("survey_items", []), 1):
            rows.append([f"X{index:02d}", item.get("content", ""), item.get("law", ""), item.get("note", ""), "待调研"])
        add_table(doc, ["编号", "专项调研内容", "法律依据", "说明", "状态"], rows, [700, 4000, 1700, 1500, 700])
        regs = "；".join(f"{r.get('name', '')}（{r.get('issuer', '')}，{r.get('status', '待核验')}）" for r in extensions[0].get("regulations", []))
        add_paragraph(doc, f"领域法规：{regs}。标注「待核验」的法规条款须经专项检索核验后引用。")
        add_heading(doc, "四、填写指引")
    else:
        add_heading(doc, "三、填写指引")
    add_paragraph(doc, "按调研表逐项访谈/查证后填写；注明证据编号与来源（手册附录1.1-1.5）。标注★必查的调研项由本项目风险清单推导，进企优先核查；标注「不适用」的调研项保留备查，如后续事实变化可恢复适用。")
    add_paragraph(doc, LAW_ABBREV_NOTE)
    path = output_dir / "01-信息调研表.docx"
    save_docx(doc, path)
    built.append(path)

    # 2. 风险识别表（手册附录2，evaluation_indicators 319 项）
    doc = new_field_doc("风险识别表（手册附录2）", "02")
    add_heading(doc, "一、用途与边界")
    add_paragraph(doc, "本表依据手册附录2编制，从数据安全管理、数据处理活动、数据安全技术、个人信息保护四方面逐项识别风险。评估方法：人员访谈/文档查验/安全核查/技术测试。")
    add_heading(doc, "二、评估指标清单")
    indicators = handbook.get("evaluation_indicators", [])
    rows = []
    for index, item in enumerate(indicators, 1):
        subclass = item.get("安全子类", "")
        text = f"{subclass} {item.get('原文评估文本', '')}"
        mark = "★必查" if any(kw in text for kw in required) else ""
        status = "待评估"
        for rule in APPLICABILITY_RULES:
            for kw in rule.get("target02", []):
                if kw in subclass:
                    status = applicability.get(f"02::{kw}", "待评估")
                    break
            if status != "待评估":
                break
        rows.append([
            item.get("指标ID", ""),
            subclass,
            item.get("原文评估文本", ""),
            item.get("评估方法", ""),
            _subclass_law(subclass),
            mark,
            status,
        ])
    if not rows:
        rows = [["", "", "暂无评估指标", "", "", "", "待评估"]]
    add_table(doc, ["编号", "子类", "评估项", "评估方法", "法律依据", "本场景必查", "状态"], rows, [600, 1000, 3200, 1000, 1400, 700, 700])
    # 领域专项评估指标
    if extensions:
        add_heading(doc, f"三、领域专项评估指标（{domain_names}）")
        add_paragraph(doc, f"以下指标由「{domain_names}」领域规则库自动追加，进企访谈与现场核查中与通用指标一并评估。")
        rows = []
        for index, item in enumerate(extensions[0].get("indicator_items", []), 1):
            rows.append([f"X{index:02d}", item.get("subclass", ""), item.get("content", ""), item.get("method", ""), item.get("law", ""), "待评估"])
        add_table(doc, ["编号", "子类", "专项评估项", "评估方法", "法律依据", "状态"], rows, [700, 1400, 3300, 1200, 1400, 700])
        add_heading(doc, "四、判定要求")
    else:
        add_heading(doc, "三、判定要求")
    add_paragraph(doc, "每项指标须结合客户事实和证据给出符合/不符合/部分符合判定；默认全部为待评估。标注★必查的指标由本项目风险清单推导，访谈与现场核查中优先确认；标注「不适用」的指标保留备查，事实变化可恢复。")
    add_paragraph(doc, LAW_ABBREV_NOTE)
    path = output_dir / "02-风险识别表.docx"
    save_docx(doc, path)
    built.append(path)

    # 3. 访谈提纲（手册方法：人员访谈）
    doc = new_field_doc("访谈提纲（手册评估方法）", "03")
    add_heading(doc, "一、访谈对象与方法")
    add_paragraph(doc, "按手册风险识别方法（人员访谈/文档查验/安全核查/技术测试）组织访谈；访谈对象覆盖管理层、数据安全、技术、法务、业务等角色。每类问题均给出追问要点与证据锚点，访谈记录须注明时间、地点、访谈人与被访谈人。")
    add_heading(doc, "二、分角色访谈问题")
    add_paragraph(doc, "问题按角色分组，供访谈前按对象选用；同一问题可由多个角色从不同角度交叉印证。")
    seq = 0
    for role, questions in INTERVIEW_QUESTIONS.items():
        add_heading(doc, role, level=2)
        rows = []
        for question, probe, evidence in questions:
            seq += 1
            rows.append([str(seq), question, probe, evidence, ""])
        add_table(doc, ["序号", "问题", "追问要点", "证据锚点", "记录"], rows, [700, 2400, 2300, 2000, 1200])
    # 模型注入：本场景重点核查问题
    if model:
        injected = []
        for r in model.get("risks", []):
            injected.append((r.get("id", ""), f"针对风险「{r.get('theme', '')}」核查：{r.get('description', '')}。追问：核实该风险描述中的事实主张是否成立，现状与应对措施。", r.get("certainty", "待核实"), r.get("level", "中")))
        for v in model.get("verificationItems", []):
            injected.append((v.get("id", ""), f"待核验：{v.get('question', '')}", f"对象：{v.get('targetRole', '')}", ""))
        if injected:
            add_heading(doc, "三、本场景重点核查问题（由模型自动生成）")
            add_paragraph(doc, "以下问题来自本项目的风险清单与待核验事项，进企访谈时优先完成核查；结论须回到风险识别表对应条目。")
            rows = [[item[0], item[1], item[2], item[3], ""] for item in injected]
            add_table(doc, ["编号", "核查问题", "确定度/对象", "等级", "记录"], rows, [1100, 3400, 1500, 900, 800])
            add_heading(doc, "四、访谈记录表")
        else:
            add_heading(doc, "三、访谈记录表")
    else:
        add_heading(doc, "三、访谈记录表")
    roles = list(INTERVIEW_QUESTIONS.keys())  # 与分角色访谈问题一致（5 角色）
    rows = [[str(i), role, "数据处理活动、安全措施、责任落实情况", ""] for i, role in enumerate(roles, 1)]
    add_table(doc, ["序号", "访谈对象", "访谈重点", "记录/证据"], rows, [700, 2000, 4200, 1800])
    add_heading(doc, "五、访谈纪律" if model and injected else "四、访谈纪律")
    add_paragraph(doc, "访谈记录须注明时间、地点、访谈人与被访谈人；涉及客户事实的陈述标注'访谈陈述'，与书面材料一致性的差异记录为'存在矛盾'；每类问题收集的证据按《进企前文件与证据清单》编号归档。")
    path = output_dir / "03-访谈提纲.docx"
    save_docx(doc, path)
    built.append(path)

    # 4. 进企前文件清单（手册附录4 报告 + 证据附件）
    doc = new_field_doc("进企前文件与证据清单", "04")
    add_heading(doc, "一、文件清单（按主题分类）")
    add_paragraph(doc, "进企调研前向企业索取；每项注明证据标准与获取渠道，索取时以《信息调研表》对应条目为准。")
    seq = 0
    for category, items in FIELD_CHECKLIST.items():
        add_heading(doc, category, level=2)
        rows = []
        for name, points, standard, channel in items:
            seq += 1
            rows.append([str(seq), name, points, standard, channel, "待收集"])
        add_table(doc, ["序号", "文件", "内容要点", "证据标准", "获取渠道", "状态"], rows, [700, 1800, 3000, 1200, 1000, 900])
    add_heading(doc, "二、现场核查动作")
    add_paragraph(doc, "进企实地调研期间同步开展以下核查，核查结果记录于风险识别表相应条目。")
    rows = [[str(i), name, points, ""] for i, (name, points, _) in enumerate(FIELD_SITE_CHECKS, 1)]
    add_table(doc, ["序号", "核查动作", "核查要点", "结果记录"], rows, [700, 1600, 4600, 1700])
    add_heading(doc, "三、证据要求")
    add_paragraph(doc, "全部证据须可追溯：注明证据编号、文件哈希、责任人与保管位置；无法提供的标注'待补充'并说明原因。访谈陈述类证据单独编号并注明访谈时间与对象，与书面证据存在差异时在风险识别表中标记。")
    path = output_dir / "04-进企前文件清单.docx"
    save_docx(doc, path)
    built.append(path)

    # 5. 数据资产盘点清单（字段级，进企第一动作）
    doc = new_field_doc("数据资产盘点清单（字段级）", "05")
    add_heading(doc, "一、用途与边界")
    add_paragraph(doc, "进企调研第一阶段完成数据资产盘点，作为信息调研表与风险识别表的事实基础。按'库/表/文件'逐项登记，字段级信息用于判断敏感级别与处理活动。")
    add_heading(doc, "二、盘点清单")
    rows = [list(row) + ["待盘点"] for row in DATA_ASSET_SAMPLE_ROWS]
    add_table(doc, ["编号", "数据资产与字段", "敏感级别", "来源与目的", "存储位置/期限", "共享对象", "删除机制", "状态"], rows, [600, 2000, 800, 1900, 1200, 900, 900, 600])
    add_heading(doc, "三、填写指引")
    add_paragraph(doc, "前五行为示例（客服场景示意），盘点时删除并替换为实际资产；敏感级别按分类分级制度认定；来源与目的须具体到系统与业务环节；删除机制注明'到期自动/人工/用户可申请'。")
    path = output_dir / "05-数据资产盘点清单.docx"
    save_docx(doc, path)
    built.append(path)

    # 6. 现场核查表（系统演示/日志抽查/安全测试）
    doc = new_field_doc("现场核查表", "06")
    add_heading(doc, "一、用途与边界")
    add_paragraph(doc, "进企实地调研期间使用，与访谈提纲、信息调研表配套；核查结论回到风险识别表对应条目。")
    add_heading(doc, "二、核查清单")
    rows = [[str(i), name, points, method, "待核查", "", ""] for i, (name, points, method, _) in enumerate(FIELD_SITE_CHECKS_DETAIL, 1)]
    add_table(doc, ["序号", "核查动作", "核查要点", "核查方式", "判定", "发现与说明", "证据编号"], rows, [600, 1700, 2300, 1100, 700, 1300, 900])
    add_heading(doc, "三、判定要求")
    add_paragraph(doc, "每项给出符合/不符合/部分符合判定并记录证据编号；不符合项须当场记录影响范围与初步整改建议，事后回到风险识别表登记。")
    path = output_dir / "06-现场核查表.docx"
    save_docx(doc, path)
    built.append(path)

    # 7. 保密承诺与进场授权函（两份文书模板）
    doc = new_field_doc("保密承诺与进场授权函", "07")
    add_heading(doc, "一、企业保密承诺函（模板）")
    add_paragraph(doc, "致：【XX律师事务所/评估机构】")
    add_paragraph(doc, "鉴于贵所接受【本企业】委托开展AI数据合规风险评估工作，我方已充分知悉在评估过程中将接触包括个人信息、商业秘密、未公开经营数据在内的敏感信息，现郑重承诺如下：")
    add_paragraph(doc, "（一）对评估过程中接触的全部信息严格保密，未经贵所书面同意，不向任何第三方披露、提供或用于本评估之外的任何目的；")
    add_paragraph(doc, "（二）评估人员仅在完成评估所必需的最小范围内接触信息，接触人员名单于评估前向贵所书面报备；")
    add_paragraph(doc, "（三）评估结束后按贵所要求归还或销毁持有的全部材料副本（含电子数据）；")
    add_paragraph(doc, "（四）因我方原因造成信息泄露的，依法承担相应责任。")
    add_paragraph(doc, "本承诺函自盖章之日起生效，长期有效。")
    add_paragraph(doc, "企业名称（盖章）：【　　　　　　　　】")
    add_paragraph(doc, "授权代表签字：　　　　　　　　日期：　　年　　月　　日")
    add_heading(doc, "二、进场调查授权书（模板）")
    add_paragraph(doc, "兹授权【XX律师事务所/评估机构】及其指派人员（名单见附件），自【　　年　　月　　日】至【　　年　　月　　日】期间进入【本企业】开展AI数据合规实地调研，具体授权范围包括：")
    add_paragraph(doc, "（一）查阅与数据处理活动相关的制度文件、系统文档与记录（涉密内容按保密承诺函处理）；")
    add_paragraph(doc, "（二）安排相关部门负责人及经办人员接受访谈；")
    add_paragraph(doc, "（三）在技术人员配合下进行系统演示与必要的现场核查；")
    add_paragraph(doc, "（四）收取评估所需文件与证据材料（复印件或电子件）。")
    add_paragraph(doc, "请各相关部门予以配合。")
    add_paragraph(doc, "企业名称（盖章）：【　　　　　　　　】")
    add_paragraph(doc, "法定代表人/授权代表签字：　　　　　　　　日期：　　年　　月　　日")
    add_heading(doc, "三、使用指引")
    add_paragraph(doc, "两份文书为模板，发出前须经律所复核并替换【】占位内容；进场人员名单作为授权书附件一并出具。")
    path = output_dir / "07-保密承诺与进场授权函.docx"
    save_docx(doc, path)
    built.append(path)

    return built
