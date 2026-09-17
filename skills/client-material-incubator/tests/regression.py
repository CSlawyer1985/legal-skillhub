# -*- coding: utf-8 -*-
"""
回归测试：一条命令跑全部校验，非 0 返回码即阻断发布。

覆盖两个方向（对应测试报告 P0 修复）：
  方向1 母版化（原文 → {{变量}}）：T9 格式零改动 + run 下划线绑定 + 变量全写入 + 真实值零残留
  方向2 生成（{{变量}} → 真实值）：T9 格式零改动 + run 下划线绑定 + 变量零残留 + 真实值全写入

用法:
  python tests/regression.py          # 全量回归，返回码非 0 即失败
"""
import sys
import os
import json
import re
import zipfile
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(HERE, "..", "scripts")
sys.path.insert(0, SCRIPTS)

import build_master as bm


# ---------------------------------------------------------------------------
# 黄金样例生成（python-docx 内联生成，可复现、不依赖提交二进制 fixture）
# ---------------------------------------------------------------------------

def make_fixture_retainer(tmpdir):
    """委托协议样例：抬头 tab 分栏 + 填空位下划线 + 跨节点日期 + 勾选框。"""
    import docx
    doc = docx.Document()

    p1 = doc.add_paragraph()  # tab 分栏（左右两栏抬头）
    p1.add_run("甲方：张伟")
    p1.add_run("\t")
    p1.add_run("乙方：浙江示范律师事务所")

    p2 = doc.add_paragraph()  # 填空位下划线（律师姓名）
    p2.add_run("承办律师：")
    r = p2.add_run("李强")
    r.font.underline = True

    p3 = doc.add_paragraph()  # 跨节点日期，部分下划线
    p3.add_run("签订日期：")
    r = p3.add_run("2026"); r.font.underline = True
    p3.add_run("年")
    p3.add_run("5")
    p3.add_run("月")
    p3.add_run("13"); p3.runs[-1].font.underline = True
    p3.add_run("日")

    p4 = doc.add_paragraph()  # 勾选框
    p4.add_run("☑ 民事诉讼　□ 行政诉讼")

    path = os.path.join(tmpdir, "委托协议_样例.docx")
    doc.save(path)
    return path


def make_fixture_poa(tmpdir):
    """授权委托书样例：两套（已填示例 + 空白手填）+ 填空位下划线。"""
    import docx
    doc = docx.Document()

    doc.add_paragraph("授权委托书")
    doc.add_paragraph("委托人：张三")
    doc.add_paragraph("授权委托书")
    p = doc.add_paragraph()
    p.add_run("委托人：")
    r = p.add_run("")
    r.font.underline = True

    path = os.path.join(tmpdir, "授权委托书_样例.docx")
    doc.save(path)
    return path


def make_fixture_overlap(tmpdir):
    """相邻填空位区间重叠：A 结尾与 B 开头共享「预付」，长词优先后短规则被跳过（v3.3.2 缺陷回归）。"""
    import docx
    doc = docx.Document()
    p = doc.add_paragraph()
    # 拆成两个 run，让两个 old 都跨节点（Pass1 不命中、走 Pass3 精确区间替换）
    p.add_run("□于本合同生效之当日预")
    p.add_run("付___/___元")
    path = os.path.join(tmpdir, "重叠填空位_样例.docx")
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 校验函数
# ---------------------------------------------------------------------------

def read_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def strip_t(xml):
    return re.sub(bm.WT_RE + r"[^<]*</w:t>", "", xml)


def check_t9(src, out):
    return strip_t(read_xml(src)) == strip_t(read_xml(out))


def check_underline_binding(src, out, mapping):
    return bm.verify_run_binding(src, out, mapping) == []


def missing_new(out, mapping):
    """返回 out 中缺失的 new 列表（空 = 全部写入）。"""
    xml = read_xml(out)
    return [m["new"] for m in mapping if m["new"] not in xml]


def has_residual_braces(out):
    """生成方向：成品是否仍含 {{ }} 变量占位符。"""
    xml = read_xml(out)
    return "{{" in xml or "}}" in xml


def find_residual_words(out, words):
    """返回 out 中仍残留的敏感词列表（空 = 零残留）。"""
    xml = read_xml(out)
    return [w for w in words if w in xml]


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def main():
    tmpdir = tempfile.mkdtemp(prefix="regression_")
    cases = [
        {
            "name": "委托协议样例",
            "src": make_fixture_retainer(tmpdir),
            "mapping": [
                {"old": "张伟", "new": "{{委托人姓名}}"},
                {"old": "李强", "new": "{{承办律师}}"},
                {"old": "2026年5月13日", "new": "{{签订日期}}"},
            ],
            "sensitive": ["张伟", "李强"],
        },
        {
            "name": "授权委托书样例",
            "src": make_fixture_poa(tmpdir),
            "mapping": [
                {"old": "张三", "new": "{{委托人姓名}}"},
            ],
            "sensitive": ["张三"],
        },
    ]

    all_pass = True
    print("=" * 70)
    print("委托材料孵化器 · 回归测试（母版化 + 生成 双方向）")
    print("=" * 70)

    for c in cases:
        map_path = os.path.join(tmpdir, c["name"] + ".json")
        with open(map_path, "w", encoding="utf-8") as f:
            json.dump(c["mapping"], f, ensure_ascii=False)

        # ---- 方向1：母版化（原文 → {{变量}}）----
        out_master = os.path.join(tmpdir, c["name"] + "_母版.docx")
        bm.build(c["src"], map_path, out_master, verify=False)

        t9 = check_t9(c["src"], out_master)
        ub = check_underline_binding(c["src"], out_master, c["mapping"])
        miss = missing_new(out_master, c["mapping"])
        resid = find_residual_words(out_master, c["sensitive"])

        print(f"\n【{c['name']} · 方向1 母版化】")
        print(f"  T9 格式零改动:      {'✅' if t9 else '❌'}")
        print(f"  run 下划线绑定:     {'✅' if ub else '❌'}")
        print(f"  变量全写入:         {'✅' if not miss else '❌ 缺 ' + ','.join(miss)}")
        print(f"  真实值零残留:       {'✅' if not resid else '❌ 残留 ' + ','.join(resid)}")
        ok1 = t9 and ub and not miss and not resid

        # ---- 方向2：生成（{{变量}} → 真实值）----
        gen_mapping = [{"old": m["new"], "new": m["old"]} for m in c["mapping"]]
        gen_map_path = os.path.join(tmpdir, c["name"] + "_gen.json")
        with open(gen_map_path, "w", encoding="utf-8") as f:
            json.dump(gen_mapping, f, ensure_ascii=False)
        out_gen = os.path.join(tmpdir, c["name"] + "_成品.docx")
        bm.build(out_master, gen_map_path, out_gen, verify=False)

        t9g = check_t9(out_master, out_gen)
        ubg = check_underline_binding(out_master, out_gen, gen_mapping)
        braces = has_residual_braces(out_gen)
        missg = missing_new(out_gen, gen_mapping)

        print(f"\n【{c['name']} · 方向2 生成】")
        print(f"  T9 格式零改动:      {'✅' if t9g else '❌'}")
        print(f"  run 下划线绑定:     {'✅' if ubg else '❌'}")
        print(f"  变量零残留:         {'✅' if not braces else '❌'}")
        print(f"  真实值全写入:       {'✅' if not missg else '❌ 缺 ' + ','.join(missg)}")
        ok2 = t9g and ubg and not braces and not missg

        if not (ok1 and ok2):
            all_pass = False

    # ---- v3.3.2 缺陷回归（dry-run 假阳性 / 重叠告警 / 跨节点降噪）----
    print("\n【v3.3.2 缺陷回归】")

    # ① dry-run 假阳性：从未出现的 old 必须判未命中（而非「不残留=命中」的假阳性）
    xml1 = read_xml(make_fixture_retainer(tmpdir))
    _, stats1 = bm.replace_in_document_xml(xml1, [
        {"old": "张伟", "new": "{{委托人姓名}}"},
        {"old": "不存在的文本XYZ123", "new": "{{幽灵变量}}"},
    ], None)
    ghost_unhit = "不存在的文本XYZ123" in stats1["unhit"]
    ghost_hit = "不存在的文本XYZ123" in stats1["hit_olds"]
    ok_ghost = ghost_unhit and not ghost_hit
    print(f"  ① dry-run 假阳性消除（从未出现→未命中）: {'✅' if ok_ghost else '❌'}")
    if not ok_ghost:
        all_pass = False

    # ② 相邻填空位区间重叠：短规则跳过并告警、其变量缺失可检出
    xml2 = read_xml(make_fixture_overlap(tmpdir))
    _, stats2 = bm.replace_in_document_xml(xml2, [
        {"old": "□于本合同生效之当日预付", "new": "{{预付款项}}"},
        {"old": "预付___/___元", "new": "{{预付款额}}"},
    ], None)
    skipped2 = stats2["skipped_overlaps"]
    short_unhit = "预付___/___元" in stats2["unhit"]
    ok_overlap = len(skipped2) >= 1 and short_unhit
    print(f"  ② 重叠跳过告警 + 短规则缺失检出（跳过 {len(skipped2)} 条）: {'✅' if ok_overlap else '❌'}")
    if not ok_overlap:
        all_pass = False

    # ③ 跨节点提醒降噪：多个含年月日 old 只产生 1 条汇总（而非逐条刷屏）
    blockers3, warnings3 = bm.check_mapping_conflicts([
        {"old": "2026年5月13日", "new": "{{日期}}"},
        {"old": "2026年8月1日", "new": "{{日期2}}"},
        {"old": "2026年12月20日", "new": "{{日期3}}"},
        {"old": "张三、李四", "new": "{{主体}}"},
    ])
    cross3 = [w for w in warnings3 if "跨 <w:t>" in w]
    ok_cross = len(cross3) <= 1
    print(f"  ③ 跨节点提醒合并为一行（{len(cross3)} 条）: {'✅' if ok_cross else '❌'}")
    if not ok_cross:
        all_pass = False

    print("\n" + "=" * 70)
    if all_pass:
        print("回归测试全部通过 ✅")
    else:
        print("回归测试存在失败 ❌（阻断发布）")
    print("=" * 70)
    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
