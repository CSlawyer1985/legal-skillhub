# -*- coding: utf-8 -*-
"""大事记表交付物自检。判据与规格一一对应，任何一条不成立即拒绝交付。"""
import zipfile, re, sys, json
from openpyxl import load_workbook
FAIL=[]
def ck(no,cond,msg):
    print(("  OK  " if cond else "  FAIL")+f" {no}  {msg}")
    if not cond: FAIL.append(no)
def dec(s): return re.sub(r'&#(\d+);',lambda m:chr(int(m.group(1))),s)

import sys as _s
DOCX=_s.argv[1] if len(_s.argv)>1 else "案件大事记.docx"
XLSX=_s.argv[2] if len(_s.argv)>2 else "案件大事记母表.xlsx"
RI=_s.argv[3] if len(_s.argv)>3 else "渲染输入.json"
z=zipfile.ZipFile(DOCX)
st=z.read("word/styles.xml").decode(); dc=z.read("word/document.xml").decode()
D=json.load(open(RI,encoding="utf-8"))
N=len(D["rows"])

print("— Word 交付版 —")
ck("W1", 'w:ascii="Times New Roman"' in st and 'w:eastAsia="宋体"' in st, "默认字体三槽：英文数字新罗马、中文宋体")
ck("W2", '<w:autoSpaceDE w:val="0"/>' in st and '<w:autoSpaceDN w:val="0"/>' in st, "autoSpace 置 0：中英文之间不自动加空")
ck("W3", st.index('<w:autoSpaceDE') < st.index('<w:spacing'), "CT_PPrBase 次序：autoSpace 在 spacing 之前")
rf=re.findall(r'<w:rFonts[^>]*?/>',dc)
ck("W4", all('w:hint="eastAsia"' in x for x in rf), f"全部 {len(rf)} 处 rFonts 带 hint=eastAsia（引号、省略号、书名号归中文）")
ck("W5", '<w:sz w:val="23"/>' in dc or 'w:val="23"' in st, "表内字号 11.5 磅")
ck("W6", '<w:sz w:val="44"/>' in dc, "标题二号字")
ck("W7", dc.count('方正小标宋简体')==1, f"只有标题用小标宋，表内一律宋体（实得 {dc.count('方正小标宋简体')}）")
mar=re.search(r'<w:pgMar[^>]*>',dc).group(0)
vals={k:int(v) for k,v in re.findall(r'w:(top|bottom|left|right)="(\d+)"',mar)}
ck("W8", set(vals.values())=={1360}, f"页边距四周 2.4cm = 1360 DXA（实得 {vals}）")
grid=re.search(r'<w:tblGrid>(.*?)</w:tblGrid>',dc,re.S)
ws=[int(x) for x in re.findall(r'w:w="(\d+)"',grid.group(1))]
ck("W9", sum(ws)==9186, f"列宽合计等于可用宽度 9186（实得 {sum(ws)}：{ws}）")
ck("W10", ws[1]>=1400, f"日期列不小于 1400，十位日期不拆行（实得 {ws[1]}）")
ck("W11", '<w:numPr>' not in dc, "序号是静态数字，不用自动编号")

def _cells(xml):
    """把 Word 表格拆成逐行逐格的纯文本，供内容判据用。"""
    rows=[]
    for tr in re.findall(r"<w:tr[ >].*?</w:tr>", xml, re.S):
        rows.append(["".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", tc))
                     for tc in re.findall(r"<w:tc>.*?</w:tc>", tr, re.S)])
    return rows
_tb=[r for r in _cells(dc) if len(r)>=5][1:]          # 去掉表头行
ck("W16", [r[0] for r in _tb]==[str(i+1) for i in range(N)],
   f"序号列是 1 起的连号（实得 {[r[0] for r in _tb][:5]}）")
_need=[k for k,r in enumerate(D["rows"]) if [p for p in (r.get("note_parts") or []) if str(p).strip()]]
ck("W17", all(_tb[k][4].strip() for k in _need if k < len(_tb)),
   f"渲染输入里有备注的行，Word 备注列不得为空（应有 {len(_need)} 行）")
ck("W12", len(re.findall(r'<w:b/>',dc))==N, f"加黑只在事项列、每行一次（实得 {len(re.findall(r'<w:b/>',dc))}，应为 {N}）")
ck("W13", dc.count('<w:top w:type="dxa" w:w="115"/>')>=N and dc.count('<w:bottom w:type="dxa" w:w="115"/>')>=N,
   "单元格上下边距 0.5 行 = 115 DXA，内文之间无段前段后")
ck("W14", '<w:color w:val="000000"' in dc or 'w:color="000000"' in dc or dc.count('w:val="000000"')>0, "表格框线黑色")
jc=dc.count('<w:jc w:val="both"/>')
ck("W15", jc>=2*N, f"事项与主要内容两端对齐（实得 {jc} 处）")

print("— Excel 母表 —")
sh=dec(zipfile.ZipFile(XLSX).read("xl/worksheets/sheet1.xml").decode())
sx=zipfile.ZipFile(XLSX).read("xl/styles.xml").decode()
wb=load_workbook(XLSX, rich_text=True); ws2=wb.active
# 列头行不写死：母表上方加了案件名标题之后，列头就不在第 1 行了。
HR = next((r for r in range(1, 6) if ws2.cell(row=r, column=1).value == "序号"), 1)
ck("X1", [c.value for c in ws2[HR]][:5]==["序号","日期","主体","事项","主要内容"], "前五列字段固定")
ck("X2", all(ws2.cell(row=i,column=1).value==f"=ROW()-{HR}" for i in range(HR+1,HR+1+N)),
   f"序号是公式且随列头行算（=ROW()-{HR}）")
ck("X3", len(re.findall(rf'<f>ROW\(\)-{HR}</f><v>\d+</v>',sh))==N,
   f"公式带缓存值且偏移随列头行（ROW()-{HR}）")
ck("X4", ws2.freeze_panes==f"B{HR+1}" and ws2.auto_filter.ref is not None,
   f"冻结列头行与编号列（B{HR+1}）且可筛选")
ck("X5", '方正小标宋简体' in sx, "表头字体保住小标宋，未被替换")
fonts=sorted({x for x in re.findall(r'<rFont val="([^"]+)"/>',sh)})
ck("X6", fonts==["Times New Roman","宋体"], f"富文本只用宋体与新罗马（实得 {fonts}）")
def _date_ok(cell):
    v=cell.value
    if hasattr(v,"__iter__") and not isinstance(v,str):      # 富文本：数字段须走新罗马
        return any(getattr(getattr(b,"font",None),"rFont",None)=="Times New Roman" for b in v if hasattr(b,"font"))
    return cell.font.name=="Times New Roman"                  # 纯数字格
ck("X7", all(_date_ok(ws2.cell(row=i,column=2)) for i in range(HR+1,HR+1+N)),
   "日期格的数字部分走新罗马（仅到月/区间档带中文标注，按混排判）")
txt=re.sub(r'<[^>]+>','',sh)
bad=re.findall(r'[\u4e00-\u9fff][ \u00a0]+[A-Za-z0-9]|[A-Za-z0-9][ \u00a0]+[\u4e00-\u9fff]',txt)
ck("X8", not bad, f"中英文之间无空格（实得 {bad[:3]}）")
ck("X9", all(ws2.cell(row=i,column=11).value in ("有书证","仅当事人陈述") for i in range(HR+1,HR+1+N)), "每行证据状态非空")
ck("X10", all(ws2.cell(row=i,column=6).value for i in range(HR+1,HR+1+N)), "每行都挂来源，一条不漏")
# 底稿有主体的行，母表第 3 列就得有主体。这一列曾经整列是空的，
# 因为 to_render 压根不产 parties，而当时没有任何判据看着它。
_want=[k for k,r in enumerate(D["rows"]) if str(r.get("parties","")).strip()]
ck("X11", all(ws2.cell(row=HR+1+k,column=3).value for k in _want),
   f"渲染输入里有主体的行，母表主体列不得为空（应有 {len(_want)} 行）")

print()
print("全部通过" if not FAIL else f"未通过：{FAIL}")
sys.exit(1 if FAIL else 0)
