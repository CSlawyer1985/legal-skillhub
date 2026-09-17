#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Robust DOCX form filling for case-archiver V3.

V3 prefers stable selectors over raw table coordinates:

* ``marker``: replace ``{{field}}`` or another explicit token, even when the
  token is split across Word runs.
* ``control_tag``: fill a Word content control by tag/title.
* ``selector.type=cell_anchor``: find a labelled cell and target the same,
  right-hand, or lower cell.
* ``anchor`` / ``loc``: V2-compatible paragraph and coordinate fallbacks.

The module writes to a temporary file, verifies the saved document, and only
replaces the requested output when all provided fields pass in strict mode.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


class TemplateFillError(RuntimeError):
    """Raised when one or more non-empty fields cannot be filled safely."""

    def __init__(self, message, report=None):
        super().__init__(message)
        self.report = report or {}


def _copy_rpr(src_run, dst_run):
    src_rpr = src_run._element.find(qn("w:rPr"))
    if src_rpr is None:
        return
    dst_rpr = dst_run._element.find(qn("w:rPr"))
    if dst_rpr is None:
        dst_rpr = OxmlElement("w:rPr")
        dst_run._element.insert(0, dst_rpr)
    for child in src_rpr:
        dst_rpr.append(deepcopy(child))


def _first_run(para):
    return para.runs[0] if para.runs else None


def _norm(text):
    return re.sub(r"[\s\u3000:：_（）()【】\[\]]+", "", str(text or "")).lower()


def _unique_row_cells(row):
    """Return row cells once per underlying XML cell (merged aliases removed)."""
    result = []
    seen = set()
    for grid_index, cell in enumerate(row.cells):
        key = cell._tc
        if key in seen:
            continue
        seen.add(key)
        result.append((grid_index, cell))
    return result


def _iter_tables(container):
    for table in getattr(container, "tables", []):
        yield table
        seen = set()
        for row in table.rows:
            for _, cell in _unique_row_cells(row):
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                yield from _iter_tables(cell)


def _story_containers(doc):
    yield doc
    seen = set()
    for section in doc.sections:
        for attr in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            container = getattr(section, attr)
            key = container._element
            if key not in seen:
                seen.add(key)
                yield container


def iter_all_paragraphs(doc):
    """Yield paragraphs in body, tables (including nested), headers and footers."""
    seen = set()
    for container in _story_containers(doc):
        for para in getattr(container, "paragraphs", []):
            key = para._p
            if key not in seen:
                seen.add(key)
                yield para
        for table in _iter_tables(container):
            for row in table.rows:
                for _, cell in _unique_row_cells(row):
                    for para in cell.paragraphs:
                        key = para._p
                        if key not in seen:
                            seen.add(key)
                            yield para


def all_tables(doc):
    tables = []
    seen = set()
    for container in _story_containers(doc):
        for table in _iter_tables(container):
            key = table._tbl
            if key not in seen:
                seen.add(key)
                tables.append(table)
    return tables


def _text_para(cell, last=False):
    paras = [p for p in cell.paragraphs if p.text.strip()]
    if not paras:
        return cell.paragraphs[0]
    return paras[-1] if last else paras[0]


def _clear_paragraph(para):
    for child in list(para._p):
        if child.tag != qn("w:pPr"):
            para._p.remove(child)


def _set_cell_text(cell, value):
    """Replace cell contents while retaining useful paragraph/run formatting."""
    target = _text_para(cell)
    proto = _first_run(target)
    for para in list(cell.paragraphs):
        if para._p is not target._p:
            para._p.getparent().remove(para._p)
    _clear_paragraph(target)
    run = target.add_run(str(value))
    if proto is not None:
        _copy_rpr(proto, run)
    return run


def _append_to_cell(cell, text):
    para = _text_para(cell, last=True)
    proto = _first_run(para)
    run = para.add_run(str(text))
    if proto is not None:
        _copy_rpr(proto, run)
    return run


def _prepend_to_cell(cell, text):
    para = _text_para(cell)
    proto = _first_run(para)
    run = para.add_run(str(text))
    if proto is not None:
        _copy_rpr(proto, run)
    p = para._p
    p.remove(run._element)
    ppr = p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(run._element)
    else:
        p.insert(0, run._element)
    return run


def _replace_once_across_runs(para, needle, replacement):
    runs = list(para.runs)
    if not runs:
        return False
    full = "".join(r.text for r in runs)
    start = full.find(needle)
    if start < 0:
        return False
    end = start + len(needle)
    positions = []
    cursor = 0
    for index, run in enumerate(runs):
        positions.append((index, cursor, cursor + len(run.text)))
        cursor += len(run.text)
    start_info = next((p for p in positions if p[1] <= start < p[2]), None)
    end_pos = max(start, end - 1)
    end_info = next((p for p in positions if p[1] <= end_pos < p[2]), start_info)
    if start_info is None or end_info is None:
        return False
    si, s0, _ = start_info
    ei, e0, _ = end_info
    soff = start - s0
    eoff = end - e0
    before = runs[si].text[:soff]
    after = runs[ei].text[eoff:]
    if si == ei:
        runs[si].text = before + replacement + after
    else:
        runs[si].text = before + replacement
        for idx in range(si + 1, ei):
            runs[idx].text = ""
        runs[ei].text = after
    return True


def replace_marker(doc, marker, value, replace_all=False, on_replace=None):
    count = 0
    for para in iter_all_paragraphs(doc):
        while marker in para.text and _replace_once_across_runs(para, marker, str(value)):
            count += 1
            if on_replace:
                on_replace(para)
            if not replace_all:
                return count
            if count > 100:
                raise TemplateFillError(f"标记符出现次数异常: {marker}")
    return count


def _story_roots(doc):
    seen = set()
    for container in _story_containers(doc):
        root = container._element
        key = root
        if key not in seen:
            seen.add(key)
            yield root


def fill_content_control(doc, tag, value, replace_all=False):
    """Fill Word content controls whose tag or alias equals ``tag``."""
    count = 0
    for root in _story_roots(doc):
        for sdt in root.iter(qn("w:sdt")):
            props = sdt.find(qn("w:sdtPr"))
            if props is None:
                continue
            tags = props.findall(qn("w:tag")) + props.findall(qn("w:alias"))
            names = [x.get(qn("w:val")) for x in tags]
            if tag not in names:
                continue
            contents = sdt.find(qn("w:sdtContent"))
            if contents is None:
                continue
            texts = list(contents.iter(qn("w:t")))
            if texts:
                texts[0].text = str(value)
                texts[0].set(qn("xml:space"), "preserve")
                for node in texts[1:]:
                    node.text = ""
            else:
                paragraphs = list(contents.iter(qn("w:p")))
                para = paragraphs[0] if paragraphs else OxmlElement("w:p")
                if not paragraphs:
                    contents.append(para)
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.set(qn("xml:space"), "preserve")
                text.text = str(value)
                run.append(text)
                para.append(run)
            count += 1
            if not replace_all:
                return count
    return count


def _matches_label(text, label, method="contains"):
    if method == "regex":
        return re.search(label, text or "") is not None
    left = _norm(text)
    right = _norm(label)
    return left == right if method == "exact" else right in left


def _cell_anchor_candidates(doc, selector):
    tables = all_tables(doc)
    table_index = selector.get("table")
    if table_index is not None:
        if not 0 <= table_index < len(tables):
            raise TemplateFillError(f"表格编号越界: {table_index}")
        selected = [(table_index, tables[table_index])]
    else:
        selected = list(enumerate(tables))
    label = selector.get("label", "")
    method = selector.get("match", "contains")
    candidates = []
    for ti, table in selected:
        for ri, row in enumerate(table.rows):
            unique = _unique_row_cells(row)
            for ui, (grid_index, cell) in enumerate(unique):
                if _matches_label(cell.text, label, method):
                    candidates.append((ti, ri, ui, grid_index, cell, unique, table))
    return candidates


def _select_cell(doc, field):
    selector = field.get("selector") or {}
    if selector.get("type") == "cell_anchor":
        candidates = _cell_anchor_candidates(doc, selector)
        if not candidates:
            raise TemplateFillError(f"未找到表格标签: {selector.get('label')}")
        occurrence = selector.get("occurrence")
        if occurrence is None and len(candidates) > 1:
            raise TemplateFillError(
                f"表格标签不唯一: {selector.get('label')}（{len(candidates)} 处），请配置 occurrence 或 table"
            )
        index = occurrence or 0
        if index < 0 or index >= len(candidates):
            raise TemplateFillError(f"表格标签 occurrence 越界: {selector.get('label')}")
        ti, ri, ui, grid_index, cell, unique, table = candidates[index]
        direction = selector.get("direction", "same")
        offset = int(selector.get("offset", 1))
        if direction == "same":
            return cell, {"table": ti, "row": ri, "grid_col": grid_index, "via": "cell_anchor"}
        if direction == "right":
            target_index = ui + offset
            if target_index >= len(unique):
                raise TemplateFillError(f"标签右侧没有可写单元格: {selector.get('label')}")
            target_grid, target = unique[target_index]
            return target, {"table": ti, "row": ri, "grid_col": target_grid, "via": "cell_anchor"}
        if direction == "below":
            target_row = ri + offset
            if target_row >= len(table.rows):
                raise TemplateFillError(f"标签下方没有可写单元格: {selector.get('label')}")
            row_cells = table.rows[target_row].cells
            if grid_index >= len(row_cells):
                raise TemplateFillError(f"标签下方网格不兼容: {selector.get('label')}")
            return row_cells[grid_index], {"table": ti, "row": target_row, "grid_col": grid_index, "via": "cell_anchor"}
        raise TemplateFillError(f"未知 cell_anchor direction: {direction}")

    loc = field.get("loc")
    if loc is None:
        raise TemplateFillError("字段缺少 marker、control_tag、selector、anchor 或 loc")
    ti, ri, ci = loc
    tables = all_tables(doc)
    try:
        cell = tables[ti].rows[ri].cells[ci]
    except (IndexError, TypeError):
        raise TemplateFillError(f"旧坐标越界: {loc}")
    guard = field.get("expected_text")
    if guard and not _matches_label(cell.text, guard, field.get("expected_match", "contains")):
        raise TemplateFillError(f"旧坐标保护校验失败: {loc} 不含 {guard}")
    return cell, {"table": ti, "row": ri, "grid_col": ci, "via": "legacy_loc"}


DATE_PLACEHOLDER = re.compile(r"^\s*年\s+月\s+日")


def _consume_after(runs, start_run_idx, start_off, count):
    index, offset, left = start_run_idx, start_off, count
    while left > 0 and index < len(runs):
        text = runs[index].text
        cut = min(left, max(0, len(text) - offset))
        if cut:
            runs[index].text = text[:offset] + text[offset + cut :]
            left -= cut
        index += 1
        offset = 0


def fill_paragraph_field(doc, anchor, value, replace_date_placeholder=True, replace_existing=False):
    if value is None or str(value).strip() == "":
        return 0
    matches = [p for p in iter_all_paragraphs(doc) if anchor in p.text]
    if not matches:
        return 0
    if len(matches) > 1:
        raise TemplateFillError(f"段落锚点不唯一: {anchor}（{len(matches)} 处）")
    para = matches[0]
    runs = list(para.runs)
    if not runs:
        para.text = para.text.replace(anchor, anchor + str(value), 1)
        return 1
    pos = para.text.index(anchor) + len(anchor)
    acc = 0
    target_index = None
    target_offset = 0
    for index, run in enumerate(runs):
        if acc + len(run.text) >= pos:
            target_index, target_offset = index, pos - acc
            break
        acc += len(run.text)
    if target_index is None:
        target_index, target_offset = len(runs) - 1, len(runs[-1].text)
    if replace_existing:
        _consume_after(runs, target_index, target_offset, len(para.text) - pos)
    elif replace_date_placeholder:
        match = DATE_PLACEHOLDER.match(para.text[pos:])
        if match:
            _consume_after(runs, target_index, target_offset, match.end())
    run = runs[target_index]
    run.text = run.text[:target_offset] + str(value) + run.text[target_offset:]
    return 1


def _fill_choice(cell, value, field):
    options = field.get("options") or []
    if isinstance(options, dict):
        selected = options.get(str(value), value)
    else:
        selected = value
    selected_values = selected if isinstance(selected, (list, tuple, set)) else [selected]
    selected_norm = {_norm(x) for x in selected_values}
    text = cell.text
    box_pattern = re.compile(r"[□☐☑☒○◯oO]\s*([^□☐☑☒○◯oO\s，,;/；]+)")
    if not box_pattern.search(text):
        _set_cell_text(cell, str(value))
        return

    def repl(match):
        label = match.group(1)
        checked = _norm(label) in selected_norm
        return ("☒" if checked else "☐") + label

    _set_cell_text(cell, box_pattern.sub(repl, text))


def _prepare_text(field, value):
    fmt = field.get("format") or {}
    overflow = fmt.get("overflow", "warn")
    max_chars = fmt.get("max_chars")
    text = str(value)
    warnings = []
    if max_chars and len(text) > int(max_chars):
        if overflow == "error":
            raise TemplateFillError(f"内容超过 max_chars={max_chars}")
        if overflow == "truncate":
            text = text[: int(max_chars)]
            warnings.append(f"内容已按 max_chars={max_chars} 截断")
        else:
            warnings.append(f"内容长度 {len(text)} 超过建议值 {max_chars}，未截断")
    return text, warnings


def _apply_paragraph_format(para, field, original_length):
    fmt = field.get("format") or {}
    align_values = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    alignment = fmt.get("alignment")
    if alignment in align_values:
        para.alignment = align_values[alignment]
    font_size = fmt.get("font_size_pt")
    max_chars = fmt.get("max_chars")
    if fmt.get("auto_shrink") and max_chars and original_length > int(max_chars):
        base = float(font_size or 10.5)
        minimum = float(fmt.get("min_font_size_pt", 8.0))
        font_size = max(minimum, base * math.sqrt(int(max_chars) / original_length))
    if font_size:
        for run in para.runs:
            run.font.size = Pt(float(font_size))


def _apply_format(cell, field, value, original_length):
    fmt = field.get("format") or {}
    max_chars = fmt.get("max_chars")
    text = str(value)

    vertical = fmt.get("vertical_alignment")
    if vertical:
        values = {
            "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
            "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
            "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
        }
        if vertical in values:
            cell.vertical_alignment = values[vertical]

    align = fmt.get("alignment")
    align_values = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    font_size = fmt.get("font_size_pt")
    if fmt.get("auto_shrink") and max_chars and original_length > int(max_chars):
        base = float(font_size or 10.5)
        minimum = float(fmt.get("min_font_size_pt", 8.0))
        font_size = max(minimum, base * math.sqrt(int(max_chars) / original_length))
    for para in cell.paragraphs:
        if align in align_values:
            para.alignment = align_values[align]
        for run in para.runs:
            if font_size:
                run.font.size = Pt(float(font_size))
    return text


def fill_table_field(doc, field, value):
    cell, location = _select_cell(doc, field)
    cur = cell.text.strip()
    original_value = str(value)
    val, warnings = _prepare_text(field, original_value)
    mode = field.get("mode", "replace" if field.get("selector") else "append")
    prefix = field.get("prefix", "")

    if mode in ("replace", "value"):
        _set_cell_text(cell, prefix + val)
    elif mode == "after_label":
        label = (field.get("selector") or {}).get("label") or field.get("label")
        if not label or not _matches_label(cur, label):
            raise TemplateFillError(f"同格标签不存在: {label}")
        separator = field.get("separator", "：")
        _set_cell_text(cell, str(label).rstrip("：:") + separator + prefix + val)
    elif mode == "prepend":
        if val not in cur:
            _prepend_to_cell(cell, prefix + val + ("  " if cur else ""))
    elif mode == "amount":
        plain = cur.replace(" ", "").replace("\u3000", "")
        if plain in ("元", ""):
            if re.search(r"免收|免费|不收", val):
                _set_cell_text(cell, val)
            elif plain == "元":
                _prepend_to_cell(cell, val.rstrip("元"))
            else:
                _set_cell_text(cell, val)
        elif val not in cur:
            _append_to_cell(cell, val)
    elif mode == "choice":
        _fill_choice(cell, val, field)
    elif mode == "append":
        if val not in cur:
            separator = "" if cur.endswith(("：", ":")) else ("：" if cur else "")
            _append_to_cell(cell, separator + prefix + val)
    else:
        raise TemplateFillError(f"未知 mode: {mode}")

    _apply_format(cell, field, val, len(original_value))
    location["written_value"] = val
    return location, warnings


def _document_text(doc):
    parts = [p.text for p in iter_all_paragraphs(doc)]
    for root in _story_roots(doc):
        parts.extend(t.text or "" for t in root.iter(qn("w:t")))
    return "\n".join(parts)


def _field_identity(field):
    return field.get("name") or field.get("marker") or field.get("control_tag") or "未命名字段"


def _apply_field(doc, field, value):
    marker = field.get("marker")
    if marker:
        written_value, warnings = _prepare_text(field, value)
        matched_paragraphs = []
        count = replace_marker(
            doc,
            marker,
            written_value,
            bool(field.get("replace_all")),
            on_replace=matched_paragraphs.append,
        )
        if not count:
            raise TemplateFillError(f"未找到标记符: {marker}")
        for para in matched_paragraphs:
            _apply_paragraph_format(para, field, len(str(value)))
        return {"via": "marker", "count": count, "written_value": written_value}, warnings

    control_tag = field.get("control_tag")
    if control_tag:
        written_value, warnings = _prepare_text(field, value)
        count = fill_content_control(doc, control_tag, written_value, bool(field.get("replace_all")))
        if not count:
            raise TemplateFillError(f"未找到内容控件: {control_tag}")
        return {"via": "content_control", "count": count, "written_value": written_value}, warnings

    selector = field.get("selector") or {}
    if selector.get("type") == "paragraph_anchor":
        anchor = selector.get("label")
        written_value, warnings = _prepare_text(field, value)
        count = fill_paragraph_field(doc, anchor, written_value, replace_existing=bool(selector.get("replace_existing")))
        if not count:
            raise TemplateFillError(f"未找到段落锚点: {anchor}")
        return {"via": "paragraph_anchor", "count": count, "written_value": written_value}, warnings

    if field.get("anchor") and not selector:
        written_value, warnings = _prepare_text(field, value)
        count = fill_paragraph_field(doc, field["anchor"], written_value, replace_existing=bool(field.get("replace_existing")))
        if not count:
            raise TemplateFillError(f"未找到段落锚点: {field['anchor']}")
        return {"via": "legacy_anchor", "count": count, "written_value": written_value}, warnings

    return fill_table_field(doc, field, value)


def fill_template(template_path, out_path, fields_values, strict=True, report_path=None):
    """Fill a template and return a structured report.

    ``fields_values`` is a list of mapping definitions with a runtime ``value``.
    Empty values are intentionally skipped. In strict mode, any non-empty field
    that cannot be located or verified prevents replacement of ``out_path``.
    """
    template_path = str(template_path)
    out_path = str(out_path)
    report = {
        "engine": "case-archiver-table-v3",
        "template": template_path,
        "output": out_path,
        "strict": bool(strict),
        "ok": True,
        "fields": [],
        "errors": [],
        "warnings": [],
    }
    doc = Document(template_path)
    written = []
    for field in fields_values:
        name = _field_identity(field)
        value = field.get("value")
        entry = {"name": name, "status": "skipped", "value": "" if value is None else str(value)}
        if value is None or str(value).strip() == "":
            report["fields"].append(entry)
            continue
        try:
            location, warnings = _apply_field(doc, field, value)
            entry.update({"status": "filled", "location": location, "warnings": warnings})
            report["warnings"].extend(f"{name}: {w}" for w in warnings)
            written.append((field, str(location.get("written_value", value)), entry))
        except Exception as exc:
            entry.update({"status": "error", "error": str(exc)})
            report["errors"].append(f"{name}: {exc}")
        report["fields"].append(entry)

    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output.with_name(f".{output.name}.tmp-{os.getpid()}.docx")
    try:
        doc.save(str(temp_path))
        saved = Document(str(temp_path))
        saved_text = _document_text(saved)
        for field, value, entry in written:
            if field.get("verify", True) and value not in saved_text:
                message = f"{entry['name']}: 写后校验未找到填充值"
                entry["status"] = "verify_error"
                entry["error"] = message
                report["errors"].append(message)
        report["ok"] = not report["errors"]
        if strict and report["errors"]:
            raise TemplateFillError("；".join(report["errors"]), report)
        os.replace(temp_path, output)
    except Exception:
        if temp_path.exists():
            temp_path.unlink()
        report["ok"] = False
        if report_path:
            Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        raise

    if report_path:
        Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def inspect_template(template_path):
    doc = Document(template_path)
    marker_re = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}|\[\[\s*([^\[\]]+?)\s*\]\]")
    tables = []
    for ti, table in enumerate(all_tables(doc)):
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for grid_index, cell in _unique_row_cells(row):
                cells.append(
                    {
                        "grid_col": grid_index,
                        "text": cell.text,
                        "markers": [next(x for x in m.groups() if x is not None) for m in marker_re.finditer(cell.text)],
                    }
                )
            rows.append({"row": ri, "cells": cells})
        tables.append({"table": ti, "rows": rows})
    controls = []
    for root in _story_roots(doc):
        for sdt in root.iter(qn("w:sdt")):
            props = sdt.find(qn("w:sdtPr"))
            if props is None:
                continue
            tags = props.findall(qn("w:tag")) + props.findall(qn("w:alias"))
            controls.extend(x.get(qn("w:val")) for x in tags if x.get(qn("w:val")))
    paragraph_markers = []
    for index, para in enumerate(iter_all_paragraphs(doc)):
        for match in marker_re.finditer(para.text):
            paragraph_markers.append({"paragraph": index, "name": next(x for x in match.groups() if x is not None)})
    return {"template": str(template_path), "tables": tables, "content_controls": sorted(set(controls)), "markers": paragraph_markers}


def main():
    parser = argparse.ArgumentParser(description="Inspect a DOCX template for V3 field mapping")
    parser.add_argument("template")
    parser.add_argument("--out", help="Write inspection JSON to this path")
    args = parser.parse_args()
    result = inspect_template(args.template)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    else:
        print(text)


if __name__ == "__main__":
    main()
