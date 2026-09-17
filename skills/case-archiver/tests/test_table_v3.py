import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import fill_docx
import onboard
import archive_case


def build_complex_template(path):
    doc = Document()
    header = doc.sections[0].header.paragraphs[0]
    header.add_run("卷宗：{{案")
    header.add_run("件名称}}")

    table = doc.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 3)).text = "案件审批表"

    table.cell(1, 0).text = "案号"
    number_cell = table.cell(1, 1).merge(table.cell(1, 3))
    number_para = number_cell.paragraphs[0]
    number_para.text = ""
    number_para.add_run("{{案")
    number_para.add_run("号}}")

    table.cell(2, 0).text = "案由"
    table.cell(2, 1).merge(table.cell(2, 3)).text = ""

    table.cell(3, 0).text = "处理方式"
    table.cell(3, 1).merge(table.cell(3, 3)).text = "□劳动仲裁  □一审诉讼  □二审诉讼"

    table.cell(4, 0).text = "案情简介"
    table.cell(4, 1).merge(table.cell(4, 3)).text = "{{案情简介}}"
    doc.save(path)


class TableV3Tests(unittest.TestCase):
    def test_archive_case_adapter_returns_structured_report(self):
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "adapter.docx"
            output = Path(temp) / "output.docx"
            doc = Document()
            doc.add_paragraph("案号：{{案号}}")
            doc.save(template)
            fields, report = archive_case.fill_table_template(
                template,
                output,
                {
                    "strict": True,
                    "fields": [
                        {"name": "案号", "source": "agent", "marker": "{{案号}}"}
                    ],
                },
                {},
                {"案号": "（2026）浙0101民初123号"},
            )
            self.assertTrue(report["ok"])
            self.assertEqual("filled", fields[0]["write_status"])
            self.assertIn("（2026）浙0101民初123号", fill_docx._document_text(Document(output)))

    def test_content_control_tag(self):
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "control.docx"
            output = Path(temp) / "output.docx"
            doc = Document()
            para = doc.add_paragraph("承办律师：")
            sdt = OxmlElement("w:sdt")
            props = OxmlElement("w:sdtPr")
            tag = OxmlElement("w:tag")
            tag.set(qn("w:val"), "承办律师")
            props.append(tag)
            content = OxmlElement("w:sdtContent")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "点击填写"
            run.append(text)
            content.append(run)
            sdt.append(props)
            sdt.append(content)
            para._p.append(sdt)
            doc.save(template)

            report = fill_docx.fill_template(
                template,
                output,
                [{"name": "承办律师", "control_tag": "承办律师", "value": "覃律师"}],
            )
            self.assertTrue(report["ok"])
            inspection = fill_docx.inspect_template(output)
            self.assertIn("承办律师", inspection["content_controls"])
            saved = Document(output)
            self.assertIn("覃律师", fill_docx._document_text(saved))

    def test_stable_selectors_merged_cells_and_choices(self):
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "template.docx"
            output = Path(temp) / "output.docx"
            build_complex_template(template)
            summary = "原告请求支付工程款。\n法院查明双方签订施工合同。\n判决被告支付相应款项。"
            report = fill_docx.fill_template(
                template,
                output,
                [
                    {"name": "案件名称", "marker": "{{案件名称}}", "value": "张三诉李四合同纠纷案"},
                    {"name": "案号", "marker": "{{案号}}", "value": "（2026）浙0101民初123号"},
                    {
                        "name": "案由",
                        "selector": {"type": "cell_anchor", "table": 0, "label": "案由", "direction": "right"},
                        "mode": "replace",
                        "value": "建设工程施工合同纠纷",
                    },
                    {
                        "name": "处理方式",
                        "selector": {"type": "cell_anchor", "table": 0, "label": "处理方式", "direction": "right"},
                        "mode": "choice",
                        "value": "一审诉讼",
                    },
                    {
                        "name": "案情简介",
                        "marker": "{{案情简介}}",
                        "value": summary,
                        "format": {"max_chars": 20, "overflow": "warn"},
                    },
                ],
            )
            self.assertTrue(report["ok"])
            self.assertTrue(report["warnings"])
            result = Document(output)
            text = "\n".join(p.text for p in fill_docx.iter_all_paragraphs(result))
            self.assertIn("张三诉李四合同纠纷案", text)
            self.assertIn("（2026）浙0101民初123号", text)
            self.assertIn("建设工程施工合同纠纷", text)
            self.assertIn("☒一审诉讼", text)
            self.assertIn("☐劳动仲裁", text)
            self.assertIn(summary, text)

    def test_ambiguous_anchor_fails_without_replacing_output(self):
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "ambiguous.docx"
            output = Path(temp) / "output.docx"
            doc = Document()
            for _ in range(2):
                table = doc.add_table(rows=1, cols=2)
                table.cell(0, 0).text = "案由"
                table.cell(0, 1).text = ""
            doc.save(template)
            with self.assertRaises(fill_docx.TemplateFillError):
                fill_docx.fill_template(
                    template,
                    output,
                    [
                        {
                            "name": "案由",
                            "selector": {"type": "cell_anchor", "label": "案由", "direction": "right"},
                            "value": "合同纠纷",
                        }
                    ],
                )
            self.assertFalse(output.exists())

    def test_legacy_coordinate_guard_fails_closed(self):
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "legacy.docx"
            output = Path(temp) / "output.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "案号"
            table.cell(0, 1).text = ""
            doc.save(template)
            with self.assertRaises(fill_docx.TemplateFillError):
                fill_docx.fill_template(
                    template,
                    output,
                    [{"name": "案号", "loc": [0, 0, 1], "expected_text": "案号", "value": "123"}],
                )

    def test_onboard_deduplicates_merged_cells(self):
        with tempfile.TemporaryDirectory() as temp:
            template = Path(temp) / "detect.docx"
            build_complex_template(template)
            spec = onboard.detect_fields(template)
            selectors = [field.get("selector") for field in spec["fields"] if field.get("selector")]
            case_reason = [item for item in selectors if item.get("label") == "案由"]
            self.assertEqual(1, len(case_reason))
            self.assertTrue(all(field.get("confidence") for field in spec["fields"]))


if __name__ == "__main__":
    unittest.main()
