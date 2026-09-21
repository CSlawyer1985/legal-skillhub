"""Contract 系共享 Word 底层 helper（对齐 NDR 统一格式底座 format_base）。"""
import os
import sys

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_here = os.path.dirname(os.path.abspath(__file__))
_ndr_ref = os.path.join(_here, '..', '..', '非争议解决公共标准', 'references')
if os.path.isdir(_ndr_ref) and _ndr_ref not in sys.path:
    sys.path.insert(0, _ndr_ref)
from format_base import *


def set_font(run, font_name, size, bold=False):
    run.font.name = font_name
    run.font.size = _resolve_size(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = MARGIN_TB
    section.bottom_margin = MARGIN_TB
    section.left_margin = MARGIN_LR
    section.right_margin = MARGIN_LR


def setup_default_style(doc, font_name, font_size):
    style = doc.styles['Normal']
    style.font.name = font_name
    style.font.size = _resolve_size(font_size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)
    style.paragraph_format.line_spacing = LINE_SPACING
    style.paragraph_format.space_after = Pt(0)


def add_body_paragraph(doc, text, font_name=FONT_VARIANT, size=BODY_SIZE):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = FIRST_INDENT
    run = p.add_run(text)
    set_font(run, font_name, size)
    return p


def add_blank_line(doc):
    return doc.add_paragraph()


def add_centered_paragraph(doc, text, font_name=FONT_VARIANT, size=BODY_SIZE, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, font_name, size, bold=bold)
    return p


def make_signature_block(doc, parties):
    for party in parties:
        label = party.get('label', '')
        name = party.get('name', '')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run('%s：%s' % (label, name))
        set_font(run, FONT_VARIANT, BODY_SIZE)
    return doc
