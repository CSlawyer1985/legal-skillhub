#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""字段出口：凡填进 casefile 的东西都要在交付物里出得来（T7）。
填了却不出现在图或表里的字段，等于让律师白填。实测抓到过三处：
整块诉请固定（J 组校验的全部产出）、决定性要件的「真伪不明的后果」、三标之三的立场方向。"""
import os, sys, json, subprocess, tempfile
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE)

# 白名单：由别处推导或只供校验用，不必单独露面
INTERNAL = {"id", "title", "response_burden", "second_burden", "figure_ref",
            "method", "claim_ref", "norm_type", "layer", "kind", "issue_kind"}

def flat(d, path=""):
    """摊平成 (路径, 文本) 对，只取字符串叶子。"""
    if isinstance(d, dict):
        for k, v in d.items():
            yield from flat(v, f"{path}.{k}" if path else k)
    elif isinstance(d, list):
        for x in d: yield from flat(x, path)
    elif isinstance(d, str):
        yield path, d

def main(case):
    out = tempfile.mkdtemp()
    subprocess.run([sys.executable, os.path.join(ROOT,"scripts","build_matrix.py"), case,
                    os.path.join(out,"m.xlsx")], capture_output=True)
    subprocess.run([sys.executable, os.path.join(ROOT,"scripts","to_figure.py"), case,
                    os.path.join(out,"f.json")], capture_output=True)
    subprocess.run([sys.executable, os.path.join(ROOT,"scripts","render_main.py"),
                    os.path.join(out,"f.json"), os.path.join(out,"f.svg")], capture_output=True)
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(out,"m.xlsx"))
    blob = ""
    for ws in wb.worksheets:
        blob += "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        blob += "\n".join(c.comment.text for row in ws.iter_rows() for c in row if c.comment)
    blob += open(os.path.join(out,"f.svg"), encoding="utf-8").read()

    d = json.load(open(case, encoding="utf-8"))
    miss = []
    for path, val in flat(d):
        key = path.split(".")[-1]
        if key in INTERNAL: continue
        s = val.strip()
        if not s or s in ("—", "否", "是", "中立"): continue
        if s[:16] not in blob: miss.append((path, s[:34]))
    ok = not miss
    print(("  OK   T7" if ok else "  FAIL T7") +
          f"  casefile 的每个字段都在交付物里出得来（缺 {len(miss)} 处）")
    for p, v in miss[:8]: print(f"       {p:<26}{v}")
    return 0 if ok else 1

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT,"examples","matrix-input.json")
    sys.exit(main(src))
