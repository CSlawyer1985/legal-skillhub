#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Maintained by Lu Lingyan, Deheng (Wuxi) Law Firm.
"""生成《待确认事实清单》.docx —— H1 事实核对互动的载体。

为什么必须逐条拆条
------------------
判决书认定的事实是对是错，只有当事人知道。但**不能让当事人自由陈述**——自由陈述必然
遗漏、必然跑偏。正确做法是把判决书"本院查明"部分逐条拆开，一条一条给当事人判断。
这份清单就是"要不要开会"的选择权：律师有空就当面过判决书，没空就直接发给当事人逐条勾。

输出：Word 文档（.docx）。**只出 Word** —— 可打印、可签名、可直接发微信。

用法
----
    python3 build_pending_facts_docx.py <payload.json> <输出.docx>

payload.json：
    {
      "案件": "张三与李四房屋买卖合同纠纷",
      "上诉人": "张三",
      "一审法院": "××市××区人民法院",
      "一审案号": "（2024）苏××××民初1234号",
      "上诉期届满日": "2026-10-01",
      "事实": [
        "判决书认定的第 1 条事实原文（照抄，不要转述）",
        "判决书认定的第 2 条事实原文"
      ],
      "落款": "××律师事务所　×××律师"
    }

「事实」数组由 agent 从判决书「本院查明」部分逐条摘录后传入，**必须是原文摘录**，
不得转述或概括——当事人要对照的就是判决书原话。
"""

import json
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HEADERS = ["序号", "判决书认定的事实（原文摘录）", "属实／不属实／部分属实",
           "如不属实，真实情况是", "有无证据", "证据名称", "在谁手上", "一审是否已提交"]
COL_W = [Cm(1.0), Cm(4.6), Cm(1.8), Cm(3.4), Cm(1.2), Cm(2.2), Cm(1.4), Cm(1.4)]

INSTRUCTIONS = [
    "填写说明",
    "1. 请对照判决书原文，逐条判断。不属实的，请在「真实情况」栏写明实际发生的情况。",
    "2. 有证据的，请写清证据名称与由谁保管。只写「事实不是这样」但没有证据的，无法作为上诉理由使用。",
    "3. 二审提交新证据需要说明为什么一审没有提出，并非所有「新事实」都能在二审提出。",
    "4. 如果您的陈述与在案证据相矛盾，可能反过来对己方不利，请如实填写。",
    "5. 另请补充：判决书完全没有提到、但对本案重要的事实有哪些（见文末补充栏）。",
]


def _font(run, name="仿宋", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _para(doc, text, name="仿宋", size=14, bold=False, align=None, indent=True, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(p.add_run(text), name, size, bold)
    return p


def build(payload, out_path):
    facts = payload.get("事实") or []
    if not facts:
        raise SystemExit("❌ payload['事实'] 为空：必须先从判决书「本院查明」部分逐条摘录原文。")

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = s.right_margin = Cm(1.8)

    _para(doc, "待确认事实清单", name="宋体", size=18, bold=True, align="center",
          indent=False, space_after=8)
    _para(doc, f"案件：{payload.get('案件', '')}", size=12, indent=False)
    _para(doc, f"上诉人：{payload.get('上诉人', '')}　　一审法院：{payload.get('一审法院', '')}　"
               f"一审案号：{payload.get('一审案号', '')}", size=12, indent=False)
    if payload.get("上诉期届满日"):
        _para(doc, f"上诉期届满日：{payload['上诉期届满日']}", size=12, indent=False)
    _para(doc, "", size=6, indent=False)

    _para(doc, "填写说明", name="宋体", size=14, bold=True, indent=False)
    for line in INSTRUCTIONS[1:]:
        _para(doc, line, size=11)

    _para(doc, "", size=6, indent=False)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(HEADERS):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        _font(p.add_run(h), "宋体", 10.5, True)

    for n, fact in enumerate(facts, 1):
        row = table.add_row()
        row.cells[0].text = ""
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.line_spacing = 1.0
        _font(p0.add_run(str(n)), "仿宋", 10.5)
        for i in range(1, len(HEADERS)):
            c = row.cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            if i == 1:
                _font(p.add_run(fact), "仿宋", 10.5)

    for row in table.rows:
        for i, w in enumerate(COL_W):
            row.cells[i].width = w

    _para(doc, "", size=8, indent=False)
    _para(doc, "补充栏", name="宋体", size=14, bold=True, indent=False)
    _para(doc, "一、判决书完全没有提到、但对本案重要的事实：", size=12)
    for _ in range(5):
        _para(doc, "　", size=12, indent=False)
    _para(doc, "二、对判决理由（「本院认为」部分）的意见：", size=12)
    _para(doc, "　□ 认可　　□ 不认可，理由是：", size=12)
    for _ in range(3):
        _para(doc, "　", size=12, indent=False)
    _para(doc, "三、本次提供的新证据：", size=12)
    _para(doc, "　证据名称／来源／形成时间／拟证明的事实：", size=12)
    for _ in range(4):
        _para(doc, "　", size=12, indent=False)

    _para(doc, "", size=10, indent=False)
    _para(doc, f"当事人签名：　　　　　　　　　　　　日期：　　年　　月　　日", size=12)
    _para(doc, payload.get("落款", ""), size=12, align="right", indent=False)

    doc.save(out_path)

    # 自检：条目数进出一致
    got = len(doc.tables[0].rows) - 1
    if got != len(facts):
        raise SystemExit(f"❌ 自检未通过：输入 {len(facts)} 条事实，输出表格 {got} 行。")
    return out_path, len(facts)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        data = json.load(fh)
    path, n = build(data, sys.argv[2])
    print(f"✅ {path}（{n} 条事实，自检通过）")
