#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Maintained by Lu Lingyan, Deheng (Wuxi) Law Firm.
"""上诉状填充器 —— 以《民事上诉状-脱敏模板.docx》为格式母版，原位填充，格式 100% 保留。

为什么不用通用渲染引擎
----------------------
`render_docx.py` 是从零 `Document()` 建文档再套样式，无法复刻母版中的细节（黑体 18pt
居中标题、仿宋 14pt 正文与加粗规则、"此致"缩进、落款居中、日期右对齐、页脚 PAGE/NUMPAGES
域名）。上诉状必须与母版逐项一致，因此这里采用**克隆母版段落为格式供体**的做法：
- 当事人栏  → 克隆母版"上诉人"段（仿宋 14pt 加粗）
- 上诉请求项 → 克隆母版"1、……"段（仿宋 14pt 不加粗）
- 事实与理由小节标题 → 克隆母版"一、……"段（仿宋 14pt 加粗）
- 事实与理由正文 → 克隆母版"错误之一：……"段（仿宋 14pt 不加粗）
段落数量按实际内容增减，母版中多余的示例段落会被删除。

输入
----
    python3 fill_appeal_docx.py <母版.docx> <输出.docx> <payload.json>

payload.json 结构见 references/上诉状填充字段表.md，示例：

    {
      "上诉人":   {"原审地位": "被告", "姓名": "张三", "身份": "张三，男，1980年1月1日出生，汉族，住××市××区××路1号。"},
      "被上诉人": [{"原审地位": "原告", "身份": "李四，男，……。"}],
      "第三人":   [{"身份": "王五，男，……。"}],
      "一审法院": "××市××区人民法院",
      "一审案号": "（2024）苏××××民初1234号",
      "二审法院": "××市中级人民法院",
      "上诉请求": ["1、撤销……；或撤销原判决，发回重审；", "2、本案一审、二审诉讼费用均由被上诉人承担。"],
      "事实与理由": [
        {"kind": "lead", "text": "开篇结论段……"},
        {"kind": "head", "text": "一、一审未查明案件基本事实，且事实认定错误。"},
        {"kind": "body", "text": "错误之一：……"},
        {"kind": "body", "text": "……"}
      ],
      "落款签名": "张三",
      "落款日期": "2026年9月16日"
    }

说明
----
- `原审地位` 只填「原告／被告／第三人」，输出为「上诉人（一审被告）：……」，**不要**自带"一审"二字。
- `一审案号` 自带全角括号时不会被重复加括号。
- `kind` 三选一：`lead`（开篇结论段）／`head`（"一、二、"小结标题，加粗）／`body`（正文，不加粗）。

自检
----
落盘前检查：① 正文不得残留 `〔` 占位符；② 上诉请求、事实与理由非空；
③ 当事人栏必备字段齐全。任一不过 → 报错并不落盘。
"""

import copy
import json
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# 母版段落索引（0 基）—— 对应《民事上诉状-脱敏模板.docx》34 段的骨架
IDX = {
    "title": 0,           # 民事上诉状
    "appellant": 2,       # 上诉人（一审……）  ← 当事人栏格式供体
    "appellee": 3,        # 被上诉人
    "third_party_first": 4,
    "third_party_last": 7,
    "preamble": 9,        # 〔上诉人〕因不服……提起上诉。
    "req_head": 10,       # 上诉请求：
    "req_first": 11,      # 1、……
    "req_second": 12,     # 2、……
    "claim_head": 14,     # 事实和理由：
    "lead": 15,           # 开篇结论段
    "head": 16,           # 一、……         ← 小结标题格式供体（加粗）
    "body_first": 17,     # 错误之一：……    ← 正文格式供体（不加粗）
    "body_last": 28,      # 综上，……（事实与理由区段末尾）
    "court": 30,          # 〔二审法院全称〕
    "sign": 32,           # 上诉人：〔……〕
    "date": 33,           # 〔　〕年〔　〕月〔　〕日
}


def _strip_runs(par):
    """删除段落内所有 run，保留段落属性（pPr）与书签等非 run 元素。"""
    for r in list(par.runs):
        r._element.getparent().remove(r._element)


def set_text(par, text, fmt_from=None):
    """写入文本，保留段落属性（pPr）。

    格式来源优先级：显式传入的 fmt_from（母版供体段落）> 该段落自身原有的首个 run。
    **必须在删除 run 之前先把 rPr 抓下来**，否则新建的 run 会退回 Normal 样式，
    造成字体、字号、加粗全部丢失。
    """
    src = fmt_from if fmt_from is not None else par._element
    rpr = None
    src_runs = src.findall(qn("w:r"))
    if src_runs:
        rpr = src_runs[0].find(qn("w:rPr"))
    _strip_runs(par)
    run = par.add_run(text)
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))
    return par


def _case_no(no):
    """案号自带全角/半角括号时不再另加括号，避免出现（（2024）……）双括号。"""
    no = (no or "").strip()
    if not no:
        return ""
    return no if no[0] in "（(" else f"（{no}）"


def clone_after(anchor_el, donor_el, text, parent):
    """在 anchor_el 之后插入一个 donor_el 的副本并写入 text，返回新段落元素与 Paragraph 对象。"""
    new_el = copy.deepcopy(donor_el)
    # 去掉副本里原有的 run，仅保留段落属性作为格式骨架
    for r in new_el.findall(qn("w:r")):
        new_el.remove(r)
    anchor_el.addnext(new_el)
    par = Paragraph(new_el, parent)
    set_text(par, text, fmt_from=donor_el)
    return new_el, par


def remove_els(els):
    for el in els:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def fill(template_path, output_path, payload):
    doc = Document(template_path)
    parent = doc.paragraphs[0]._parent
    paras = doc.paragraphs

    # ---- 先固化格式供体（母版段落元素，必须在删除之前深拷贝）----
    d_party = copy.deepcopy(paras[IDX["appellant"]]._element)   # 当事人栏：仿宋 14pt 加粗
    d_req = copy.deepcopy(paras[IDX["req_first"]]._element)     # 请求项：仿宋 14pt 不加粗
    d_head = copy.deepcopy(paras[IDX["head"]]._element)         # 小结标题：加粗
    d_body = copy.deepcopy(paras[IDX["body_first"]]._element)   # 正文：不加粗

    missing = [k for k in ("上诉人", "被上诉人", "一审法院", "一审案号", "二审法院",
                           "上诉请求", "事实与理由") if not payload.get(k)]
    if missing:
        raise SystemExit(f"❌ payload 缺字段：{', '.join(missing)}")

    # ---- 一、当事人栏 ----
    a = payload["上诉人"]
    set_text(paras[IDX["appellant"]],
             f"上诉人（一审{a.get('原审地位', '')}）：{a['身份']}",
             fmt_from=d_party)

    # 母版中"被上诉人 + 第三人×2 + 两行提示"整段删掉，再按实际人数重建
    remove_els([paras[i]._element
                for i in range(IDX["appellee"], IDX["third_party_last"] + 1)])
    anchor = paras[IDX["appellant"]]._element
    for ap in payload["被上诉人"]:
        anchor, _ = clone_after(anchor, d_party,
                                f"被上诉人（一审{ap.get('原审地位', '')}）：{ap['身份']}", parent)
    for tp in payload.get("第三人", []):
        anchor, _ = clone_after(anchor, d_party, f"第三人：{tp['身份']}", parent)

    # ---- 二、不服前提句 ----
    who = payload.get("上诉人简称") or a.get("姓名") or "上诉人"
    set_text(paras[IDX["preamble"]],
             f"{who}因不服{payload['一审法院']}{_case_no(payload['一审案号'])}民事判决，提起上诉。",
             fmt_from=d_party)

    # ---- 三、上诉请求 ----
    remove_els([paras[IDX["req_first"]]._element, paras[IDX["req_second"]]._element])
    anchor = paras[IDX["req_head"]]._element
    for item in payload["上诉请求"]:
        anchor, _ = clone_after(anchor, d_req, item, parent)

    # ---- 四、事实与理由（整段重建）----
    region = [p._element for p in paras[IDX["lead"]:IDX["body_last"] + 1]]
    remove_els(region)
    anchor = paras[IDX["claim_head"]]._element
    donor_map = {"lead": d_body, "head": d_head, "body": d_body}
    for item in payload["事实与理由"]:
        kind = item.get("kind", "body")
        if kind not in donor_map:
            raise SystemExit(f"❌ 事实与理由的 kind 只能是 lead/head/body，收到：{kind}")
        anchor, _ = clone_after(anchor, donor_map[kind], item["text"], parent)

    # ---- 五、此致 / 落款 ----
    set_text(paras[IDX["court"]], payload["二审法院"])
    sign = payload.get("落款签名")
    set_text(paras[IDX["sign"]], f"上诉人：{sign}" if sign else "上诉人：　　　　　　　　（签名/盖章）")
    set_text(paras[IDX["date"]], payload.get("落款日期") or "　　　　年　　月　　日")

    # ---- 自检 ----
    full = "\n".join(p.text for p in doc.paragraphs)
    errs = []
    if "〔" in full or "〕" in full:
        left = [l for l in full.splitlines() if "〔" in l]
        errs.append("残留占位符：" + " | ".join(left[:3]))
    if not payload["上诉请求"]:
        errs.append("上诉请求为空")
    if not payload["事实与理由"]:
        errs.append("事实与理由为空")
    for key in ("一审法院", "一审案号", "二审法院"):
        if payload[key] not in full:
            errs.append(f"{key} 未写入正文（{payload[key]}）")
    if errs:
        raise SystemExit("❌ 自检未通过，未落盘：\n  - " + "\n  - ".join(errs))

    doc.save(output_path)
    return output_path


def Paras(el, parent):
    return Paragraph(el, parent)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    with open(sys.argv[3], encoding="utf-8") as fh:
        data = json.load(fh)
    out = fill(sys.argv[1], sys.argv[2], data)
    print(f"✅ {out}（自检通过）")
