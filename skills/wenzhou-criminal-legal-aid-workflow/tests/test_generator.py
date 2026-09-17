from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
GENERATOR = ROOT / "scripts" / "generate_criminal_legal_aid_set.py"
TIMELINE = ROOT / "scripts" / "extract_case_timeline.py"
PLATFORM_CHECK = ROOT / "scripts" / "check_platform_compatibility.py"
CHARGE_VALIDATOR = ROOT / "scripts" / "validate_charge_library.py"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_text(path: Path) -> str:
    values: list[str] = []
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(zf.read(name))
            values.extend(node.text or "" for node in root.iter() if node.tag in {W_NS + "t", W_NS + "instrText"})
    return "\n".join(values)


def nonempty_paragraph_xml(path: Path) -> list[ET.Element]:
    with ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    return [
        paragraph
        for paragraph in root.iter(W_NS + "p")
        if "".join(node.text or "" for node in paragraph.iter(W_NS + "t")).strip()
    ]


def first_text_run(paragraph: ET.Element) -> ET.Element:
    return next(
        run
        for run in paragraph.iter(W_NS + "r")
        if "".join(node.text or "" for node in run.iter(W_NS + "t")).strip()
    )


class GeneratorTests(unittest.TestCase):
    def base_case(self) -> dict:
        return {
            "recipient_name": "测试人员",
            "recipient_id_no": "",
            "charge": "开设赌场罪",
            "stage": "审判阶段",
            "aid_center": "鹿城区法律援助中心",
            "handling_agency": "测试人民法院",
            "prosecuting_agency": "测试人民检察院",
            "custody_status": "取保候审",
            "lawyer": "李鸿鸿",
            "lawyer_phone": "",
            "law_firm": "浙江光正大律师事务所",
            "meeting_place": "浙江光正大律师事务所",
            "progress": ["归档", "取保候审"],
        }

    def generate(
        self,
        case: dict,
        mode: str = "archive",
        external_pack: bool = False,
        template_pack_dir: Path | None = None,
    ) -> Path:
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-skill-test-"))
        source = temp / "case.json"
        source.write_text(json.dumps(case, ensure_ascii=False), encoding="utf-8")
        command = [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out"), "--mode", mode]
        if external_pack:
            command.extend(["--template-pack-dir", str(ROOT / "assets" / "template-packs" / "wenzhou")])
        if template_pack_dir is not None:
            command.extend(["--template-pack-dir", str(template_pack_dir)])
        subprocess.run(
            command,
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        return next((temp / "out").iterdir())

    def test_non_custodial_meeting_uses_direct_contact(self) -> None:
        out = self.generate(self.base_case(), "current")
        meeting = next(out.glob("*会见笔录*.docx"))
        text = docx_text(meeting)
        self.assertIn("可通过电话或到浙江光正大律师事务所与律师联系", text)
        self.assertNotIn("通过看守所", text)
        self.assertNotIn("最近在看守所", text)
        self.assertNotIn("刑事拘留后", text)
        self.assertIn("侦查机关是否依法告知你有权委托辩护人", text)
        self.assertNotRegex(text, r"\{\{[^{}]+\}\}")

    def test_trial_keeps_court_and_prosecutor_separate(self) -> None:
        out = self.generate(self.base_case(), "current")
        trial = next(out.glob("*庭审笔录*.docx"))
        text = docx_text(trial)
        self.assertIn("测试人民法院", text)
        self.assertIn("测试人民检察院", text)
        self.assertNotIn("温州市测试人民检察院", text)
        self.assertNotIn("刑事审判第1庭", text)
        self.assertNotIn("鹿检   部刑诉", text)
        self.assertNotIn("辩：好的", text)

    def test_trial_outline_has_no_civil_case_residue(self) -> None:
        out = self.generate(self.base_case(), "current")
        outline = next(out.glob("*出庭提纲*.docx"))
        text = docx_text(outline)
        self.assertIn("辩护人发问提纲", text)
        self.assertNotIn("原告", text)
        self.assertNotIn("代理意见", text)

    def test_reading_note_uses_blank_fields_not_instruction_residue(self) -> None:
        out = self.generate(self.base_case(), "current")
        note = next(out.glob("*阅卷笔录*.docx"))
        text = docx_text(note)
        self.assertIn("供述、辩解摘录及证明内容", text)
        self.assertNotIn("X年X月X日", text)
        self.assertNotIn("供述及辩解摘抄提炼证明了什么", text)

    def test_blank_personal_phone_is_not_rendered_as_field(self) -> None:
        out = self.generate(self.base_case(), "current")
        authorization = next(out.glob("*委托书*.docx"))
        text = docx_text(authorization)
        self.assertNotIn("电话：", text)
        self.assertIn("受委托律师：李鸿鸿", text)

    def test_closing_report_lawyer_type_is_not_lawyer_name(self) -> None:
        out = self.generate(self.base_case(), "archive")
        closing = next(out.glob("*结案报告表*.docx"))
        text = docx_text(closing)
        self.assertIn("社会律师", text)

    def test_wenzhou_typography_is_applied_to_title_and_body(self) -> None:
        out = self.generate(self.base_case(), "current")
        authorization = next(out.glob("*委托书*.docx"))
        title, body = nonempty_paragraph_xml(authorization)[:2]
        title_run = first_text_run(title)
        body_run = first_text_run(body)

        title_fonts = title_run.find(W_NS + "rPr/" + W_NS + "rFonts")
        body_fonts = body_run.find(W_NS + "rPr/" + W_NS + "rFonts")
        title_size = title_run.find(W_NS + "rPr/" + W_NS + "sz")
        body_size = body_run.find(W_NS + "rPr/" + W_NS + "sz")
        title_spacing = title.find(W_NS + "pPr/" + W_NS + "spacing")
        body_spacing = body.find(W_NS + "pPr/" + W_NS + "spacing")

        self.assertEqual(title_fonts.get(W_NS + "eastAsia"), "黑体")
        self.assertEqual(title_size.get(W_NS + "val"), "32")
        self.assertEqual(body_fonts.get(W_NS + "eastAsia"), "仿宋_GB2312")
        self.assertEqual(body_size.get(W_NS + "val"), "28")
        self.assertEqual(title_spacing.get(W_NS + "line"), "570")
        self.assertEqual(title_spacing.get(W_NS + "lineRule"), "exact")
        self.assertEqual(body_spacing.get(W_NS + "line"), "570")
        self.assertEqual(body_spacing.get(W_NS + "lineRule"), "exact")

    def test_default_pack_rejects_unlisted_aid_center(self) -> None:
        case = self.base_case()
        case["aid_center"] = "其他地区法律援助中心"
        with self.assertRaises(subprocess.CalledProcessError):
            self.generate(case, "current")

    def test_all_stages_generate_without_unresolved_tokens(self) -> None:
        expected_minimum = {"侦查阶段": 3, "审查起诉阶段": 4, "审判阶段": 6}
        for stage, minimum in expected_minimum.items():
            with self.subTest(stage=stage):
                case = self.base_case()
                case["stage"] = stage
                out = self.generate(case, "all")
                generated = list(out.glob("*.docx"))
                self.assertGreaterEqual(len(generated), minimum)
                for path in generated:
                    self.assertNotRegex(docx_text(path), r"\{\{[^{}]+\}\}")

    def test_external_template_pack_is_supported(self) -> None:
        out = self.generate(self.base_case(), "current", external_pack=True)
        self.assertTrue(any(out.glob("*.docx")))

    def test_embedded_templates_survive_binary_asset_filtering(self) -> None:
        source_pack = ROOT / "assets" / "template-packs" / "wenzhou"
        filtered_pack = Path(tempfile.mkdtemp(prefix="legal-aid-filtered-pack-"))
        shutil.copy2(source_pack / "manifest.json", filtered_pack / "manifest.json")
        for encoded in source_pack.glob("*.b64.txt"):
            shutil.copy2(encoded, filtered_pack / encoded.name)
        out = self.generate(self.base_case(), "current", template_pack_dir=filtered_pack)
        self.assertTrue(any(out.glob("*.docx")))

    def test_plea_note_does_not_assume_special_identity(self) -> None:
        case = self.base_case()
        case["include_plea_note"] = True
        out = self.generate(case, "current")
        plea = next(out.glob("*认罪认罚笔录*.docx"))
        text = docx_text(plea)
        self.assertNotIn("未成年人、聋哑人", text)
        self.assertIn("测试人民检察院", text)

    def test_generation_creates_lawyer_review_checklist(self) -> None:
        case = self.base_case()
        case["assignment_date"] = "2026-07-01"
        out = self.generate(case, "archive")
        checklist = out / "00-生成结果律师复核清单.md"
        self.assertTrue(checklist.is_file())
        content = checklist.read_text(encoding="utf-8")
        self.assertIn("文件已生成，只表示系统完成了文书底稿制作", content)
        self.assertIn("起诉意见书作为阶段起点", content)
        self.assertIn("起诉书作为阶段起点", content)

    def test_dry_run_does_not_create_output(self) -> None:
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-dry-run-"))
        source = temp / "case.json"
        output = temp / "out"
        source.write_text(json.dumps(self.base_case(), ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(output), "--dry-run"],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        self.assertIn("只预览，不创建文件", result.stdout)
        self.assertFalse(output.exists())

    def test_timeline_extractor_marks_all_dates_unverified(self) -> None:
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-timeline-test-"))
        materials = temp / "materials"
        materials.mkdir()
        (materials / "流程记录.md").write_text(
            "2026年7月1日 接受法律援助指派\n"
            "2026年7月5日 会见并制作会见笔录\n"
            "2026年8月3日 检察院提起公诉\n"
            "2026年8月23日 律师收到起诉书\n",
            encoding="utf-8",
        )
        output = temp / "timeline"
        subprocess.run(
            [sys.executable, str(TIMELINE), str(materials), "--output-dir", str(output), "--ocr", "never"],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        payload = json.loads((output / "timeline_candidates.json").read_text(encoding="utf-8"))
        self.assertEqual({item["status"] for item in payload["candidates"]}, {"待律师核实"})
        dated_events = {(item["date"], item["event"]) for item in payload["candidates"]}
        self.assertIn(("2026-08-03", "审判阶段开始"), dated_events)
        self.assertIn(("2026-08-23", "收到或送达文书"), dated_events)
        self.assertNotEqual("2026-08-03", "2026-08-23")

    def test_platform_compatibility_report_is_machine_readable(self) -> None:
        result = subprocess.run(
            [sys.executable, str(PLATFORM_CHECK), "--platform", "opencode", "--json"],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        report = json.loads(result.stdout)
        self.assertTrue(report["core_generation_ready"])
        self.assertEqual(report["platform"], "opencode")
        self.assertIn("目录名应与 SKILL.md 的 name 保持一致。", report["platform_notes"])
        self.assertIn("SKILL.md", report["required_files"])

    def test_common_charge_library_covers_four_categories(self) -> None:
        library = json.loads((ROOT / "references" / "charge-law-library.json").read_text(encoding="utf-8"))
        charges = library["charges"]
        self.assertGreaterEqual(len(charges), 20)
        self.assertEqual(
            {entry["category"] for entry in charges.values()},
            {"侵犯财产", "侵犯公民人身权利", "危害公共安全", "妨害社会管理秩序"},
        )
        for charge in ("盗窃罪", "故意杀人罪", "交通肇事罪", "帮助信息网络犯罪活动罪"):
            self.assertIn(charge, charges)
            self.assertTrue(charges[charge]["articles"])
            self.assertGreater(len(charges[charge]["meeting_summary"]), 45)

    def test_representative_built_in_charges_generate_without_manual_basis(self) -> None:
        for charge in ("盗窃罪", "故意伤害罪", "危险驾驶罪", "寻衅滋事罪"):
            with self.subTest(charge=charge):
                case = self.base_case()
                case["charge"] = charge
                out = self.generate(case, "current")
                meeting = next(out.glob("*会见笔录*.docx"))
                text = docx_text(meeting)
                self.assertIn(charge, text)
                self.assertIn("《中华人民共和国刑法》", text)
                self.assertNotIn("请结合案件具体罪名补充", text)

    def test_unknown_charge_still_requires_verified_manual_basis(self) -> None:
        case = self.base_case()
        case["charge"] = "测试罪名罪"
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-unknown-charge-"))
        source = temp / "case.json"
        source.write_text(json.dumps(case, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out")],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("尚未收入内置基础说明库", result.stderr)
        self.assertFalse((temp / "out").exists())

    def test_charge_library_validator_passes(self) -> None:
        result = subprocess.run(
            [sys.executable, str(CHARGE_VALIDATOR)],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        self.assertIn("23个罪名", result.stdout)

    def test_archive_mode_excludes_stage_base_docs(self) -> None:
        out = self.generate(self.base_case(), "archive")
        names = {p.name for p in out.iterdir()}
        self.assertFalse([n for n in names if "委托书" in n], "归档模式不得生成委托书，应使用已有原件")
        self.assertFalse([n for n in names if "会见笔录" in n], "归档模式不得生成会见笔录，应使用已有原件")
        self.assertTrue(any("承办通报" in n for n in names))
        self.assertTrue(any("结案报告表" in n for n in names))
        self.assertTrue(any("归档目录" in n for n in names))

    def test_archive_keeps_remarks_blank_by_default(self) -> None:
        # 对外语书备注栏默认留空；填报依据等工作过程说明只进内部底稿（见 SKILL.md 边界）。
        out = self.generate(self.base_case(), "archive")
        doc = next(out.glob("*结案报告表*.docx"))
        from docx import Document as _Document

        for table in _Document(str(doc)).tables:
            for row in table.rows:
                cells = row.cells
                for i, cell in enumerate(cells[:-1]):
                    if cell.text.replace("\n", "").replace(" ", "").strip() == "备注":
                        self.assertEqual(cells[i + 1].text.strip(), "")

    def test_evidence_statement_uses_lawyer_verified_content(self) -> None:
        case = self.base_case()
        case["stage"] = "侦查阶段"
        case["handling_agency"] = "测试公安分局"
        case["evidence_statement_content"] = "因本案没有证据需自行取证，故无调查取证证据。"
        out = self.generate(case, "archive")
        statement = next(out.glob("*调查取证情况说明*.docx"))
        text = docx_text(statement)
        self.assertIn("因本案没有证据需自行取证", text)
        self.assertNotIn("请根据实际情况填写", text)

    def test_evidence_statement_missing_content_is_internal_warning(self) -> None:
        case = self.base_case()
        case["stage"] = "侦查阶段"
        case["handling_agency"] = "测试公安分局"
        out = self.generate(case, "archive")
        statement = next(out.glob("*调查取证情况说明*.docx"))
        self.assertNotIn("请根据实际情况填写", docx_text(statement))
        checklist = (out / "00-生成结果律师复核清单.md").read_text(encoding="utf-8")
        self.assertIn("尚未完成，不可直接提交", checklist)
        self.assertIn("调查取证", checklist)

    def test_closing_never_infers_meeting_from_assignment(self) -> None:
        for summary in ("", "经核实的基本案情"):
            case = self.base_case()
            case.update(assignment_date="2026-09-01", summary=summary)
            out = self.generate(case)
            content = docx_text(next(out.glob("*结案报告表*.docx")))
            self.assertNotIn("约谈受援人", content)
            self.assertNotIn("告知权利义务，案件风险", content)
            self.assertNotIn("会见日期为", content)
            self.assertIn("会见情况：", content)

    def test_closing_meeting_date_alone_does_not_create_actions(self) -> None:
        case = self.base_case()
        case["meeting_date"] = "2026-09-02"
        out = self.generate(case)
        content = docx_text(next(out.glob("*结案报告表*.docx")))
        self.assertNotIn("会见日期为", content)
        self.assertNotIn("告知权利义务，案件风险", content)

    def test_closing_uses_verified_meeting_without_basic_summary(self) -> None:
        case = self.base_case()
        case.update(meeting_date="2026-09-02",
                    meeting_summary="在律师事务所会见受援人，核实相关证据。",
                    defense_opinions="经核实的辩护意见")
        out = self.generate(case)
        content = docx_text(next(out.glob("*结案报告表*.docx")))
        self.assertIn("2026年9月2日", content)
        self.assertIn(case["meeting_summary"], content)
        self.assertIn(case["defense_opinions"], content)
        self.assertNotIn("看守所", content)

    def test_progress_report_defaults_are_blank_rows(self) -> None:
        out = self.generate(self.base_case(), "archive")
        report = next(out.glob("*承办通报*.docx"))
        text = docx_text(report)
        for fabricated in ("接受指派", "联系受援人", "委托手续", "会见笔录制作", "沟通案情", "告知案件结果"):
            self.assertNotIn(fabricated, text)
        self.assertNotIn("{{assignment_date}}", text)
        checklist = (out / "00-生成结果律师复核清单.md").read_text(encoding="utf-8")
        self.assertIn("案件承办通报默认输出空白记录行", checklist)

    def test_investigation_stage_plea_event_is_rejected(self) -> None:
        case = self.base_case()
        case["stage"] = "侦查阶段"
        case["handling_agency"] = "测试公安分局"
        case["progress"] = ["认罪认罚"]
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-plea-gate-"))
        source = temp / "case.json"
        source.write_text(json.dumps(case, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out")],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("侦查阶段不生成认罪认罚见证笔录", result.stderr)
        self.assertFalse((temp / "out").exists())

    def test_all_output_titles_use_title_font(self) -> None:
        out = self.generate(self.base_case(), "all")
        for path in sorted(out.glob("*.docx")):
            with self.subTest(doc=path.name):
                paragraphs = nonempty_paragraph_xml(path)
                self.assertTrue(paragraphs, path.name)
                run = first_text_run(paragraphs[0])
                fonts = run.find(W_NS + "rPr/" + W_NS + "rFonts")
                size = run.find(W_NS + "rPr/" + W_NS + "sz")
                self.assertEqual(fonts.get(W_NS + "eastAsia"), "黑体", path.name)
                self.assertEqual(size.get(W_NS + "val"), "32", path.name)

    def test_failed_validation_quarantines_outputs(self) -> None:
        source_pack = ROOT / "assets" / "template-packs" / "wenzhou"
        broken_pack = Path(tempfile.mkdtemp(prefix="legal-aid-broken-pack-"))
        shutil.copy2(source_pack / "manifest.json", broken_pack / "manifest.json")
        for item in source_pack.iterdir():
            if item.suffix in {".docx", ".xlsx"} or item.name.endswith(".b64.txt"):
                shutil.copy2(item, broken_pack / item.name)
        note = Document(broken_pack / "reading_note.docx")
        note.add_paragraph("原告测试残留")
        note.save(broken_pack / "reading_note.docx")
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-quarantine-"))
        source = temp / "case.json"
        source.write_text(json.dumps(self.base_case(), ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out"),
             "--template-pack-dir", str(broken_pack)],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 2)
        case_dir = next((temp / "out").iterdir())
        leaked = list(case_dir.glob("*.docx")) + list(case_dir.glob("*.xlsx"))
        self.assertEqual(leaked, [])
        quarantine = case_dir / "未通过校验-勿用"
        self.assertTrue(quarantine.is_dir())
        self.assertTrue(any(quarantine.glob("*.docx")))
        self.assertIn("隔离", result.stderr)

    def test_template_pack_missing_anchor_fails_fast(self) -> None:
        source_pack = ROOT / "assets" / "template-packs" / "wenzhou"
        mutated_pack = Path(tempfile.mkdtemp(prefix="legal-aid-anchor-pack-"))
        shutil.copy2(source_pack / "manifest.json", mutated_pack / "manifest.json")
        for item in source_pack.iterdir():
            if item.suffix in {".docx", ".xlsx"} or item.name.endswith(".b64.txt"):
                shutil.copy2(item, mutated_pack / item.name)
        meeting = Document(mutated_pack / "meeting_trial.docx")
        for paragraph in meeting.paragraphs:
            if "依法接受" in paragraph.text:
                set_paragraph_blank(paragraph)
        meeting.save(mutated_pack / "meeting_trial.docx")
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-anchor-case-"))
        source = temp / "case.json"
        source.write_text(json.dumps(self.base_case(), ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out"),
             "--template-pack-dir", str(mutated_pack)],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("锚点", result.stderr)
        self.assertFalse((temp / "out").exists())


def set_paragraph_blank(paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


class V6RegressionTests(unittest.TestCase):
    def base_case(self) -> dict:
        return {
            "recipient_name": "测试人员",
            "charge": "开设赌场罪",
            "stage": "审判阶段",
            "aid_center": "鹿城区法律援助中心",
            "handling_agency": "测试人民法院",
            "prosecuting_agency": "测试人民检察院",
            "custody_status": "取保候审",
            "lawyer": "李鸿鸿",
            "lawyer_phone": "",
            "law_firm": "浙江光正大律师事务所",
            "meeting_place": "浙江光正大律师事务所",
            "progress": ["归档", "取保候审"],
        }

    def test_progress_report_entry_overflow_is_rejected(self) -> None:
        case = self.base_case()
        case["case_progress_entries"] = [
            {"date": "2026-08-01", "method": "电话", "content": f"事项{i}"} for i in range(12)
        ]
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-entry-overflow-"))
        source = temp / "case.json"
        source.write_text(json.dumps(case, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out")],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("将被丢弃", result.stderr)
        self.assertFalse((temp / "out").exists())

    def test_quarantine_remnant_is_flagged_on_successful_rerun(self) -> None:
        source_pack = ROOT / "assets" / "template-packs" / "wenzhou"
        broken_pack = Path(tempfile.mkdtemp(prefix="legal-aid-remnant-pack-"))
        shutil.copy2(source_pack / "manifest.json", broken_pack / "manifest.json")
        for item in source_pack.iterdir():
            if item.suffix in {".docx", ".xlsx"} or item.name.endswith(".b64.txt"):
                shutil.copy2(item, broken_pack / item.name)
        note = Document(broken_pack / "reading_note.docx")
        note.add_paragraph("原告测试残留")
        note.save(broken_pack / "reading_note.docx")
        base = Path(tempfile.mkdtemp(prefix="legal-aid-remnant-out-"))
        source = base / "case.json"
        source.write_text(json.dumps(self.base_case(), ensure_ascii=False), encoding="utf-8")
        out = base / "out"
        broken = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(out),
             "--template-pack-dir", str(broken_pack)],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        self.assertEqual(broken.returncode, 2)
        good = subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(out)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=True,
        )
        self.assertIn("未通过校验-勿用", good.stderr)
        case_dir = next(out.iterdir())
        checklist = (case_dir / "00-生成结果律师复核清单.md").read_text(encoding="utf-8")
        self.assertIn("历史隔离文件", checklist)
        self.assertTrue((case_dir / "未通过校验-勿用").is_dir())

    def test_review_checklist_is_differentiated_by_doc_type(self) -> None:
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-checklist-diff-"))
        source = temp / "case.json"
        source.write_text(json.dumps(self.base_case(), ensure_ascii=False), encoding="utf-8")
        subprocess.run(
            [sys.executable, str(GENERATOR), "--input", str(source), "--output-dir", str(temp / "out")],
            cwd=ROOT,
            check=True,
            capture_output=True,
            text=True,
        )
        case_dir = next((temp / "out").iterdir())
        checklist = (case_dir / "00-生成结果律师复核清单.md").read_text(encoding="utf-8")
        self.assertIn("`02-会见笔录-法援-审判阶段.docx`：", checklist)
        self.assertIn("羁押口径", checklist)
        self.assertIn("摘录与卷宗原文逐页核对", checklist)
        self.assertIn("发问、质证和辩护意见提纲", checklist)
        self.assertIn("庭审程序记载与实际开庭一致", checklist)

    def test_charge_library_stale_reviewed_on_warns(self) -> None:
        library = json.loads((ROOT / "references" / "charge-law-library.json").read_text(encoding="utf-8"))
        library["reviewed_on"] = "2026-01-01"
        temp = Path(tempfile.mkdtemp(prefix="legal-aid-stale-library-"))
        stale = temp / "stale-library.json"
        stale.write_text(json.dumps(library, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(CHARGE_VALIDATOR), str(stale)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=True,
        )
        self.assertIn("已超过90天", result.stdout)
        self.assertEqual(result.returncode, 0)


if __name__ == "__main__":
    unittest.main()
