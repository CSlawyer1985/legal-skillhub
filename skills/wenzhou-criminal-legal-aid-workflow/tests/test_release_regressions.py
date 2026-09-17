from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
import test_generator as fixtures
from test_generator import ROOT, docx_text

sys.path.insert(0, str(ROOT / "scripts"))
import extract_case_timeline as timeline
import generate_criminal_legal_aid_set as generator
import validate_template_pack as pack_validator


class ReleaseRegressionTests(unittest.TestCase):
    base_case = fixtures.GeneratorTests.base_case
    generate = fixtures.GeneratorTests.generate

    def local_pack(self):
        temp = tempfile.TemporaryDirectory(prefix="legal-aid-local-test-")
        self.addCleanup(temp.cleanup)
        root = Path(temp.name) / "pack"
        shutil.copytree(ROOT / "assets/template-packs/wenzhou", root)
        return root

    def test_ocr_success_returns_text_chunk(self):
        for languages, expected in [({"chi_sim", "eng"}, "chi_sim+eng"), ({"chi_sim"}, "chi_sim")]:
            with self.subTest(languages=languages), \
                 patch.object(timeline.shutil, "which", return_value="/test/tesseract"), \
                 patch.object(timeline, "tesseract_languages", return_value=languages), \
                 patch.object(timeline, "run_command", return_value=subprocess.CompletedProcess([], 0, "2026年8月3日 会见笔录", "")) as run:
                chunk, error = timeline.ocr_image(Path("test.png"), "第1页")
                self.assertIsNone(error)
                self.assertIn("2026年8月3日", chunk.text)
                self.assertIn(expected, run.call_args.args[0])
                chunks, warnings = timeline.extract_file(Path("test.png"), "auto", 200)
                self.assertEqual(len(chunks), 1)
                self.assertFalse(warnings)

    def test_ocr_without_chinese_model_does_not_run(self):
        with patch.object(timeline.shutil, "which", return_value="/test/tesseract"), \
             patch.object(timeline, "tesseract_languages", return_value={"eng"}), \
             patch.object(timeline, "run_command") as run:
            chunk, error = timeline.ocr_image(Path("test.png"), "第1页")
            self.assertIsNone(chunk)
            self.assertIn("chi_sim", error)
            run.assert_not_called()

    def test_all_plea_aliases_rejected_in_investigation(self):
        for event in ("认罪认罚", "认罪", "plea"):
            for field in ("progress", "events"):
                case = self.base_case()
                case.update(stage="侦查阶段", progress=[])
                case[field] = [event]
                with self.subTest(event=event, field=field):
                    with self.assertRaisesRegex(ValueError, "侦查阶段不生成"):
                        generator.validate_case_data(case, {})
                    with self.assertRaisesRegex(ValueError, "侦查阶段不生成"):
                        generator.build_doc_plan(case, "current")

    def test_local_template_uses_variables_without_wenzhou_sentences(self):
        root = self.local_pack()
        manifest = json.loads((root / "manifest.json").read_text())
        manifest["renderers"] = {"meeting_trial": "placeholders"}
        (root / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
        doc = Document()
        doc.add_heading("会见笔录", 0)
        doc.add_paragraph("本地模板结构保持不变")
        doc.add_paragraph("被会见人：{{recipient_name}}；涉嫌罪名：{{charge}}")
        doc.add_paragraph("法律说明：{{charge_legal_basis}}")
        doc.add_paragraph("答：")
        doc.save(root / "meeting_trial.docx")
        self.assertEqual(pack_validator.validate(root), [])
        out = self.generate(self.base_case(), "current", template_pack_dir=root)
        content = docx_text(next(out.glob("*会见笔录*.docx")))
        self.assertIn("本地模板结构保持不变", content)
        self.assertIn("测试人员", content)
        self.assertIn("第三百零三条", content)
        self.assertNotIn("{{", content)
        self.assertNotIn("最近在看守所", content)

    def test_unknown_or_missing_local_variables_fail_before_generation(self):
        root = self.local_pack()
        manifest = json.loads((root / "manifest.json").read_text())
        manifest["renderers"] = {"meeting_trial": "placeholders"}
        (root / "manifest.json").write_text(json.dumps(manifest))
        doc = Document()
        doc.add_paragraph("{{unsupported_field}}")
        doc.save(root / "meeting_trial.docx")
        errors = "\n".join(pack_validator.validate(root))
        self.assertIn("缺少必需变量", errors)
        self.assertIn("含未知变量", errors)

    def test_corrected_legacy_wording_does_not_require_old_typo(self):
        root = self.local_pack()
        doc = Document(root / "meeting_trial.docx")
        for paragraph in doc.paragraphs:
            if "你收否" in paragraph.text:
                paragraph.text = paragraph.text.replace("你收否", "你是否")
            if paragraph.text.strip().startswith("行辩论，"):
                paragraph.text = ""
        doc.save(root / "meeting_trial.docx")
        self.assertEqual(pack_validator.validate(root), [])
        out = self.generate(self.base_case(), "current", template_pack_dir=root)
        self.assertNotIn("你收否", docx_text(next(out.glob("*会见笔录*.docx"))))

    def test_statements_allow_procedures_not_performed(self):
        out = self.generate(self.base_case(), "archive")
        checklist = (out / "00-生成结果律师复核清单.md").read_text()
        self.assertIn("未开展时据实说明", checklist)
        self.assertIn("未召开时据实说明", checklist)
        self.assertNotIn("确已实际开展调查取证", checklist)
        self.assertNotIn("确已实际召开庭前会议", checklist)

    def test_malformed_progress_entries_rejected(self):
        for entries in ("错误字段类型", ["不是对象"]):
            with self.assertRaisesRegex(ValueError, "记录对象数组"):
                generator.validate_progress_entries({"case_progress_entries": entries}, {})

    def test_library_count_and_unknown_charge_gate(self):
        self.assertEqual(len(generator.load_charge_library()), 23)
        case = self.base_case()
        case["charge"] = "未收录的测试罪名"
        with self.assertRaisesRegex(ValueError, "罪名法律依据"):
            generator.validate_case_data(case, {})
        case["charge_legal_basis"] = "测试字段：须由律师另行核验，不是真实法律说明。"
        generator.validate_case_data(case, {})


if __name__ == "__main__":
    unittest.main()
