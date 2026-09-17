#!/usr/bin/env python3
from __future__ import annotations

import argparse
import atexit
import base64
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PACK_ROOT = SKILL_DIR / "assets" / "template-packs"
DEFAULT_TEMPLATE_PACK = "wenzhou"
CHARGE_LIBRARY_PATH = SKILL_DIR / "references" / "charge-law-library.json"

STAGE_ALIASES = {
    "侦查": "侦查阶段",
    "侦查阶段": "侦查阶段",
    "公安": "侦查阶段",
    "公安阶段": "侦查阶段",
    "审查起诉": "审查起诉阶段",
    "审查起诉阶段": "审查起诉阶段",
    "检察院": "审查起诉阶段",
    "检察院阶段": "审查起诉阶段",
    "起诉阶段": "审查起诉阶段",
    "审判": "审判阶段",
    "审判阶段": "审判阶段",
    "法院": "审判阶段",
    "法院阶段": "审判阶段",
    "一审": "审判阶段",
    "一审阶段": "审判阶段",
}
STAGE_FILE_LABELS = {
    "侦查阶段": "侦查阶段",
    "审查起诉阶段": "审查起诉阶段",
    "审判阶段": "审判阶段",
}
STAGE_TEMPLATE_KEYS = {
    "侦查阶段": "investigation",
    "审查起诉阶段": "prosecution",
    "审判阶段": "trial",
}
STAGE_AUTH_VALIDITY = {
    "侦查阶段": "侦查阶段结束止",
    "审查起诉阶段": "审查起诉阶段结束止",
    "审判阶段": "审理阶段结束止",
}
STAGE_ROLE_PHRASE = {
    "侦查阶段": "侦查阶段的辩护人",
    "审查起诉阶段": "审查起诉阶段的辩护人",
    "审判阶段": "一审审判阶段的辩护人",
}
STAGE_RECIPIENT_LABEL = {
    "侦查阶段": "犯罪嫌疑人",
    "审查起诉阶段": "犯罪嫌疑人",
    "审判阶段": "被告人",
}
REQUIRED_TEMPLATE_KEYS = {
    "authorization_investigation",
    "authorization_prosecution",
    "authorization_trial",
    "meeting_investigation",
    "meeting_prosecution",
    "meeting_trial",
    "plea_meeting",
    "reading_note",
    "trial_outline",
    "trial_record",
    "evidence_statement",
    "post_judgment_visit",
    "case_progress_report",
    "closing_report",
    "investigation_statement",
    "prosecution_statement",
    "pretrial_meeting_statement",
    "archive_directory",
}
TEMPLATES: dict[str, Path] = {}
ACTIVE_TEMPLATE_PACK: dict = {}
MATERIALIZED_TEMPLATE_ROOT: Path | None = None
BASE_DOCS = {
    "侦查阶段": ["authorization", "meeting"],
    "审查起诉阶段": ["authorization", "meeting", "reading_note"],
    "审判阶段": ["authorization", "meeting", "reading_note", "trial_outline", "trial_record"],
}
ARCHIVE_DOCS = {
    "侦查阶段": ["case_progress_report", "closing_report", "investigation_statement", "evidence_statement"],
    "审查起诉阶段": ["case_progress_report", "closing_report", "prosecution_statement"],
    "审判阶段": [
        "evidence_statement",
        "pretrial_meeting_statement",
        "case_progress_report",
        "closing_report",
    ],
}
EVENT_DOCS = {
    "认罪认罚": "plea_meeting",
    "认罪": "plea_meeting",
    "plea": "plea_meeting",
    "承办通报": "case_progress_report",
    "通报": "case_progress_report",
    "结案": "closing_report",
    "结案报告": "closing_report",
    "调查取证": "evidence_statement",
    "庭前会议": "pretrial_meeting_statement",
    "判决后会见": "post_judgment_visit",
    "回访": "post_judgment_visit",
    "判决后回访": "post_judgment_visit",
    "归档目录": "archive_directory",
}
# 生成器填充逻辑依赖的模板原文锚点：模板包与脚本按版本配套的硬约束。
# 校验失败说明模板被改版或替换后未复验，继续生成会静默产出未清洗的文书。
TEMPLATE_ANCHORS: dict[str, tuple[str, ...]] = {
    "authorization_investigation": ("根据法律的规定，经", "本委托书有效期自即日起至"),
    "authorization_prosecution": ("根据法律的规定，经", "本委托书有效期自即日起至"),
    "authorization_trial": ("根据法律的规定，经", "本委托书有效期自即日起至"),
    "meeting_investigation": (
        "依法接受", "是否同意", "侦查机关认为你涉嫌", "所涉嫌的罪名是",
        "律师的业务活动主要是", "最近在看守所", "通过看守所",
    ),
    "meeting_prosecution": (
        "依法接受", "是否同意", "公诉机关认为你涉嫌", "所涉嫌的罪名是",
        "律师的业务活动主要是", "最近在看守所", "通过看守所",
    ),
    "meeting_trial": (
        "依法接受", "是否同意", "{{charge_legal_basis}}", "《法律援助法》第35条",
        "检察院的起诉书", "律师的主要活动是", "通过近亲属联系我",
    ),
    "plea_meeting": ("会见人：", "被会见人：", "公诉机关：", "审查起诉阶段认罪认罚"),
    "reading_note": ("阅卷地点：",),
    "trial_outline": ("辩护人", "立提纲人"),
    "trial_record": ("（敲法槌）", "起诉书副本有没有收到"),
    "evidence_statement": ("说明人：",),
    "pretrial_meeting_statement": ("说明人：",),
    "investigation_statement": ("说明人：",),
    "prosecution_statement": ("说明人：",),
    "post_judgment_visit": ("会见人：", "案由："),
    "case_progress_report": ("承办情况",),
    "closing_report": ("承办机构", "所处阶段", "结案日", "承办情况"),
}
# 复核清单按文书类型差异化：每类文书写明该次交付最容易出错的那个检查点。
DOC_REVIEW_POINTS = {
    "authorization": "指派机构、承办阶段与委托书有效期一致；受援人身份信息与原件核对；签名、日期栏完整。",
    "meeting": "到案经过与自首、自动投案情节表述属实；羁押口径（在押/取保候审）与地点、联系方式一致；罪名解释与现行法源一致；答复栏已经受援人逐页核对签名捺印。",
    "plea_meeting": "具结书签署的自愿性、知情性已确认；量刑建议告知与检察院出具文本一致；公诉机关、检察员信息核对无误。",
    "reading_note": "摘录与卷宗原文逐页核对无遗漏；证据清单完整；阅卷时间、地点如实填写。",
    "trial_outline": "发问、质证和辩护意见提纲与阅卷、会见结论一致。",
    "trial_record": "庭审程序记载与实际开庭一致；合议庭、公诉人信息据实记录。",
    "evidence_statement": "先核实是否实际开展调查取证；未开展时据实说明，材料缺失不等于程序未发生，不得虚构取证行为。",
    "pretrial_meeting_statement": "先核实是否实际召开庭前会议；未召开时据实说明，材料缺失不等于程序未发生，不得虚构会议及笔录。",
    "investigation_statement": "内容反映侦查阶段真实办理情况，无默认套话。",
    "prosecution_statement": "内容反映审查起诉阶段真实办理情况，无默认套话。",
    "post_judgment_visit": "裁判送达或生效后进行回访；服判息诉、上诉权利告知和执行情况据实记录。",
    "case_progress_report": "承办记录逐条据实填写，无虚构工作痕迹。",
    "closing_report": "承办结果、结案日期与法律援助中心系统一致。",
    "archive_directory": "与援助中心系统目录核对；起诉意见书、起诉书按阶段分别归位。",
}
OUTPUT_NAMES = {
    "authorization": "委托书-法援",
    "meeting": "会见笔录-法援",
    "reading_note": "阅卷笔录-法援",
    "plea_meeting": "认罪认罚笔录-法援",
    "trial_outline": "出庭提纲-法援",
    "trial_record": "庭审笔录-法援",
    "evidence_statement": "调查取证情况说明-法援",
    "pretrial_meeting_statement": "庭前会议情况说明-法援",
    "post_judgment_visit": "回访笔录-法援-判决后",
    "case_progress_report": "案件承办通报-法援",
    "closing_report": "结案报告表-法援",
    "investigation_statement": "侦查阶段情况说明-法援",
    "prosecution_statement": "审查起诉阶段情况说明-法援",
    "archive_directory": "刑事承办卷归档目录",
}
TITLE_TEXTS = {
    "授权委托书",
    "会见笔录",
    "审判阶段（辩护）会见笔录",
    "阅卷笔录",
    "出庭提纲",
    "刑事庭审笔录",
    "情况说明",
    "法律援助案件承办情况通报/报告记录",
    "结案报告表",
}


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as fh:
        return json.load(fh)


def load_charge_library() -> dict[str, dict]:
    if not CHARGE_LIBRARY_PATH.is_file():
        raise ValueError(f"常见罪名基础说明库不存在：{CHARGE_LIBRARY_PATH}")
    payload = load_json(CHARGE_LIBRARY_PATH)
    charges = payload.get("charges")
    if not isinstance(charges, dict) or not charges:
        raise ValueError("常见罪名基础说明库中的 charges 必须是非空对象。")
    invalid = [name for name, entry in charges.items() if not isinstance(entry, dict) or not text(entry.get("meeting_summary"))]
    if invalid:
        raise ValueError("常见罪名基础说明库条目缺少 meeting_summary：" + "、".join(invalid))
    return charges


def materialize_embedded_template(candidate: Path) -> Path | None:
    encoded = candidate.with_name(candidate.name + ".b64.txt")
    if not encoded.is_file():
        return None
    try:
        payload = base64.b64decode("".join(encoded.read_text(encoding="ascii").split()), validate=True)
    except Exception as exc:
        raise ValueError(f"嵌入模板无法解码：{encoded}：{exc}") from exc
    if not payload.startswith(b"PK\x03\x04"):
        raise ValueError(f"嵌入模板不是有效的 OOXML ZIP：{encoded}")
    global MATERIALIZED_TEMPLATE_ROOT
    if MATERIALIZED_TEMPLATE_ROOT is None:
        MATERIALIZED_TEMPLATE_ROOT = Path(tempfile.mkdtemp(prefix="legal-aid-template-assets-"))
        atexit.register(shutil.rmtree, MATERIALIZED_TEMPLATE_ROOT, ignore_errors=True)
    target = MATERIALIZED_TEMPLATE_ROOT / candidate.name
    target.write_bytes(payload)
    return target


def load_template_pack(pack_name: str, pack_dir: str | None = None) -> tuple[dict, dict[str, Path]]:
    root = Path(pack_dir).expanduser().resolve() if pack_dir else (TEMPLATE_PACK_ROOT / pack_name).resolve()
    manifest_path = root / "manifest.json"
    if not manifest_path.is_file():
        raise ValueError(f"模板包缺少 manifest.json：{manifest_path}")
    manifest = load_json(manifest_path)
    mapping = manifest.get("templates")
    if not isinstance(mapping, dict):
        raise ValueError("模板包 manifest.json 的 templates 必须是对象。")
    missing = sorted(REQUIRED_TEMPLATE_KEYS - set(mapping))
    if missing:
        raise ValueError("模板包缺少模板键：" + "、".join(missing))
    templates: dict[str, Path] = {}
    for key in sorted(REQUIRED_TEMPLATE_KEYS):
        candidate = (root / str(mapping[key])).resolve()
        if root not in candidate.parents:
            raise ValueError(f"模板路径越出模板包目录：{key}")
        expected_suffix = ".xlsx" if key == "archive_directory" else ".docx"
        if candidate.suffix.lower() != expected_suffix:
            raise ValueError(f"模板扩展名错误：{key} 应为 {expected_suffix}")
        resolved = candidate if candidate.is_file() else materialize_embedded_template(candidate)
        if resolved is None:
            raise ValueError(f"模板文件及其嵌入副本均不存在：{key} -> {candidate}")
        templates[key] = resolved
    validate_template_content(manifest, templates)
    return manifest, templates


def renderer_for(manifest: dict, key: str) -> str:
    renderers = manifest.get("renderers", {})
    if not isinstance(renderers, dict):
        raise ValueError("manifest.renderers 必须是对象。")
    if set(renderers) - REQUIRED_TEMPLATE_KEYS:
        raise ValueError("manifest.renderers 包含未知模板键。")
    renderer = renderers.get(key, "wenzhou-legacy")
    if renderer not in {"wenzhou-legacy", "placeholders"}:
        raise ValueError(f"{key} 的 renderer 不支持：{renderer}")
    if key == "archive_directory" and renderer != "wenzhou-legacy":
        raise ValueError("archive_directory 是直接复制的 Excel 母版，不支持变量填充。")
    return renderer


def validate_template_content(manifest: dict, templates: dict[str, Path]) -> None:
    """旧温州模板检查填充锚点；当地变量模板只检查变量契约。"""
    anchor_errors: list[str] = []
    allowed = {"{{" + key + "}}" for key in (
        "aid_center", "handling_agency", "prosecuting_agency", "court", "law_firm",
        "lawyer", "lawyer_phone", "recipient_label", "recipient_name", "recipient_id_no",
        "charge", "charge_legal_basis", "meeting_place", "meeting_date", "meeting_start",
        "meeting_end", "assignment_date", "closing_date", "case_result", "summary",
        "sign_date_line", "case_progress_entries", "meeting_summary", "defense_opinions",
        "evidence_statement_content", "pretrial_meeting_statement_content",
    )}
    for key in sorted(templates):
        renderer = renderer_for(manifest, key)
        if key == "archive_directory":
            continue
        joined = "\n".join(paragraph.text for paragraph in iter_all_paragraphs(Document(templates[key])))
        if renderer == "placeholders":
            required = {"{{recipient_name}}"}
            if key.startswith("meeting_") or key == "plea_meeting":
                required |= {"{{charge}}", "{{charge_legal_basis}}"}
            if key == "case_progress_report":
                required.add("{{case_progress_entries}}")
            if key == "closing_report":
                required |= {"{{case_result}}", "{{summary}}"}
            tokens = set(re.findall(r"\{\{[^{}]+\}\}", joined))
            if required - tokens:
                anchor_errors.append(f"{key} 缺少必需变量：{'、'.join(sorted(required - tokens))}")
            if tokens - allowed:
                anchor_errors.append(f"{key} 含未知变量：{'、'.join(sorted(tokens - allowed))}")
            continue
        anchors = TEMPLATE_ANCHORS.get(key, ())
        missing_anchors = [anchor for anchor in anchors if anchor not in joined]
        if missing_anchors:
            anchor_errors.append(f"{key} 缺少：{'、'.join(missing_anchors)}")
    if anchor_errors:
        raise ValueError(
            "模板校验失败（旧模板原文锚点或当地模板变量不匹配）；当地新模板请按指南配置 placeholders，不必恢复温州原句：\n- "
            + "\n- ".join(anchor_errors)
        )


def text(value, default: str = "") -> str:
    if value is None:
        return default
    value = str(value).strip()
    return value if value else default


def truthy(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return text(value).lower() in {"1", "true", "yes", "y", "是", "有", "需要"}


def normalize_stage(value: object) -> str:
    raw = text(value)
    return STAGE_ALIASES.get(raw, raw)


def validate_case_data(case_data: dict, template_pack: dict) -> None:
    missing: list[str] = []
    for key, label in {
        "recipient_name": "受援人姓名",
        "charge": "涉嫌罪名",
        "stage": "案件阶段（侦查阶段 / 审查起诉阶段 / 审判阶段）",
        "aid_center": "指派机构",
        "handling_agency": "经办机关",
        "lawyer": "当前案件承办律师",
        "law_firm": "当前案件承办律师事务所",
    }.items():
        if not text(case_data.get(key)):
            missing.append(label)
    stage = normalize_stage(case_data.get("stage"))
    if text(case_data.get("stage")) and stage not in STAGE_TEMPLATE_KEYS:
        missing.append("可识别的案件阶段（侦查阶段 / 审查起诉阶段 / 审判阶段）")
    aid_center = text(case_data.get("aid_center"))
    supported_centers = template_pack.get("assigning_institutions", [])
    if supported_centers and "*" not in supported_centers and aid_center and aid_center not in supported_centers:
        missing.append("当前模板包支持的指派机构：" + " / ".join(supported_centers))
    charge = text(case_data.get("charge"))
    charge_library = load_charge_library()
    if charge and charge not in charge_library and not text(case_data.get("charge_legal_basis")):
        missing.append(
            "罪名法律依据：该罪名尚未收入内置基础说明库，请核验对应《刑法》条款及必要司法解释后填写 charge_legal_basis"
        )
    if stage == "侦查阶段" and any(
        EVENT_DOCS.get(token) == "plea_meeting"
        for token in collect_progress_tokens(case_data, "current")
    ):
        missing.append(
            "阶段与事件组合：侦查阶段不生成认罪认罚见证笔录；如需记录侦查阶段的认罪认罚权利告知，"
            "请移除 progress/events 中的「认罪认罚」，改在会见笔录中据实记载"
        )
    if missing:
        raise ValueError("以下刑事法援手续信息缺失或不匹配，请先补充后再生成：\n- " + "\n- ".join(missing))


def set_paragraph_text(paragraph, new_text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def remove_page_breaks(paragraph) -> None:
    for br in list(paragraph._p.iter(qn("w:br"))):
        parent = br.getparent()
        if parent is not None:
            parent.remove(br)


def remove_trailing_empty_paragraphs(doc: Document) -> None:
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        remove_paragraph(doc.paragraphs[-1])


def set_runs_black(paragraph) -> None:
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.highlight_color = None


def clear_static_footer_page_numbers(doc: Document) -> None:
    seen_parts: set[str] = set()
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            part_name = str(footer.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            for paragraph in footer.paragraphs:
                if re.fullmatch(r"\s*\d+\s*", paragraph.text):
                    set_paragraph_text(paragraph, "")


def scrub_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""


def replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    original = paragraph.text
    updated = original
    for old, new in replacements.items():
        updated = updated.replace(old, new)
    if updated != original:
        set_paragraph_text(paragraph, updated)


def iter_all_paragraphs(doc: Document):
    def iter_container(container):
        for paragraph in container.paragraphs:
            yield paragraph
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from iter_container(cell)

    yield from iter_container(doc)
    seen_parts: set[str] = set()
    for section in doc.sections:
        for container in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            yield from iter_container(container)


def normalized_title_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def is_document_title(value: str) -> bool:
    normalized = normalized_title_text(value)
    if normalized in {normalized_title_text(item) for item in TITLE_TEXTS}:
        return True
    if normalized.endswith("法律援助中心"):
        return True
    return normalized.endswith(
        (
            "法律援助中心侦查阶段会见笔录",
            "法律援助中心审查起诉阶段会见笔录",
            "法律援助中心审判阶段（辩护）会见笔录",
            "法律援助中心回访笔录",
        )
    )


def set_run_typeface(run, font_name: str, size_pt: float) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(key), font_name)
    r_fonts.set(qn("w:hint"), "eastAsia")


def paragraph_is_in_table(paragraph) -> bool:
    parent = paragraph._p.getparent()
    while parent is not None:
        if parent.tag == qn("w:tc"):
            return True
        parent = parent.getparent()
    return False


def apply_template_typography(doc: Document) -> None:
    typography = ACTIVE_TEMPLATE_PACK.get("typography")
    if not isinstance(typography, dict):
        return
    body_font = text(typography.get("body_font"), "仿宋_GB2312")
    body_size = float(typography.get("body_size_pt", 14))
    title_font = text(typography.get("title_font"), "黑体")
    title_size = float(typography.get("title_size_pt", 16))
    line_spacing = float(typography.get("line_spacing_pt", 28.5))

    normal_style = doc.styles["Normal"]
    normal_style.font.name = body_font
    normal_style.font.size = Pt(body_size)
    normal_r_pr = normal_style._element.get_or_add_rPr()
    normal_r_fonts = normal_r_pr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        normal_r_fonts.set(qn(key), body_font)
    normal_r_fonts.set(qn("w:hint"), "eastAsia")
    normal_style.paragraph_format.line_spacing = Pt(line_spacing)

    for paragraph in iter_all_paragraphs(doc):
        if not paragraph_is_in_table(paragraph):
            paragraph.paragraph_format.line_spacing = Pt(line_spacing)
        title = is_document_title(paragraph.text)
        font_name = title_font if title else body_font
        size_pt = title_size if title else body_size
        for run in paragraph.runs:
            set_run_typeface(run, font_name, size_pt)


def replace_everywhere(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in iter_all_paragraphs(doc):
        replace_in_paragraph(paragraph, replacements)


def blank_id(case_data: dict) -> str:
    return text(case_data.get("recipient_id_no"), "                  ")


def derive_prosecuting_agency(case_data: dict, stage: str) -> str:
    explicit = text(case_data.get("prosecuting_agency"))
    if explicit:
        return explicit
    if stage == "审查起诉阶段":
        return text(case_data.get("handling_agency"))
    return ""


def is_non_custodial(case_data: dict) -> bool:
    values: list[str] = [text(case_data.get("custody_status"))]
    for key in ("progress", "events"):
        value = case_data.get(key)
        if isinstance(value, list):
            values.extend(text(item) for item in value)
        else:
            values.append(text(value))
    joined = " ".join(values)
    return any(token in joined for token in ("取保候审", "未羁押", "监视居住"))


def case_context(case_data: dict) -> dict[str, str]:
    stage = normalize_stage(case_data.get("stage"))
    lawyer = text(case_data.get("lawyer"))
    handling_agency = text(case_data.get("handling_agency"))
    non_custodial = is_non_custodial(case_data)
    default_place = f"{text(case_data.get('law_firm'))}（具体地点待确认）" if non_custodial else "□鹿城区看守所号会见室 □其他地点："
    return {
        "stage": stage,
        "recipient_label": STAGE_RECIPIENT_LABEL[stage],
        "recipient": text(case_data.get("recipient_name")),
        "charge": text(case_data.get("charge")),
        "charge_legal_basis": text(case_data.get("charge_legal_basis")),
        "aid_center": text(case_data.get("aid_center")),
        "lawyer": lawyer,
        "lawyer_phone": text(case_data.get("lawyer_phone")),
        "law_firm": text(case_data.get("law_firm")),
        "law_firm_website": text(case_data.get("law_firm_website")),
        "record_lawyer": text(case_data.get("record_lawyer"), lawyer),
        "meeting_place": text(case_data.get("meeting_place"), default_place),
        "non_custodial": non_custodial,
        "meeting_date": text(case_data.get("meeting_date")),
        "meeting_start": text(case_data.get("meeting_start")),
        "meeting_end": text(case_data.get("meeting_end")),
        "handling_agency": handling_agency,
        "prosecuting_agency": derive_prosecuting_agency(case_data, stage),
        "prosecutor": text(case_data.get("prosecutor"), "                       "),
        "court": text(case_data.get("court"), handling_agency if stage == "审判阶段" else "人民法院"),
        "assignment_date": text(case_data.get("assignment_date")),
        "closing_date": text(case_data.get("closing_date")),
        "case_result": text(case_data.get("case_result")),
        "summary": text(case_data.get("summary")),
        "sign_date_line": text(case_data.get("sign_date_line"), "年   月  日"),
    }


def common_replacements(case_data: dict) -> dict[str, str]:
    c = case_context(case_data)
    stage = c["stage"]
    recipient = c["recipient"]
    charge = c["charge"]
    lawyer = c["lawyer"]
    aid_center = c["aid_center"]
    handling_agency = c["handling_agency"]
    recipient_label = c["recipient_label"]
    return {
        "{{aid_center}}": aid_center,
        "{{law_firm}}": c["law_firm"],
        "{{lawyer}}": lawyer,
        "{{lawyer_phone}}": c["lawyer_phone"],
        "{{recipient_label}}": recipient_label,
        "{{recipient_name}}": recipient,
        "{{recipient_id_no}}": blank_id(case_data),
        "{{charge}}": charge,
        "{{charge_legal_basis}}": charge_explanation(case_data),
        "{{handling_agency}}": handling_agency,
        "{{prosecuting_agency}}": c["prosecuting_agency"],
        "{{court}}": c["court"],
        "{{meeting_place}}": c["meeting_place"],
        "{{meeting_date}}": c["meeting_date"],
        "{{meeting_start}}": c["meeting_start"],
        "{{meeting_end}}": c["meeting_end"],
        "{{assignment_date}}": c["assignment_date"],
        "{{closing_date}}": c["closing_date"],
        "{{case_result}}": c["case_result"],
        "{{summary}}": c["summary"],
        "{{sign_date_line}}": c["sign_date_line"],
        "{{meeting_summary}}": text(case_data.get("meeting_summary")),
        "{{defense_opinions}}": text(case_data.get("defense_opinions")),
        "{{evidence_statement_content}}": text(case_data.get("evidence_statement_content")),
        "{{pretrial_meeting_statement_content}}": text(case_data.get("pretrial_meeting_statement_content")),
        "{{case_progress_entries}}": "\n".join(
            "｜".join(text(entry.get(field)) for field in ("date", "method", "content", "note"))
            for entry in (case_data.get("case_progress_entries") or [])
        ),
        "温州市鹿城区法律援助中心": aid_center,
        "鹿城区法律援助中心": aid_center,
        "XX律师事务所": c["law_firm"],
        "浙江***律师事务所": c["law_firm"],
        "浙江光正大律师事务所李鸿鸿律师": f"{c['law_firm']}{lawyer}律师",
        "浙江光正大律师事务所李鸿鸿律师为     罪案件犯罪嫌疑人       （公民身份号码：                  ） 的指定辩护人。": (
            f"{c['law_firm']}{lawyer}律师为{charge}案件{recipient_label}{recipient}"
            f"（公民身份号码：{blank_id(case_data)}）的指定辩护人。"
        ),
        "本委托书有效期自即日起至侦查阶段结束止。": f"本委托书有效期自即日起至{STAGE_AUTH_VALIDITY[stage]}。",
        "本委托书有效期自即日起至审查起诉阶段结束止。": f"本委托书有效期自即日起至{STAGE_AUTH_VALIDITY[stage]}。",
        "本委托书有效期自即日起至审理阶段结束止。": f"本委托书有效期自即日起至{STAGE_AUTH_VALIDITY[stage]}。",
        "网址： HYPERLINK \"http://www.gzdls.com/\" www.gzdls.com": f"网址：{c['law_firm_website']}",
        "时间:X年 X月 X日 X时X 分至 x 时  x 分": meeting_time_line(c, prefix="时间："),
        "时间：X年 X月 X日 X时X 分至 x 时  x 分": meeting_time_line(c, prefix="时间："),
        "地点：浙江***律师事务所": f"地点：{c['law_firm']}",
        "案由：                       记录人：": f"案由：{charge}                       记录人：",
        "温州市鹿城区人民检察院": c["prosecuting_agency"],
        "鹿城区人民检察院": c["prosecuting_agency"],
        "温州市鹿城区人民法院": c["court"],
        "鹿城区人民法院": c["court"],
        "***人民法院": c["court"],
        "被告人   犯    一案": f"被告人{recipient}犯{charge}一案",
        "被告人      辩护": f"被告人{recipient}辩护",
        "被告人：   ，": f"被告人：{recipient}，",
        "被告人：               ": f"被告人：{recipient}               ",
        "被告人              到庭": f"被告人{recipient}到庭",
        "经办机关：": f"经办机关：{handling_agency}",
        "承办人：李鸿鸿": f"承办人：{lawyer}",
        "说明人：李鸿鸿": f"说明人：{lawyer}",
        "说明人：        ": f"说明人：{lawyer}",
        "一审阶段的辩护人": STAGE_ROLE_PHRASE[stage],
        "一审审判阶段   的辩护人": STAGE_ROLE_PHRASE[stage],
        "审查起诉阶段的辩护人": STAGE_ROLE_PHRASE[stage],
        "侦查阶段的辩护人": STAGE_ROLE_PHRASE[stage],
        "主要代理意见/辩护意见": "主要辩护意见",
    }


def meeting_time_line(c: dict[str, str], prefix: str = "时间:") -> str:
    date = c["meeting_date"] or "____年__月__日 "
    start = c["meeting_start"] or "__时__分"
    end = c["meeting_end"] or "__时__分"
    return f"{prefix}{date}{start}至{end}"


def charge_explanation(case_data: dict) -> str:
    c = case_context(case_data)
    if c["charge_legal_basis"]:
        return c["charge_legal_basis"]
    entry = load_charge_library().get(c["charge"], {})
    return text(
        entry.get("meeting_summary"),
        "请结合案件具体罪名补充《中华人民共和国刑法》及相关司法解释、指导意见中关于犯罪构成和量刑幅度的规定。",
    )


def fill_authorization(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    seen_sign_date = False
    for paragraph in list(doc.paragraphs):
        stripped = paragraph.text.strip()
        if stripped.startswith("根据法律的规定，经") and "指定辩护人" in stripped:
            set_paragraph_text(
                paragraph,
                f"根据法律的规定，经{c['aid_center']}指派，{c['law_firm']}{c['lawyer']}律师为{c['charge']}案件{c['recipient_label']}{c['recipient']}（公民身份号码：{blank_id(case_data)}）的指定辩护人。",
            )
        elif stripped.startswith("本委托书有效期自即日起至"):
            set_paragraph_text(paragraph, f"本委托书有效期自即日起至{STAGE_AUTH_VALIDITY[c['stage']]}。")
        elif stripped.startswith("电话：") or stripped.startswith("电话:"):
            if c["lawyer_phone"]:
                set_paragraph_text(paragraph, f"电话：{c['lawyer_phone']}")
                set_runs_black(paragraph)
            else:
                remove_paragraph(paragraph)
        elif stripped.startswith("委托方：") and "受托方：" in stripped:
            set_paragraph_text(paragraph, f"委托方（签名）：    受托方：{c['law_firm']}")
        elif stripped.startswith("受委托律师："):
            set_paragraph_text(paragraph, f"受委托律师：{c['lawyer']}")
        elif "HYPERLINK" in paragraph.text or stripped.startswith("网址：") or stripped.startswith("网址:"):
            set_paragraph_text(paragraph, f"网址：{c['law_firm_website']}")
        elif re.fullmatch(r"年\s+月\s+日", stripped):
            set_paragraph_text(paragraph, f"                              {c['sign_date_line']}")
            seen_sign_date = True
        elif seen_sign_date and not stripped:
            remove_paragraph(paragraph)


def fill_meeting(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    stage = c["stage"]
    stage_title = {
        "侦查阶段": "侦查阶段会见笔录",
        "审查起诉阶段": "审查起诉阶段会见笔录",
        "审判阶段": "审判阶段（辩护）会见笔录",
    }[stage]
    for idx, paragraph in enumerate(doc.paragraphs):
        stripped = paragraph.text.strip()
        if idx <= 1 and "法律援助中心" in stripped:
            set_paragraph_text(paragraph, f"{c['aid_center']}{stage_title}")
        elif stripped.startswith("审判阶段（辩护）会见笔录"):
            set_paragraph_text(paragraph, "")
        elif stripped.startswith("时间:") and (stage == "审判阶段" or any([c["meeting_date"], c["meeting_start"], c["meeting_end"]])):
            set_paragraph_text(paragraph, meeting_time_line(c, prefix="时间:"))
        elif stripped.startswith("时间：") and (stage == "审判阶段" or any([c["meeting_date"], c["meeting_start"], c["meeting_end"]])):
            set_paragraph_text(paragraph, meeting_time_line(c, prefix="时间："))
        elif stripped.startswith("地点:") and (text(case_data.get("meeting_place")) or c["non_custodial"]):
            set_paragraph_text(paragraph, f"地点:{c['meeting_place']}")
        elif stripped.startswith("地点：") and (text(case_data.get("meeting_place")) or c["non_custodial"]):
            set_paragraph_text(paragraph, f"地点：{c['meeting_place']}")
        elif stage == "审判阶段" and ("其他地点:" in stripped or "其他地点：" in stripped):
            set_paragraph_text(paragraph, "")
        elif stage == "审判阶段" and re.match(r"^\s*答[:：]", paragraph.text):
            set_paragraph_text(paragraph, trial_blank_answer(paragraph.text))
        elif stage == "审判阶段" and "《法律援助法》第35条" in stripped:
            set_paragraph_text(paragraph, "律师:你依法获得了法律援助，根据《法律援助法》第35条等相关法律规定，你享有以下权利和义务。")
        elif "会见人:" in stripped and "被会见人" in stripped:
            set_paragraph_text(paragraph, f"    会见人: {c['lawyer']}  被会见人:   {c['recipient']}")
        elif "记录人:" in stripped and "涉嫌罪名" in stripped:
            set_paragraph_text(paragraph, f"    记录人：{c['record_lawyer']}    涉嫌罪名：{c['charge']}")
        elif stage == "审判阶段" and "你收否收到了检察院的起诉书" in stripped:
            set_paragraph_text(paragraph, "律师：你是否收到了检察院的起诉书？何时收到？")
        elif stripped.startswith("律师:") and "依法接受" in stripped and "是否同意" in stripped:
            set_paragraph_text(
                paragraph,
                f"律师:{c['law_firm']}依法接受{c['aid_center']}的指派，并指派本所{c['lawyer']}律师担任你{STAGE_ROLE_PHRASE[stage]}，提供无偿的法律服务（不收取任何费用），直至{STAGE_AUTH_VALIDITY[stage]}，你是否同意我担任你的辩护人？",
            )
        elif stripped.startswith("律师:根据《中华人民共和国刑事诉讼法》的规定，律师的业务活动主要是") or (
            stage == "审判阶段" and "律师的主要活动是" in stripped
        ):
            set_paragraph_text(paragraph, stage_work_notice(stage))
        elif stage == "审判阶段" and stripped.startswith("律师:你是什么时间、什么地点"):
            if c["non_custodial"]:
                set_paragraph_text(paragraph, "律师：你于何时、何地获知本案，并于何时因涉嫌什么罪名被采取何种强制措施？目前强制措施为何？")
            else:
                set_paragraph_text(paragraph, "律师：你是什么时间、什么地点，因涉嫌什么罪名被采取刑事拘留、逮捕等强制措施？")
        elif stage == "审判阶段" and stripped.startswith("拘留的？"):
            set_paragraph_text(paragraph, "")
        elif "侦查机关认为你涉嫌" in stripped and "所涉嫌的罪名是" in stripped:
            set_paragraph_text(
                paragraph,
                f"律师:侦查机关认为你涉嫌{c['charge']}，现在给你解释你所涉嫌的罪名的法律规定，若有疑问，请向我提出，我会尽量向你解释清楚法律上的规定。你所涉嫌的罪名是：{c['charge']}。法律的规定是这样的：{charge_explanation(case_data)}以上规定，你是否清楚了？",
            )
        elif "公诉机关认为你涉嫌" in stripped and "所涉嫌的罪名是" in stripped:
            set_paragraph_text(
                paragraph,
                f"律师:公诉机关认为你涉嫌{c['charge']}，我现在给你解释你所涉嫌的罪名的法律规定，若有疑问，请向我提出，我会尽量向你解释清楚法律上的规定。你所涉嫌的罪名是：{c['charge']}。法律的规定是这样的：{charge_explanation(case_data)}以上规定，你是否清楚了？",
            )
        elif stage == "审判阶段" and "经核验的相关法律规定为" in stripped:
            paragraph.paragraph_format.keep_together = True
        elif stripped.startswith("根据刑法第二百三十四条"):
            set_paragraph_text(paragraph, "")
        elif stage == "审判阶段" and "你向侦查机关供诉的事实" in stripped:
            set_paragraph_text(paragraph, "律师：你向侦查机关供述的事实与你现在说的是否一致？")
        elif "按照认罪认罚的量刑判决有期徒刑6个月到一年" in stripped:
            set_paragraph_text(
                paragraph,
                "律师：结合起诉书指控事实、相关证据材料以及你的供述和辩解，现辩护人告知你本案中可能存在的事实认定、证据采信、量刑幅度、认罪认罚、退赔谅解等风险，具体需结合阅卷和会见情况进一步分析，你清楚了吗？",
            )
        elif "我将你为你作" in stripped:
            set_paragraph_text(paragraph, "律师：根据我们刚才关于案情的沟通，结合相应的法律规定，我将根据事实和法律为你提出相应的辩护意见，你是否同意？")
            paragraph.paragraph_format.keep_together = True
        elif stage == "审判阶段" and "根据我们刚才关于案情的沟通" in stripped:
            paragraph.paragraph_format.keep_together = True
        elif stage == "审判阶段" and stripped.startswith("行辩论，"):
            set_paragraph_text(paragraph, "")
        if c["non_custodial"] and "最近在看守所的生活情况" in paragraph.text:
            set_paragraph_text(paragraph, "律师：你目前身体和生活情况如何？是否受到不法侵害？有无物品被扣押或其他事项需要律师协助？")
        elif c["non_custodial"] and "侦查机关对你刑事拘留后" in paragraph.text:
            set_paragraph_text(paragraph, "律师：在侦查过程中，有无刑讯逼供、诱供、威胁、非法限制人身自由或者剥夺申辩权等违法情形？")
        elif c["non_custodial"] and "你被刑事拘留后" in paragraph.text and "告知你有权" in paragraph.text and any(token in paragraph.text for token in ("辩护", "律师")):
            set_paragraph_text(paragraph, "律师：侦查机关是否依法告知你有权委托辩护人？何时告知？")
        elif c["non_custodial"] and any(token in paragraph.text for token in ("通过看守所", "通过近亲属联系我")):
            set_paragraph_text(paragraph, f"律师：今天会见到此结束。如需继续沟通，可通过电话或到{c['law_firm']}与律师联系。你是否清楚？")
    if doc.tables:
        table = doc.tables[0]
        if len(table.rows) >= 2 and len(table.rows[0].cells) >= 2 and len(table.rows[1].cells) >= 2:
            set_paragraph_text(table.rows[0].cells[0].paragraphs[0], f"会见人: {c['lawyer']}")
            set_paragraph_text(table.rows[0].cells[1].paragraphs[0], f"被会见人: {c['recipient']}")
            set_paragraph_text(table.rows[1].cells[0].paragraphs[0], f"记录人：{c['record_lawyer']}")
            set_paragraph_text(table.rows[1].cells[1].paragraphs[0], f"涉嫌罪名: {c['charge']}")
    ensure_arrival_question(doc)
    replace_everywhere(doc, common_replacements(case_data))
    for paragraph in iter_all_paragraphs(doc):
        set_runs_black(paragraph)
    clear_static_footer_page_numbers(doc)


def ensure_arrival_question(doc: Document) -> None:
    full_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    if "如何到案" in full_text:
        return
    anchor = None
    for paragraph in doc.paragraphs:
        text_value = paragraph.text.strip()
        if "什么时间、什么地点" in text_value and any(token in text_value for token in ("拘留", "逮捕", "取保候审", "强制措施")):
            anchor = paragraph
            break
    if anchor is None:
        for paragraph in doc.paragraphs:
            if "所涉嫌的罪名" in paragraph.text:
                anchor = paragraph
                break
    if anchor is None:
        return
    question = anchor.insert_paragraph_before("律师:请你陈述一下本案你是如何到案的？这关系到是否存在自首、自动投案等量刑情节。")
    answer = anchor.insert_paragraph_before("答:                                                                      ")
    question.style = anchor.style
    answer.style = anchor.style


def trial_blank_answer(original: str) -> str:
    prefix = "答：" if "答：" in original else "答:"
    indent = original[: len(original) - len(original.lstrip())]
    return f"{indent}{prefix}                                                                      "


def stage_work_notice(stage: str) -> str:
    if stage == "侦查阶段":
        return "律师:根据《中华人民共和国刑事诉讼法》的规定，侦查阶段律师的主要工作包括：向侦查机关了解涉嫌罪名和案件有关情况，提出意见；会见犯罪嫌疑人并了解案件情况；提供法律咨询，代理申诉、控告；申请变更强制措施；根据案件情况申请调取证据或提出法律意见。你听清楚了吗？"
    if stage == "审查起诉阶段":
        return "律师:根据《中华人民共和国刑事诉讼法》的规定，审查起诉阶段律师的主要工作包括：查阅、摘抄、复制案卷材料，向检察机关了解案件情况；会见犯罪嫌疑人并核实证据；提供法律咨询，代理申诉、控告；申请变更强制措施或者提出羁押必要性审查意见；根据案件情况向检察机关提交辩护意见、申请调取证据或者提出不起诉、罪轻、从轻减轻处罚等意见。你听清楚了吗？"
    return "律师:根据《中华人民共和国刑事诉讼法》的规定，审判阶段律师的主要工作包括：查阅案卷材料，会见被告人并核实案件事实和证据；围绕起诉书指控事实、证据和法律适用准备辩护意见；参加庭前会议、开庭审理，进行发问、质证、辩论；根据案件情况提交书面辩护意见，维护你的诉讼权利。你听清楚了吗？"


def fill_plea_meeting(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped.startswith("时间：") and any([c["meeting_date"], c["meeting_start"], c["meeting_end"]]):
            set_paragraph_text(paragraph, meeting_time_line(c, prefix="时间："))
        elif stripped.startswith("地点：") and text(case_data.get("meeting_place")):
            set_paragraph_text(paragraph, f"地点：{c['meeting_place']}")
        elif stripped.startswith("会见人："):
            set_paragraph_text(paragraph, f"会见人：{c['law_firm']}（法律援助中心）{c['lawyer']}律师")
        elif stripped.startswith("被会见人："):
            special_label = text(case_data.get("special_recipient_label"))
            special_text = f"（{special_label}）" if special_label else ""
            set_paragraph_text(paragraph, f"被会见人：{c['recipient']}{special_text} 性别：     年龄：      ")
        elif stripped.startswith("案由："):
            set_paragraph_text(paragraph, f"案由：{c['charge']}                       记录人：{c['record_lawyer']}")
        elif stripped.startswith("公诉机关："):
            set_paragraph_text(paragraph, f"公诉机关：{c['prosecuting_agency']}   检察员：{c['prosecutor']}")
        elif c["stage"] != "审查起诉阶段" and "审查起诉阶段认罪认罚" in stripped:
            set_paragraph_text(paragraph, paragraph.text.replace("审查起诉阶段认罪认罚", f"{c['stage']}认罪认罚"))
        elif c["non_custodial"] and "你是什么时间、什么地点" in stripped and any(token in stripped for token in ("拘留", "逮捕")):
            set_paragraph_text(paragraph, "问：你于何时、何地获知本案，并于何时因涉嫌何罪名被采取何种强制措施？目前强制措施为何？")
        elif c["non_custodial"] and "侦查机关对你刑事拘留后" in stripped:
            set_paragraph_text(paragraph, "问：在侦查过程中，有无刑讯逼供、诱供、威胁、非法限制人身自由或者剥夺申辩权等违法情形？")
    replace_everywhere(doc, common_replacements(case_data))


def fill_reading_note(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped == "阅卷地点：":
            set_paragraph_text(paragraph, f"阅卷地点：{c['law_firm']}")
        elif stripped == "涉嫌罪名：":
            set_paragraph_text(paragraph, f"涉嫌罪名：{c['charge']}")
        elif stripped.startswith("被告人基本情况") and c["stage"] == "审判阶段":
            set_paragraph_text(paragraph, f"被告人基本情况：{c['recipient']}")
        elif "时间：X年X月X日X时X分至X年X月X日X时X分，地点：X" in stripped:
            label = stripped.split("（", 1)[0]
            set_paragraph_text(paragraph, f"{label}（时间：____年__月__日__时__分至____年__月__日__时__分，地点：________）")
        elif "供述及辩解摘抄提炼证明了什么" in stripped:
            set_paragraph_text(paragraph, "供述、辩解摘录及证明内容：")
            set_runs_black(paragraph)
    replace_everywhere(doc, common_replacements(case_data))


def fill_trial_outline(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    replacements = {
        "一、辩护人（原告）发问": "一、辩护人发问提纲",
        "二、辩护人（原告）质证意见": "二、质证意见提纲",
        "三、被告人及辩护人出示证据（原告举证）：": "三、被告人及辩护人拟出示证据：",
        "四、辩护意见（原告代理意见）：": "四、辩护意见提纲：",
    }
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped in replacements:
            set_paragraph_text(paragraph, replacements[stripped])
        elif stripped.startswith("立提纲人："):
            set_paragraph_text(paragraph, f"立提纲人：{c['lawyer']}")
        elif re.fullmatch(r"年\s+月\s+日", stripped):
            set_paragraph_text(paragraph, c["sign_date_line"])
        remove_page_breaks(paragraph)
    replace_everywhere(doc, common_replacements(case_data))
    for paragraph in list(doc.paragraphs):
        if not paragraph.text.strip():
            remove_paragraph(paragraph)
    remove_trailing_empty_paragraphs(doc)


def fill_trial_record(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped.startswith("地点：") and text(case_data.get("hearing_place")):
            set_paragraph_text(paragraph, f"地点：{text(case_data.get('hearing_place'))}")
        elif stripped.startswith("案由："):
            set_paragraph_text(paragraph, f"案由：{c['charge']}")
        elif stripped.startswith("公诉人：") and text(case_data.get("prosecutor")):
            set_paragraph_text(paragraph, f"公诉人：{c['prosecutor']}")
        elif stripped.startswith("被告人：") and "指定辩护人：" in stripped:
            set_paragraph_text(paragraph, f"被告人：{c['recipient']}    指定辩护人：{c['lawyer']}")
        elif "（敲法槌）" in stripped and "现在开庭" in stripped:
            set_paragraph_text(paragraph, f"审判长（以下简称“审”）：（敲法槌）{c['court']}现在开庭。传被告人到庭。")
        elif "起诉书副本有没有收到" in stripped:
            set_paragraph_text(paragraph, f"审：{c['prosecuting_agency']}的起诉书副本有没有收到？何时收到的？")
            paragraph.paragraph_format.page_break_before = True
        elif stripped.startswith("审：根据《中华人民共和国刑事诉讼法》的规定") and "提起公诉" in stripped:
            set_paragraph_text(
                paragraph,
                f"审：根据《中华人民共和国刑事诉讼法》的规定，{c['court']}依法开庭审理由{c['prosecuting_agency']}提起公诉的被告人{c['recipient']}涉嫌{c['charge']}一案。本案审判人员、书记员及公诉人信息据实记录。受{c['aid_center']}指派，{c['law_firm']}{c['lawyer']}律师出庭为被告人{c['recipient']}辩护。",
            )
        elif stripped.startswith("公：") and "起诉书查明认定的案件事实" in stripped:
            set_paragraph_text(paragraph, f"公：宣读{c['prosecuting_agency']}起诉书（文号：________），主要内容如下：")
        elif stripped.startswith("本院认为") and "起诉书宣读完毕" in stripped:
            set_paragraph_text(paragraph, "公：起诉书宣读完毕。")
        elif stripped.startswith("审：根据") and "建议，本案适用简易程序" in stripped:
            set_paragraph_text(paragraph, "审：关于本案适用何种审理程序，听取被告人及辩护人的意见。")
        elif stripped.startswith("公：本院认为"):
            set_paragraph_text(paragraph, "公：")
        elif stripped.startswith("辩：好的"):
            set_paragraph_text(paragraph, "辩：")
        elif stripped == "三、法庭辩论":
            paragraph.paragraph_format.keep_with_next = True
        elif stripped.startswith("尊敬的审判长"):
            set_paragraph_text(paragraph, "辩护意见要点：")
        elif "以上辩护意见恳请法庭予以采纳" in stripped:
            set_paragraph_text(paragraph, "")
        remove_page_breaks(paragraph)
    replace_everywhere(doc, common_replacements(case_data))
    clear_static_footer_page_numbers(doc)
    remove_trailing_empty_paragraphs(doc)


def fill_post_judgment_visit(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped.startswith("时间：") and any([c["meeting_date"], c["meeting_start"], c["meeting_end"]]):
            set_paragraph_text(paragraph, meeting_time_line(c, prefix="时间："))
        elif stripped.startswith("地点：") and text(case_data.get("meeting_place")):
            set_paragraph_text(paragraph, f"地点：{c['meeting_place']}")
        elif stripped == "会见人：":
            set_paragraph_text(paragraph, f"会见人：{c['lawyer']}")
        elif stripped == "被会见人：":
            set_paragraph_text(paragraph, f"被会见人：{c['recipient']}")
        elif stripped == "案由：":
            set_paragraph_text(paragraph, f"案由：{c['charge']}")
    replace_everywhere(doc, common_replacements(case_data))


def fill_statement_doc(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    content = text(case_data.get("evidence_statement_content"))
    for paragraph in doc.paragraphs:
        stripped = paragraph.text.strip()
        if stripped == "说明人：":
            set_paragraph_text(paragraph, f"说明人：{c['lawyer']}")
        elif stripped.startswith("二〇") or re.fullmatch(r"\d{4}年\d{1,2}月\d{1,2}日", stripped):
            set_paragraph_text(paragraph, c["sign_date_line"])
        elif stripped.startswith("请根据实际情况填写调查取证"):
            # 未核实正文留空；填写提醒只进入独立复核清单。
            set_paragraph_text(paragraph, content)
        elif stripped.startswith("请根据实际情况填写是否召开庭前会议"):
            set_paragraph_text(paragraph, text(case_data.get("pretrial_meeting_statement_content")))
    replace_everywhere(doc, common_replacements(case_data))


def _cn_date(value: str) -> str:
    """ISO 日期转中文日期；非 ISO 格式原样返回。"""
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", value.strip())
    if not m:
        return value.strip()
    year, month, day = m.group(1), int(m.group(2)), int(m.group(3))
    return f"{year}年{month}月{day}日"


def fill_closing_summary_cell(cell, case_data: dict, c: dict) -> None:
    """结案报告表"承办情况小结"按模板四段式逐段填写，保留法援中心制式结构：
    1 会见情况；2 基本案情；3 主要辩护意见；4 案件结果。
    模板编号为 Word 自动编号，段落文本不含编号前缀，故按内容关键词匹配。
    辩护意见未提供时保留模板原句，不预填意见内容。"""
    opinions = text(case_data.get("defense_opinions"))
    interview = text(case_data.get("meeting_date"))
    meeting_summary = text(case_data.get("meeting_summary"))
    for paragraph in cell.paragraphs:
        stripped = paragraph.text.strip()
        if "约谈受援人" in stripped or stripped.startswith("会见情况"):
            # 日期不等于已实施证明；只写律师提供的实际会见内容。
            value = f"会见情况：{meeting_summary}" if meeting_summary else "会见情况："
            if meeting_summary and interview:
                value = f"会见情况：会见日期为{_cn_date(interview)}。{meeting_summary}"
            set_paragraph_text(paragraph, value)
        elif stripped.startswith("基本案情") and c["summary"]:
            set_paragraph_text(paragraph, f"基本案情：{c['summary']}")
        elif stripped.startswith("主要辩护意见") and opinions:
            set_paragraph_text(paragraph, f"主要辩护意见：{opinions}")
        elif stripped.startswith("案件结果") and c["case_result"]:
            set_paragraph_text(paragraph, f"案件结果：{c['case_result']}")


def fill_label_tables(doc: Document, case_data: dict) -> None:
    c = case_context(case_data)
    label_values = {
        "承办机构": c["law_firm"],
        "承办人分类": text(case_data.get("lawyer_type"), "社会律师"),
        "承办人": c["lawyer"],
        "受援人": c["recipient"],
        "案由": c["charge"],
        "办案机关": c["handling_agency"],
        "（单位）": c["handling_agency"],
        "所处阶段": c["stage"],
        "指派日期": c["assignment_date"],
        "结案日": c["closing_date"],
        "援助形式": text(case_data.get("aid_type"), "刑事辩护"),
        "承办结果": c["case_result"],
        "备注": text(case_data.get("remarks")),
    }
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            # 先采集写入前的单元格原文：合并单元格会被重复返回，写入后的值文本
            # 可能含其他标签词（如备注里出现"受援人"），据原文匹配才不会连锁覆盖。
            original = [c.text.replace("\n", "").replace(" ", "").strip() for c in cells]
            i = 0
            while i < len(cells) - 1:
                key = original[i]
                if len(key) > 12:
                    # 标签格都是2-6字短词；长文本格（如小结提示句"约谈受援人…"）
                    # 本身含标签字样，不作为标签候选，避免误写入合并单元格。
                    i += 1
                    continue
                matched = False
                for label, value in label_values.items():
                    if value and label in key:
                        set_paragraph_text(cells[i + 1].paragraphs[0], value)
                        matched = True
                        break
                i += 2 if matched else 1
            if cells:
                row_text = "".join(cell.text for cell in cells)
                if "承办情况" in row_text:
                    fill_closing_summary_cell(cells[-1], case_data, c)


def fill_case_progress_report(doc: Document, case_data: dict) -> None:
    replace_everywhere(doc, common_replacements(case_data))
    if not doc.tables:
        return
    table = doc.tables[0]
    entries = case_data.get("case_progress_entries")
    entries = entries if isinstance(entries, list) else []
    # 承办通报是要提交法援中心的对外材料：先清空模板自带的全部示例记录行，
    # 再回填律师在 case_progress_entries 中据实提供的条目；生成器不预填任何承办行为。
    for offset in range(1, len(table.rows)):
        row = table.rows[offset]
        blank_row = [str(offset), "", "", "", ""]
        for cell, value in zip(row.cells, blank_row):
            set_paragraph_text(cell.paragraphs[0], value)
    for offset, entry in enumerate(entries, start=1):
        if offset >= len(table.rows):
            break
        row = table.rows[offset]
        values = [
            str(entry.get("code", offset)),
            text(entry.get("date")),
            text(entry.get("method")),
            text(entry.get("content")),
            text(entry.get("note")),
        ]
        for cell, value in zip(row.cells, values):
            set_paragraph_text(cell.paragraphs[0], value)


def fill_generic_doc(doc: Document, case_data: dict, doc_key: str) -> None:
    if doc_key == "plea_meeting":
        fill_plea_meeting(doc, case_data)
    elif doc_key == "reading_note":
        fill_reading_note(doc, case_data)
    elif doc_key == "trial_outline":
        fill_trial_outline(doc, case_data)
    elif doc_key == "trial_record":
        fill_trial_record(doc, case_data)
    elif doc_key == "post_judgment_visit":
        fill_post_judgment_visit(doc, case_data)
    elif doc_key in {
        "evidence_statement",
        "pretrial_meeting_statement",
        "investigation_statement",
        "prosecution_statement",
    }:
        fill_statement_doc(doc, case_data)
    elif doc_key == "case_progress_report":
        fill_case_progress_report(doc, case_data)
    elif doc_key == "closing_report":
        replace_everywhere(doc, common_replacements(case_data))
        fill_label_tables(doc, case_data)
    else:
        replace_everywhere(doc, common_replacements(case_data))


def collect_progress_tokens(case_data: dict, mode: str) -> set[str]:
    tokens: set[str] = {mode}
    for key in ("progress", "events"):
        value = case_data.get(key)
        if isinstance(value, str):
            tokens.update(part.strip() for part in re.split(r"[,，/、;；\s]+", value) if part.strip())
        elif isinstance(value, list):
            tokens.update(text(part) for part in value if text(part))
    bool_event_map = {
        "include_plea_note": "认罪认罚",
        "plea": "认罪认罚",
        "include_archive_docs": "归档",
        "archive": "归档",
        "include_post_judgment_visit": "回访",
        "post_judgment_visit": "回访",
        "include_archive_directory": "归档目录",
    }
    for key, token in bool_event_map.items():
        if truthy(case_data.get(key)):
            tokens.add(token)
    return tokens


def build_doc_plan(case_data: dict, mode: str) -> list[str]:
    stage = normalize_stage(case_data.get("stage"))
    tokens = collect_progress_tokens(case_data, mode)
    if stage == "侦查阶段" and any(EVENT_DOCS.get(token) == "plea_meeting" for token in tokens):
        raise ValueError("侦查阶段不生成认罪认罚见证笔录；权利告知请记入会见笔录。")
    # 归档模式只生成归档文书：委托书、会见笔录等办理阶段文书使用已有原件，不在归档时重复生成。
    if mode == "archive":
        plan: list[str] = []
    else:
        plan = list(BASE_DOCS[stage])
    if mode in {"archive", "all"} or {"归档", "结案", "结案报告"} & tokens:
        plan.extend(ARCHIVE_DOCS[stage])
    if mode == "all" and stage == "审判阶段":
        plan.append("post_judgment_visit")
    if mode in {"archive", "all"} or "归档目录" in tokens:
        plan.append("archive_directory")
    for token, doc_key in EVENT_DOCS.items():
        if token in tokens:
            plan.append(doc_key)
    return dedupe_preserve_order(plan)


def dedupe_preserve_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def template_key_for(doc_key: str, stage: str) -> str:
    if doc_key == "authorization":
        return f"authorization_{STAGE_TEMPLATE_KEYS[stage]}"
    if doc_key == "meeting":
        return f"meeting_{STAGE_TEMPLATE_KEYS[stage]}"
    return doc_key


def output_name_for(doc_key: str, stage: str, order_no: int) -> str:
    ext = ".xlsx" if doc_key == "archive_directory" else ".docx"
    stage_suffix = f"-{STAGE_FILE_LABELS[stage]}" if doc_key in {"authorization", "meeting"} else ""
    return f"{order_no:02d}-{OUTPUT_NAMES[doc_key]}{stage_suffix}{ext}"


def generate_doc(case_data: dict, case_dir: Path, doc_key: str, order_no: int) -> Path:
    stage = normalize_stage(case_data.get("stage"))
    template_key = template_key_for(doc_key, stage)
    template = TEMPLATES[template_key]
    output = case_dir / output_name_for(doc_key, stage, order_no)
    if doc_key == "archive_directory":
        shutil.copy2(template, output)
        return output
    doc = Document(template)
    if renderer_for(ACTIVE_TEMPLATE_PACK, template_key) == "placeholders":
        # 只替换变量，不运行温州原句替换、表格定位或段落清理。
        replace_everywhere(doc, {k: v for k, v in common_replacements(case_data).items() if k.startswith("{{")})
    elif doc_key == "authorization":
        fill_authorization(doc, case_data)
    elif doc_key == "meeting":
        fill_meeting(doc, case_data)
    else:
        fill_generic_doc(doc, case_data, doc_key)
    apply_template_typography(doc)
    scrub_document_properties(doc)
    doc.save(output)
    return output


def validate_created_docs(paths: list[Path], case_data: dict) -> None:
    c = case_context(case_data)
    allowed_charge = c["charge"]
    stale_tokens = [
        "XX律师事务所",
        "浙江***律师事务所",
        "***人民法院",
        "第条第款",
        "法律规定规定",
        "法律咨咨询",
        "我将你为你作",
        "人民法院提起公诉",
        "法院提起公诉",
        "人民法院指派检察员",
        "法院指派检察员",
        "人民法院鹿检",
        "法院鹿检",
        "原告",
        "代理意见",
        "X年X月X日",
        "供述及辩解摘抄提炼证明了什么",
        "辩：好的",
        "刑事审判第1庭",
        "鹿检   部刑诉",
        "本院认为，………",
        "请根据实际情况填写",
    ]
    errors: list[str] = []
    for path in paths:
        if path.suffix.lower() != ".docx":
            continue
        doc = Document(path)
        full_text = "\n".join(paragraph.text for paragraph in iter_all_paragraphs(doc))
        hits = [token for token in stale_tokens if token and token in full_text]
        if re.search(r"\{\{[^{}]+\}\}", full_text):
            hits.append("存在未替换模板变量")
        if f"{allowed_charge}{allowed_charge}" in full_text or f"{allowed_charge}罪" in full_text and allowed_charge.endswith("罪"):
            hits.append("疑似罪名重复")
        if c["non_custodial"]:
            hits.extend(token for token in ("刑事拘留后", "通过看守所", "最近在看守所") if token in full_text)
        if hits:
            errors.append(f"{path.name}: " + "、".join(dict.fromkeys(hits)))
    if errors:
        raise ValueError("生成文档未通过脏内容检查，请先清理模板或替换规则：\n- " + "\n- ".join(errors))


def build_case_folder_name(case_data: dict) -> str:
    return f"{text(case_data.get('recipient_name'))}-{text(case_data.get('charge'))}-{normalize_stage(case_data.get('stage'))}"


def write_review_checklist(case_dir: Path, created: list[Path], case_data: dict, plan: list[str]) -> Path:
    c = case_context(case_data)
    checklist = case_dir / "00-生成结果律师复核清单.md"
    date_fields = (
        ("法律援助指派日期", c["assignment_date"]),
        ("会见日期", c["meeting_date"]),
        ("行政结案日期", c["closing_date"]),
    )
    lines = [
        "# 生成结果律师复核清单",
        "",
        "> 文件已生成，只表示系统完成了文书底稿制作，不表示相关程序行为、送达、会见、阅卷、开庭、裁判或结案已经发生。",
        "",
        "## 本次生成范围",
        "",
        f"- 案件阶段：{c['stage']}（待承办律师确认）",
        f"- 指派机构：{c['aid_center']}（待承办律师确认）",
        f"- 经办机关：{c['handling_agency']}（待承办律师确认）",
        f"- 羁押状态：{text(case_data.get('custody_status'), '未填写')}（待承办律师确认）",
    ]
    quarantine_dir = case_dir / "未通过校验-勿用"
    pending = []
    if "closing_report" in plan:
        for field, label in (
            ("meeting_date", "真实会见日期（不得由指派日期推定）"),
            ("meeting_summary", "已核实的实际会见内容；未填写时结案报告会见段留空"),
            ("summary", "基本案情"), ("defense_opinions", "主要辩护意见"),
            ("case_result", "案件结果"),
        ):
            if not text(case_data.get(field)):
                pending.append(label)
    for key, field, label in (
        ("evidence_statement", "evidence_statement_content", "调查取证情况说明正文"),
        ("pretrial_meeting_statement", "pretrial_meeting_statement_content", "庭前会议情况说明正文"),
    ):
        if key in plan and not text(case_data.get(field)):
            pending.append(label + "（文书正文已留空，核实后补齐）")
    if pending:
        lines.extend(["", "## 尚未完成，不可直接提交", ""])
        lines.extend(f"- [ ] {label}" for label in pending)
    if quarantine_dir.is_dir() and any(quarantine_dir.iterdir()):
        lines.append(
            "- ⚠ 检测到历史隔离文件：`未通过校验-勿用/`。本次生成未使用其中的任何内容；确认无用后删除，勿与本次产物混用。"
        )
    lines.extend(
        [
            "",
            "## 日期核对",
            "",
        ]
    )
    for label, value in date_fields:
        lines.append(f"- [ ] {label}：{value or '未填写'}（须对照原件或律师确认记录）")
    lines.extend(
        [
            "- [ ] 分开核对司法机关程序行为日期、律师实际收件日期和法援行政结案日期。",
            "- [ ] 审查起诉阶段以起诉意见书作为阶段起点；审判阶段以起诉书作为阶段起点，两份文书不得机械合并。",
            "",
            "## 文件逐项复核",
            "",
        ]
    )
    for path, doc_key in zip(created, plan):
        point = DOC_REVIEW_POINTS.get(doc_key)
        lines.append(f"- [ ] `{path.name}`：{point}" if point else f"- [ ] `{path.name}`")
    lines.extend(
        [
            "",
            "## 交付前硬校验",
            "",
            "- 案件承办通报默认输出空白记录行：承办记录须逐条据实填写真实承办行为、日期和方式，不得保留任何预填示例或虚构承办行为。",
            "- [ ] 逐项核对姓名、罪名、案号、机关、阶段、地点和程序状态。",
            "- [ ] 会见答复、供述辩解、认罪认罚态度、辩护意见和裁判结果均有原始材料依据。",
            "- [ ] OCR或自动提取内容均已回看原页；未确认内容仍标为“待核实”。",
            "- [ ] 签名、日期、盖章和系统上传所需页码已经人工检查。",
            "- [ ] 正文仿宋_GB2312四号、主标题黑体三号、固定行距28.5磅，并完成逐页视觉检查。",
            "- [ ] 不含其他案件内容、个人私章、律师证号、个人手机号或无关受援人信息。",
            "- [ ] 先保留可修改Word；律师确认后再导出最终PDF上传。",
            "",
        ]
    )
    checklist.write_text("\n".join(lines), encoding="utf-8")
    return checklist


def error_advice(message: str) -> str:
    if "锚点" in message or "不配套" in message:
        return "原温州模板须与旧填充逻辑配套；当地新模板请按指南选择 placeholders 并设置必需变量，不必恢复温州原句。运行 validate_template_pack.py 后再试生成。"
    if "manifest.json" in message or "模板" in message:
        return "先运行 python3 scripts/doctor.py 和模板包校验；其他地区须导入完整当地模板包。"
    if "信息缺失" in message or "不匹配" in message:
        return "打开案件JSON，按错误列表补齐或核实字段；不要用推测值填充。"
    if "脏内容检查" in message:
        return "根据列出的文件和残留词检查模板或替换规则，修正后重新生成。"
    if "No module named" in message or "docx" in message:
        return "安装依赖：python3 -m pip install -r requirements.txt，然后运行 doctor.py 复查。"
    return "先运行 python3 scripts/doctor.py；仍无法解决时保留完整错误信息和输入字段名进行排查。"


def quarantine_failed_outputs(case_dir: Path, created: list[Path]) -> Path:
    """校验失败即隔离：未通过脏内容检查的文件移入「勿用」子目录，避免被当作可用文书取走。"""
    failed_dir = case_dir / "未通过校验-勿用"
    failed_dir.mkdir(parents=True, exist_ok=True)
    for path in created:
        if path.is_file():
            shutil.move(str(path), failed_dir / path.name)
    return failed_dir


def validate_progress_entries(case_data: dict, templates: dict[str, Path]) -> None:
    """承办通报表格行数有限：条目超行会在生成时被丢弃，必须在落盘前拦截。"""
    entries = case_data.get("case_progress_entries")
    if entries is None:
        return
    if not isinstance(entries, list) or any(not isinstance(entry, dict) for entry in entries):
        raise ValueError("case_progress_entries 必须是记录对象数组；无记录请使用 []。")
    if not entries:
        return
    if renderer_for(ACTIVE_TEMPLATE_PACK, "case_progress_report") == "placeholders":
        return
    template = templates.get("case_progress_report")
    if template is None:
        return
    tables = Document(template).tables
    capacity = max(len(tables[0].rows) - 1, 0) if tables else 0
    if len(entries) > capacity:
        raise ValueError(
            f"案件承办通报数据行共 {capacity} 行，case_progress_entries 提供了 {len(entries)} 条，"
            f"超出 {len(entries) - capacity} 条将被丢弃；请精简到 {capacity} 条以内，"
            "或与法律援助中心确认能否分表提交后重新生成。"
        )


def main() -> None:
    parser = argparse.ArgumentParser(description="生成刑事法律援助全流程文书底稿。")
    parser.add_argument("--input", required=True, help="Path to case JSON")
    parser.add_argument("--output-dir", required=True, help="Base output directory")
    parser.add_argument(
        "--template-pack",
        default=DEFAULT_TEMPLATE_PACK,
        help="Bundled template-pack id under assets/template-packs (default: wenzhou).",
    )
    parser.add_argument(
        "--template-pack-dir",
        help="External template-pack directory. When set, it overrides --template-pack.",
    )
    parser.add_argument(
        "--mode",
        choices=["current", "initial", "archive", "all"],
        default="current",
        help="current/initial generate stage base docs; archive generates only closing/archive docs (authorization & meeting use existing originals); all adds every stage-relevant doc.",
    )
    parser.add_argument("--dry-run", action="store_true", help="只校验输入并显示文书计划，不创建任何输出文件。")
    args = parser.parse_args()

    global ACTIVE_TEMPLATE_PACK
    ACTIVE_TEMPLATE_PACK, templates = load_template_pack(args.template_pack, args.template_pack_dir)
    TEMPLATES.clear()
    TEMPLATES.update(templates)
    case_data = dict(load_json(Path(args.input)))
    case_data["stage"] = normalize_stage(case_data.get("stage"))
    validate_case_data(case_data, ACTIVE_TEMPLATE_PACK)
    validate_progress_entries(case_data, templates)

    plan = build_doc_plan(case_data, args.mode)
    if args.dry_run:
        print("校验通过；本次只预览，不创建文件。")
        print(f"案件阶段：{case_data['stage']}")
        print("拟生成文件：")
        for order_no, doc_key in enumerate(plan, start=1):
            print(f"- {output_name_for(doc_key, case_data['stage'], order_no)}")
        return

    case_dir = Path(args.output_dir) / build_case_folder_name(case_data)
    case_dir.mkdir(parents=True, exist_ok=True)

    created: list[Path] = []
    try:
        for doc_key in plan:
            created.append(generate_doc(case_data, case_dir, doc_key, len(created) + 1))
        validate_created_docs(created, case_data)
    except ValueError:
        failed_dir = quarantine_failed_outputs(case_dir, created)
        print(f"已将 {len(created)} 个未通过交付前校验的文件隔离到：{failed_dir}", file=sys.stderr)
        print("上述文件请勿使用；修正模板或输入后重新生成。", file=sys.stderr)
        raise
    checklist = write_review_checklist(case_dir, created, case_data, plan)

    stale_quarantine = case_dir / "未通过校验-勿用"
    if stale_quarantine.is_dir() and any(stale_quarantine.iterdir()):
        print(f"注意：输出目录中存在上次未通过校验的历史文件：{stale_quarantine}", file=sys.stderr)
        print("本次生成未使用其中的内容；确认无用后清理，勿与本次产物混用。", file=sys.stderr)

    print(checklist)
    for path in created:
        print(path)


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, json.JSONDecodeError, ValueError) as exc:
        print("生成失败：本次运行未产出可交付的正式文书，请勿使用任何已输出文件。", file=sys.stderr)
        print(f"问题：{exc}", file=sys.stderr)
        print(f"处理建议：{error_advice(str(exc))}", file=sys.stderr)
        raise SystemExit(2) from None
