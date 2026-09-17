#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path

from docx import Document


TEXT_SUFFIXES = {".md", ".txt", ".json", ".csv"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES | {".pdf", ".docx"}
DATE_PATTERNS = (
    re.compile(r"(?<!\d)(20\d{2})\s*[年./-]\s*(1[0-2]|0?[1-9])\s*[月./-]\s*(3[01]|[12]\d|0?[1-9])\s*日?"),
    re.compile(r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)"),
)
EVENT_RULES = (
    ("收到或送达文书", ("收到", "签收", "送达")),
    ("接受法律援助指派", ("指派通知书", "接受指派", "指派材料")),
    ("办理委托手续", ("授权委托书", "委托手续", "办理委托")),
    ("会见", ("会见笔录", "会见介绍信", "会见")),
    ("阅卷", ("阅卷笔录", "阅卷")),
    ("审查起诉阶段开始", ("起诉意见书", "移送审查起诉")),
    ("审判阶段开始", ("起诉书", "提起公诉")),
    ("认罪认罚", ("认罪认罚具结书", "认罪认罚")),
    ("调查取证", ("调查取证", "调取证据", "取证")),
    ("庭前会议", ("庭前会议" ,)),
    ("开庭审理", ("出庭通知书", "开庭", "庭审笔录")),
    ("作出裁判", ("判决书", "裁定书", "判决", "裁定")),
    ("结案", ("结案报告", "结案")),
)


@dataclass(frozen=True)
class TextChunk:
    source: str
    locator: str
    text: str
    extraction_method: str
    extraction_confidence: str


@dataclass(frozen=True)
class TimelineCandidate:
    date: str
    event: str
    context: str
    source: str
    locator: str
    extraction_method: str
    confidence: str
    status: str = "待律师核实"


def run_command(command: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(command, capture_output=True, text=True, errors="replace", check=False)


def read_text(path: Path) -> list[TextChunk]:
    content = path.read_text(encoding="utf-8", errors="ignore")
    chunks: list[TextChunk] = []
    for number, line in enumerate(content.splitlines(), start=1):
        if line.strip():
            chunks.append(TextChunk(path.name, f"第{number}行", line.strip(), "文本直接读取", "高"))
    return chunks


def read_docx(path: Path) -> list[TextChunk]:
    doc = Document(path)
    chunks: list[TextChunk] = []
    for number, paragraph in enumerate(doc.paragraphs, start=1):
        if paragraph.text.strip():
            chunks.append(TextChunk(path.name, f"正文第{number}段", paragraph.text.strip(), "DOCX直接读取", "高"))
    for table_no, table in enumerate(doc.tables, start=1):
        for row_no, row in enumerate(table.rows, start=1):
            for cell_no, cell in enumerate(row.cells, start=1):
                text_value = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if text_value:
                    chunks.append(
                        TextChunk(
                            path.name,
                            f"表{table_no}第{row_no}行第{cell_no}列",
                            text_value,
                            "DOCX直接读取",
                            "高",
                        )
                    )
    return chunks


def usable_pdf_text(value: str) -> bool:
    compact = re.sub(r"\s+", "", value)
    meaningful = sum(1 for char in compact if "\u4e00" <= char <= "\u9fff" or char.isalnum())
    return meaningful >= 20


def pdf_text_layer(path: Path) -> tuple[list[TextChunk], str | None]:
    executable = shutil.which("pdftotext")
    if not executable:
        return [], "未找到 pdftotext，无法读取PDF文字层"
    result = run_command([executable, "-layout", str(path), "-"])
    if result.returncode != 0:
        message = result.stderr.strip() or "未知错误"
        return [], f"PDF文字层提取失败：{message}"
    chunks: list[TextChunk] = []
    for page_no, page_text in enumerate(result.stdout.split("\f"), start=1):
        if usable_pdf_text(page_text):
            chunks.append(TextChunk(path.name, f"第{page_no}页", page_text, "PDF文字层", "高"))
    return chunks, None


def tesseract_languages() -> set[str]:
    executable = shutil.which("tesseract")
    if not executable:
        return set()
    result = run_command([executable, "--list-langs"])
    if result.returncode != 0:
        return set()
    return {line.strip() for line in result.stdout.splitlines()[1:] if line.strip()}


def ocr_image(path: Path, locator: str) -> tuple[TextChunk | None, str | None]:
    executable = shutil.which("tesseract")
    if not executable:
        return None, "未找到 tesseract，无法识别扫描图片"
    languages = tesseract_languages()
    if "chi_sim" not in languages:
        return None, (
            "未安装Tesseract简体中文模型chi_sim，已跳过OCR识别；"
            "英文模型对中文扫描件的识别结果不可用，请安装中文语言包后重新扫描"
        )
    language = "chi_sim+eng" if "eng" in languages else "chi_sim"
    result = run_command([executable, str(path), "stdout", "-l", language, "--psm", "6"])
    if result.returncode != 0:
        return None, f"OCR失败：{result.stderr.strip() or '未知错误'}"
    if not result.stdout.strip():
        return None, "OCR未识别到文字"
    return TextChunk(path.name, locator, result.stdout, f"本地Tesseract OCR（{language}）", "低"), None


def pdf_page_count(path: Path) -> int | None:
    executable = shutil.which("pdfinfo")
    if not executable:
        return None
    result = run_command([executable, str(path)])
    if result.returncode != 0:
        return None
    match = re.search(r"^Pages:\s+(\d+)", result.stdout, re.MULTILINE)
    return int(match.group(1)) if match else None


def ocr_pdf(path: Path, max_pages: int) -> tuple[list[TextChunk], list[str]]:
    renderer = shutil.which("pdftoppm")
    if not renderer:
        return [], ["未找到 pdftoppm，无法把扫描PDF渲染为图片"]
    if not shutil.which("tesseract"):
        return [], ["未找到 tesseract，无法识别扫描PDF"]
    page_count = pdf_page_count(path)
    if page_count and page_count > max_pages:
        return [], [f"扫描PDF共{page_count}页，超过单次OCR上限{max_pages}页；请分卷或调整 --max-ocr-pages"]
    chunks: list[TextChunk] = []
    issues: list[str] = []
    with tempfile.TemporaryDirectory(prefix="legal-aid-timeline-ocr-") as tmp:
        prefix = Path(tmp) / "page"
        result = run_command([renderer, "-jpeg", "-r", "180", str(path), str(prefix)])
        if result.returncode != 0:
            return [], [f"扫描PDF渲染失败：{result.stderr.strip() or '未知错误'}"]
        images = sorted(Path(tmp).glob("page-*.jpg"))
        if len(images) > max_pages:
            return [], [f"扫描PDF渲染出{len(images)}页，超过单次OCR上限{max_pages}页"]
        for page_no, image in enumerate(images, start=1):
            chunk, issue = ocr_image(image, f"第{page_no}页")
            if chunk:
                chunks.append(TextChunk(path.name, chunk.locator, chunk.text, chunk.extraction_method, chunk.extraction_confidence))
            if issue:
                issues.append(f"第{page_no}页：{issue}")
    return chunks, issues


def extract_file(path: Path, ocr_mode: str, max_ocr_pages: int) -> tuple[list[TextChunk], list[str]]:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES:
            return read_text(path), []
        if suffix == ".docx":
            return read_docx(path), []
        if suffix == ".pdf":
            chunks, issue = pdf_text_layer(path)
            if chunks or ocr_mode == "never":
                return chunks, [issue] if issue else []
            ocr_chunks, ocr_issues = ocr_pdf(path, max_ocr_pages)
            return ocr_chunks, ([issue] if issue else []) + ocr_issues
        if suffix in IMAGE_SUFFIXES:
            if ocr_mode == "never":
                return [], ["已按 --ocr never 跳过图片OCR"]
            chunk, issue = ocr_image(path, "整页")
            return ([chunk] if chunk else []), ([issue] if issue else [])
    except Exception as exc:
        return [], [f"读取失败：{exc}"]
    return [], [f"暂不支持的文件类型：{suffix or '无扩展名'}"]


def normalize_date(year: str, month: str, day: str) -> str | None:
    try:
        return date(int(year), int(month), int(day)).isoformat()
    except ValueError:
        return None


def classify_event(context: str, source: str) -> tuple[str, bool]:
    haystack = f"{source} {context}"
    for event, keywords in EVENT_RULES:
        if any(keyword in haystack for keyword in keywords):
            return event, True
    return "日期事项待分类", False


def compact_context(value: str, start: int, end: int, width: int = 80) -> str:
    left = max(0, start - width)
    right = min(len(value), end + width)
    return re.sub(r"\s+", " ", value[left:right]).strip()


def candidates_from_chunk(chunk: TextChunk) -> list[TimelineCandidate]:
    candidates: list[TimelineCandidate] = []
    for pattern in DATE_PATTERNS:
        for match in pattern.finditer(chunk.text):
            normalized = normalize_date(*match.groups())
            if not normalized:
                continue
            context = compact_context(chunk.text, match.start(), match.end())
            event, classified = classify_event(context, chunk.source)
            if chunk.extraction_method.startswith(("文本", "DOCX", "PDF")) and classified:
                confidence = "较高"
            elif chunk.extraction_confidence == "低":
                confidence = "低"
            else:
                confidence = "中"
            candidates.append(
                TimelineCandidate(
                    date=normalized,
                    event=event,
                    context=context,
                    source=chunk.source,
                    locator=chunk.locator,
                    extraction_method=chunk.extraction_method,
                    confidence=confidence,
                )
            )
    return candidates


def filename_candidates(path: Path) -> list[TimelineCandidate]:
    result: list[TimelineCandidate] = []
    for pattern in DATE_PATTERNS:
        for match in pattern.finditer(path.stem):
            normalized = normalize_date(*match.groups())
            if not normalized:
                continue
            event, _ = classify_event(path.stem, path.name)
            result.append(
                TimelineCandidate(
                    date=normalized,
                    event=event,
                    context=f"日期来自文件名：{path.name}",
                    source=path.name,
                    locator="文件名",
                    extraction_method="文件名识别",
                    confidence="低",
                )
            )
    return result


def dedupe_candidates(values: list[TimelineCandidate]) -> list[TimelineCandidate]:
    unique: dict[tuple[str, str, str, str, str], TimelineCandidate] = {}
    for item in values:
        key = (item.date, item.event, item.source, item.locator, item.context)
        unique[key] = item
    return sorted(unique.values(), key=lambda item: (item.date, item.event, item.source, item.locator))


def markdown_escape(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def write_outputs(
    output_dir: Path,
    root: Path,
    candidates: list[TimelineCandidate],
    issues: list[dict[str, str]],
    scanned_files: list[str],
) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "timeline_candidates.json"
    md_path = output_dir / "案件流程台账-待核实.md"
    payload = {
        "materials_root": str(root),
        "status": "候选时间线，未经律师确认，不得直接作为正式案件事实",
        "scanned_files": scanned_files,
        "candidates": [asdict(item) for item in candidates],
        "issues": issues,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = [
        "# 案件流程台账（待核实）",
        "",
        "> 本台账由本地材料扫描/OCR自动生成。所有事项均为候选线索；在律师核对原件前，不得写入正式文书或表述为已经发生的案件事实。",
        "",
        "## 候选时间线",
        "",
        "| 日期 | 候选事项 | 来源 | 位置 | 提取方式 | 置信度 | 状态 | 上下文 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    if candidates:
        for item in candidates:
            lines.append(
                "| " + " | ".join(
                    markdown_escape(value)
                    for value in (
                        item.date,
                        item.event,
                        item.source,
                        item.locator,
                        item.extraction_method,
                        item.confidence,
                        item.status,
                        item.context,
                    )
                ) + " |"
            )
    else:
        lines.append("| — | 未提取到候选日期 | — | — | — | — | 待人工检查 | — |")

    lines.extend(["", "## 扫描问题与遗漏风险", ""])
    if issues:
        lines.extend(f"- `{item['source']}`：{item['issue']}" for item in issues)
    else:
        lines.append("- 未检测到技术性读取问题；仍需人工核对材料完整性和日期含义。")
    lines.extend(
        [
            "",
            "## 律师确认要求",
            "",
            "- [ ] 对照原件确认日期、行为主体、程序含义和文书形成/收件关系。",
            "- [ ] 分开确认司法机关程序行为日期、律师实际收件日期和法援行政结案日期。",
            "- [ ] 起诉意见书只作为审查起诉阶段起点；起诉书作为审判阶段起点。",
            "- [ ] OCR低置信度事项逐页复核，不以文件名日期替代文书正文。",
            "- [ ] 经确认后再转录到正式流程台账或案件输入JSON；未确认事项继续保留“待核实”。",
            "",
        ]
    )
    md_path.write_text("\n".join(lines), encoding="utf-8")
    return md_path, json_path


def main() -> int:
    parser = argparse.ArgumentParser(description="从本地案件材料提取候选日期并生成待核实流程台账。")
    parser.add_argument("materials_dir", type=Path, help="案件材料目录；只读扫描，不修改原文件")
    parser.add_argument("--output-dir", type=Path, required=True, help="候选台账输出目录，必须位于Skill包之外")
    parser.add_argument("--ocr", choices=["auto", "never"], default="auto", help="扫描PDF/图片的本地OCR策略")
    parser.add_argument("--max-ocr-pages", type=int, default=200, help="单个扫描PDF允许OCR的最大页数")
    args = parser.parse_args()

    root = args.materials_dir.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    if not root.is_dir():
        raise SystemExit(f"扫描失败：案件材料目录不存在或不是文件夹：{root}\n处理建议：检查路径后重新运行。")
    if root == output_dir or root in output_dir.parents:
        raise SystemExit("扫描失败：输出目录不能位于案件材料目录内部。\n处理建议：请把 --output-dir 指向独立的工作底稿目录。")
    if args.max_ocr_pages < 1:
        raise SystemExit("扫描失败：--max-ocr-pages 必须大于0。")

    candidates: list[TimelineCandidate] = []
    issues: list[dict[str, str]] = []
    scanned_files: list[str] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.is_symlink():
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        relative = path.relative_to(root)
        scanned_files.append(str(relative))
        chunks, file_issues = extract_file(path, args.ocr, args.max_ocr_pages)
        candidates.extend(filename_candidates(path))
        for chunk in chunks:
            source_chunk = TextChunk(str(relative), chunk.locator, chunk.text, chunk.extraction_method, chunk.extraction_confidence)
            candidates.extend(candidates_from_chunk(source_chunk))
        issues.extend({"source": str(relative), "issue": issue} for issue in file_issues if issue)

    candidates = dedupe_candidates(candidates)
    md_path, json_path = write_outputs(output_dir, root, candidates, issues, scanned_files)
    print(f"扫描完成：{len(scanned_files)}个文件，提取{len(candidates)}条候选日期。")
    print(f"待核实台账：{md_path}")
    print(f"机器可读结果：{json_path}")
    if issues:
        print(f"注意：有{len(issues)}项读取/OCR问题，请查看台账中的‘扫描问题与遗漏风险’。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
