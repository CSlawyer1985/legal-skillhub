"""
律师事务所函生成脚本（告知法院代理关系）
"""
import argparse
import json
import os
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from user_config import load_config


def set_font(run, name_cn='仿宋_GB2312', name_en='Times New Roman', size=16, bold=False):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name_cn)


def add_p(doc, text, alignment=WD_ALIGN_PARAGRAPH.LEFT, font_cn='仿宋_GB2312',
          font_en='Times New Roman', size=16, bold=False, space_before=0, space_after=0,
          first_line_indent=None):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    run = p.add_run(text)
    set_font(run, font_cn, font_en, size, bold)
    return p


def generate(config, court_name, plaintiff, defendant, case_number, cause, stage, lawyer_name, output_dir, numbering_file):
    os.makedirs(output_dir, exist_ok=True)
    firm = config.get('law_firm', '律师事务所')

    numbering = {"year": 2026, "firm_abbr": "友安", "counters": {"intro_letter": 0, "law_firm_letter": 0, "contract": 0}}
    if os.path.exists(numbering_file):
        with open(numbering_file, 'r', encoding='utf-8') as f:
            numbering = json.load(f)
    current = numbering["counters"].get("law_firm_letter", 0)
    next_no = current + 1
    numbering["counters"]["law_firm_letter"] = next_no
    os.makedirs(os.path.dirname(os.path.abspath(numbering_file)), exist_ok=True)
    with open(numbering_file, 'w', encoding='utf-8') as f:
        json.dump(numbering, f, ensure_ascii=False, indent=2)
    letter_no = f'〔{numbering["year"]}〕{numbering["firm_abbr"]}民字第{next_no:04d}号'

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(3.7)
        section.bottom_margin = Cm(3.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.6)

    add_p(doc, firm, alignment=WD_ALIGN_PARAGRAPH.CENTER, font_cn='黑体', size=22, bold=True, space_after=4)
    add_p(doc, letter_no, alignment=WD_ALIGN_PARAGRAPH.CENTER, size=16, space_after=12)
    add_p(doc, f'{court_name}：', space_after=8)
    add_p(doc, '', size=8, space_after=4)
    text = f'贵院受理的{plaintiff}诉{defendant}的{cause}案件（案号：{case_number}），现{plaintiff}已委托本所{lawyer_name}律师为其{stage}阶段诉讼代理人，特此函告。'
    add_p(doc, text, first_line_indent=0.74, space_after=16)
    add_p(doc, '', size=8, space_after=4)
    add_p(doc, '此致', first_line_indent=0.74, space_after=20)
    add_p(doc, court_name, space_after=30)
    add_p(doc, f'{firm}（公章）', alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=16, space_after=4)
    add_p(doc, '    年  月  日', alignment=WD_ALIGN_PARAGRAPH.RIGHT, size=16)

    addr = config.get('address', '')
    phone = config.get('phone', '')
    add_p(doc, '', size=6, space_after=2)
    if addr:
        add_p(doc, f'地址：{addr}', size=14, space_after=1)
    if phone:
        add_p(doc, f'电话：{phone}', size=14)

    first_party = plaintiff.split('、')[0].split('与')[0].strip()
    filename = f'{output_dir}/律师事务所函_{first_party}.docx'
    doc.save(filename)
    print(f'已生成：{filename}')
    print(f'文号：{letter_no}')
    return filename


def main():
    parser = argparse.ArgumentParser(description='生成律师事务所函')
    parser.add_argument('--user-config', required=True)
    parser.add_argument('--court-name', required=True)
    parser.add_argument('--plaintiff', required=True)
    parser.add_argument('--defendant', required=True)
    parser.add_argument('--case-number', required=True)
    parser.add_argument('--cause', required=True)
    parser.add_argument('--stage', default='一审')
    parser.add_argument('--lawyer-name', default=None)
    parser.add_argument('--output-dir', default=None)

    args = parser.parse_args()
    config = load_config(args.user_config)
    if not config:
        print('错误：未找到用户配置，请先录入信息。')
        sys.exit(1)

    output_dir = args.output_dir or os.path.join(os.path.dirname(args.user_config), '..', 'output')
    numbering_file = config.get('numbering_file', os.path.join(output_dir, 'numbering.json'))
    lawyer_name = args.lawyer_name or config.get('name', '')
    generate(config, args.court_name, args.plaintiff, args.defendant,
             args.case_number, args.cause, args.stage, lawyer_name,
             output_dir, numbering_file)


if __name__ == '__main__':
    main()
