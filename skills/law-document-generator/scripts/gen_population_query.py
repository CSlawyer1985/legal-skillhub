"""
人口信息查询全套文书生成
每个被查询人：授权委托书1页 + 人口信息查询申请表1页（查询人姓名留空）
"""
import argparse
import json
import os
import sys
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from user_config import load_config


def fill_cell(cell, text, sz=14, bold=False, center=True, font_name='宋体'):
    tc = cell._tc
    for p_elem in tc.findall(qn('w:p')):
        tc.remove(p_elem)
    p = etree.SubElement(tc, qn('w:p'))
    if center:
        pPr = etree.SubElement(p, qn('w:pPr'))
        jc = etree.SubElement(pPr, qn('w:jc'))
        jc.set(qn('w:val'), 'center')
    r = etree.SubElement(p, qn('w:r'))
    rPr = etree.SubElement(r, qn('w:rPr'))
    rf = etree.SubElement(rPr, qn('w:rFonts'))
    for a in ['w:eastAsia', 'w:ascii', 'w:hAnsi', 'w:cs']:
        rf.set(qn(a), font_name)
    sz_elem = etree.SubElement(rPr, qn('w:sz'))
    sz_elem.set(qn('w:val'), str(sz * 2))
    szCs = etree.SubElement(rPr, qn('w:szCs'))
    szCs.set(qn('w:val'), str(sz * 2))
    if bold:
        etree.SubElement(rPr, qn('w:b'))
    t = etree.SubElement(r, qn('w:t'))
    t.set(qn('xml:space'), 'preserve')
    t.text = text


def insert_page_break_xml(body):
    pb_p = etree.SubElement(body, qn('w:p'))
    pb_r = etree.SubElement(pb_p, qn('w:r'))
    pb_br = etree.SubElement(pb_r, qn('w:br'))
    pb_br.set(qn('w:type'), 'page')


def add_run_songti(p, text, size=14):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = '宋体'
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = etree.SubElement(rPr, qn('w:rFonts'))
    for a in ['w:eastAsia', 'w:ascii', 'w:hAnsi', 'w:cs']:
        rf.set(qn(a), '宋体')
    return run


def make_entrustment_doc(config, pl_name, pl_id, def_name, def_id):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    firm = config.get('law_firm', '律师事务所')
    lawyer = config.get('name', '律师')
    license_no = config.get('license_number', '')
    address = config.get('address', '')
    phone = config.get('phone', '')

    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_songti(p, '授权委托书', size=16)
    p.runs[0].bold = True

    # 委托人
    table1 = doc.add_table(rows=2, cols=2)
    table1.style = 'Table Grid'
    fill_cell(table1.rows[0].cells[0], '委托人：', center=False)
    fill_cell(table1.rows[0].cells[1], pl_name, center=False)
    fill_cell(table1.rows[1].cells[0], f'公民身份号码/统一社会信用代码：{pl_id}', center=False, sz=12)

    # 受托人信息标头
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    add_run_songti(p, '受托人信息：')

    table2 = doc.add_table(rows=5, cols=2)
    table2.style = 'Table Grid'
    for i, (label, value) in enumerate([
        ('律师事务所', firm),
        ('律师姓名', lawyer),
        ('执业证号', license_no),
        ('联系地址', address),
        ('联系电话', phone),
    ]):
        fill_cell(table2.rows[i].cells[0], f'{label}：', center=False)
        fill_cell(table2.rows[i].cells[1], value, center=False)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    add_run_songti(p, f'本所接受本案原告{pl_name}的委托，指派{lawyer}律师担任代理人，现因办理诉讼事宜，根据《中华人民共和国律师法》规定，需查询以下信息。')

    # 被查询人
    table3 = doc.add_table(rows=2, cols=2)
    table3.style = 'Table Grid'
    fill_cell(table3.rows[0].cells[0], '被查询人姓名：', center=False)
    fill_cell(table3.rows[0].cells[1], def_name, center=False)
    fill_cell(table3.rows[1].cells[0], '被查询人身份号码：', center=False)
    fill_cell(table3.rows[1].cells[1], def_id, center=False)

    p = doc.add_paragraph()
    add_run_songti(p, '查询信息类型：□ 外省市户籍人口信息  □ 外省市户籍人口信息在沪居住信息')

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    add_run_songti(p, '本人承诺查询上述信息仅供办理前述事项使用。如将查询到的信息挪为他用或泄露他人的，本人愿意承担全部法律责任。')

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    add_run_songti(p, '此  呈')
    p = doc.add_paragraph()
    add_run_songti(p, '上海市公安局              派出所')
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    add_run_songti(p, f'委托人：{pl_name}')
    p = doc.add_paragraph()
    add_run_songti(p, '日  期：    年    月    日')
    return doc


def make_query_form_doc(config, def_name, def_id):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    firm = config.get('law_firm', '')
    office_phone = config.get('office_phone', '')
    phone = config.get('phone', '')
    license_no = config.get('license_number', '')

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_songti(p, '人口信息查询申请表', size=16)
    p.runs[0].bold = True

    table1 = doc.add_table(rows=4, cols=4)
    table1.style = 'Table Grid'
    fill_cell(table1.rows[0].cells[0], '查询人单位：', center=False)
    fill_cell(table1.rows[0].cells[1], firm, center=False)
    fill_cell(table1.rows[0].cells[2], '联系电话（单位）：', center=False)
    fill_cell(table1.rows[0].cells[3], office_phone, center=False)
    fill_cell(table1.rows[1].cells[0], '联系电话（本人）：', center=False)
    fill_cell(table1.rows[1].cells[1], phone, center=False)
    fill_cell(table1.rows[1].cells[2], '证件号码：', center=False)
    fill_cell(table1.rows[1].cells[3], license_no, center=False)
    fill_cell(table1.rows[2].cells[0], '查询人姓名：', center=False)
    fill_cell(table1.rows[2].cells[1], '', center=False)  # 留空供手签
    table1.rows[2].cells[2].merge(table1.rows[2].cells[3])
    table1.rows[3].cells[0].merge(table1.rows[3].cells[3])
    fill_cell(table1.rows[3].cells[0],
              '本人保证以上填写内容及所提交的申请材料真实、合法、有效。本人承诺查询上述信息仅供办理诉讼事务使用，不将查询到的信息挪为他用或泄露他人，否则愿意承担相应的法律责任。',
              center=False, sz=10)

    table2 = doc.add_table(rows=4, cols=2)
    table2.style = 'Table Grid'
    fill_cell(table2.rows[0].cells[0], '被查询人姓名', center=True)
    fill_cell(table2.rows[0].cells[1], def_name, center=False)
    fill_cell(table2.rows[1].cells[0], '被查询人身份号码', center=True)
    fill_cell(table2.rows[1].cells[1], def_id, center=False)
    fill_cell(table2.rows[2].cells[0], '查询用途', center=True)
    fill_cell(table2.rows[2].cells[1], '诉讼', center=False)
    fill_cell(table2.rows[3].cells[0], '查询内容', center=True)
    fill_cell(table2.rows[3].cells[1], '身份信息', center=False)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    add_run_songti(p, '查询人签名：_______________')
    p = doc.add_paragraph()
    add_run_songti(p, '日    期：    年    月    日')
    return doc


def append_doc_with_pagebreak(target_doc, source_doc):
    body = target_doc.element.body
    insert_page_break_xml(body)
    for child in list(source_doc.element.body):
        if child.tag == qn('w:sectPr'):
            continue
        body.append(deepcopy(child))


def generate(config, pl_name, pl_id, defendants, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    if not defendants:
        print('错误：请至少提供一个被查询人。')
        return

    def_name, def_id = defendants[0]
    master = make_entrustment_doc(config, pl_name, pl_id, def_name, def_id)
    qdoc = make_query_form_doc(config, def_name, def_id)
    append_doc_with_pagebreak(master, qdoc)

    for def_name, def_id in defendants[1:]:
        edoc = make_entrustment_doc(config, pl_name, pl_id, def_name, def_id)
        append_doc_with_pagebreak(master, edoc)
        qdoc = make_query_form_doc(config, def_name, def_id)
        append_doc_with_pagebreak(master, qdoc)

    first_def = defendants[0][0]
    suffix = '等' if len(defendants) > 1 else ''
    filename = f'{output_dir}/人口信息查询_{pl_name}_诉_{first_def}{suffix}.docx'
    master.save(filename)
    print(f'已生成：{filename}')
    print(f'被查询人数量：{len(defendants)}')

    list_file = f'{output_dir}/人口信息查询_{pl_name}_被查询人清单.txt'
    with open(list_file, 'w', encoding='utf-8') as f:
        f.write(f'原告/委托人：{pl_name}（{pl_id}）\n')
        f.write(f'代理律师：{config.get("name", "")}（{config.get("license_number", "")}）\n')
        f.write(f'律所：{config.get("law_firm", "")}\n')
        f.write(f'电话：{config.get("phone", "")}\n\n')
        f.write('被查询人信息：\n')
        for i, (dn, did) in enumerate(defendants, 1):
            f.write(f'  {i}. {dn}，身份证号：{did}\n')
    print(f'已生成清单：{list_file}')
    return filename


def main():
    parser = argparse.ArgumentParser(description='生成人口信息查询文书')
    parser.add_argument('--user-config', required=True)
    parser.add_argument('--plaintiff-name', required=True)
    parser.add_argument('--plaintiff-id', required=True)
    parser.add_argument('--defendants', action='append', dest='defs',
                        help='被查询人：姓名|身份证号')
    parser.add_argument('--defendant', help='被查询人姓名（单被告模式）')
    parser.add_argument('--defendant-id', help='被查询人身份证号（单被告模式）')
    parser.add_argument('--output-dir', default=None)

    args = parser.parse_args()
    config = load_config(args.user_config)
    if not config:
        print('错误：未找到用户配置，请先录入信息。')
        sys.exit(1)

    output_dir = args.output_dir or os.path.join(os.path.dirname(args.user_config), '..', 'output')

    defendants = []
    if args.defs:
        for d in args.defs:
            parts = d.split('|', 1)
            if len(parts) == 2:
                defendants.append((parts[0], parts[1]))
    elif args.defendant and args.defendant_id:
        defendants.append((args.defendant, args.defendant_id))

    if not defendants:
        print('错误：请提供被查询人信息。')
        sys.exit(1)

    generate(config, args.plaintiff_name, args.plaintiff_id, defendants, output_dir)


if __name__ == '__main__':
    main()
