#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""polish_docx.py — 律所版式 docx 后处理（ClientScope Pro 配套脚本）

用法: python3 polish_docx.py <docx路径> [--toc-title 目录] [--body-pt 12] [--table-pt 10.5]

做什么（对应 references/format-spec.md §1 排版标准）:
  1. Normal 样式: 宋体+Times New Roman、12pt、1.5倍行距、两端对齐、段前段后0
  2. Heading 1/2/3: 宋体加粗黑色、14/12/12pt
  3. 按文本正则把段落挂接标题样式（第X部分→H1居中；一、→H2；（一）开头的正文不挂样式仅缩进）
  4. 全部正文段落统一首行缩进2字符（firstLineChars=200）；封面段与标题不缩进
  5. 所有表格: 单元格10.5pt、1.2倍行距、宋体、去缩进
  6. 在目录标题后插入可点击 TOC 域（TOC \\o "1-3" \\h \\z \\u, dirty=true）
  7. settings.xml 加 updateFields=true（打开文档自动更新目录页码）

依赖: pip install python-docx（建议在受管 venv 内安装）
"""
import re, sys, argparse
import docx
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH


def set_fonts(el, ea="宋体", ascii_="Times New Roman"):
    rPr = el.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr'); el.insert(0, rPr)
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    rf.set(qn('w:ascii'), ascii_); rf.set(qn('w:hAnsi'), ascii_); rf.set(qn('w:eastAsia'), ea)


def polish(path, toc_title='目录', body_pt=12.0, table_pt=10.5):
    d = docx.Document(path)

    # 1) Normal
    st = d.styles['Normal']
    set_fonts(st.element); st.font.size = Pt(body_pt)
    pf = st.paragraph_format
    pf.line_spacing = 1.5; pf.space_after = Pt(0); pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 2) Headings
    for hn, sz in [('Heading 1', body_pt + 2), ('Heading 2', body_pt), ('Heading 3', body_pt)]:
        try:
            hs = d.styles[hn]; set_fonts(hs.element)
            hs.font.size = Pt(sz); hs.font.bold = True; hs.font.color.rgb = RGBColor(0, 0, 0)
        except KeyError:
            pass

    pat1 = re.compile(r'^第[一二三四五六七八九十]+部分')
    pat2 = re.compile(r'^[一二三四五六七八九十]+、')
    cover_keys = ('关于', '之', '报告', '二〇二', '律师首次拜访', '内部参考')
    toc_title_idx = None

    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if t.replace('\u00a0', '').replace(' ', '') == toc_title:
            toc_title_idx = i
            p.style = d.styles['Normal']  # 目录标题不进TOC
            for r in p.runs:
                r.font.bold = True; r.font.size = Pt(body_pt + 2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            continue
        if pat1.match(t):
            p.style = d.styles['Heading 1']
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.bold = True; set_fonts(r._element)
            continue
        if pat2.match(t) and len(t) < 60:
            p.style = d.styles['Heading 2']
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.bold = True; set_fonts(r._element)
            continue
        # 封面短句（居中且含关键词）不缩进
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER and any(k in t for k in cover_keys) and len(t) < 40:
            p.paragraph_format.first_line_indent = Pt(0)
            continue
        # 正文：统一首行缩进2字符
        ppr = p._element.get_or_add_pPr()
        for ind in ppr.findall(qn('w:ind')):
            ppr.remove(ind)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLineChars'), '200')
        ind.set(qn('w:firstLine'), str(int(body_pt * 2 * 20)))
        ppr.append(ind)

    # 5) 表格
    for tb in d.tables:
        for row in tb.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.2
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    ppr = p._element.get_or_add_pPr()
                    for ind in ppr.findall(qn('w:ind')):
                        ppr.remove(ind)
                    for r in p.runs:
                        r.font.size = Pt(table_pt); set_fonts(r._element)

    # 6) TOC 域
    if toc_title_idx is not None:
        toc_p = d.paragraphs[toc_title_idx]
        new_p = OxmlElement('w:p')
        toc_p._element.addnext(new_p)

        def fld(t, dirty=True):
            e = OxmlElement('w:r'); c = OxmlElement('w:fldChar')
            c.set(qn('w:fldCharType'), t)
            if dirty:
                c.set(qn('w:dirty'), 'true')
            e.append(c); return e

        instr = OxmlElement('w:r')
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
        it.text = ' TOC \\o "1-3" \\h \\z \\u '
        instr.append(it)
        sep = OxmlElement('w:r'); c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'separate'); sep.append(c)
        holder = OxmlElement('w:r'); wt = OxmlElement('w:t')
        wt.text = '目录将在打开文档时自动生成（如未生成请全选后按F9更新域）'
        holder.append(wt)
        for e in (fld('begin'), instr, sep, holder, fld('end')):
            new_p.append(e)

    # 7) 打开时更新域
    settings = d.settings.element
    if settings.find(qn('w:updateFields')) is None:
        uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true'); settings.append(uf)

    d.save(path)
    return path


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('--toc-title', default='目录')
    ap.add_argument('--body-pt', type=float, default=12)
    ap.add_argument('--table-pt', type=float, default=10.5)
    a = ap.parse_args()
    print('polished:', polish(a.docx, a.toc_title, a.body_pt, a.table_pt))
