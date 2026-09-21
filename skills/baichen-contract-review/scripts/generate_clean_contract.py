#!/usr/bin/env python3
"""
generate_clean_contract.py — 修正版干净合同 Word 文档生成工具

属于 合同审查 三文档交付体系中的「文档三」。
将审查阶段的所有修改一次性落位到干净文本，无修订痕迹，可直接签署。

用法（由 Agent 调用）
    from generate_clean_contract import generate_clean_contract
import os

    sections = [
        {"type": "title", "text": "退休返聘用工协议"},
        {"type": "article", "text": "第一条  协议性质与法律适用"},
        {"type": "body", "text": "1.1 甲乙双方一致确认..."},
        {"type": "signature", "party_a": "甲方名称", "party_b": "乙方名称"},
    ]

    generate_clean_contract(sections, output_path)

支持的 section type
- title: 居中大标题 (22pt 黑体加粗)
- subtitle: 居中副标题 (16pt 黑体加粗)
- chapter: 居中章标题 (14pt 黑体加粗)
- article: 条标题 (12pt 黑体加粗，首行缩进)
- body: 正文段落 (12pt 仿宋，首行缩进)
- body_no_indent: 正文段落 (无缩进)
- center: 居中文本 (12pt 仿宋)
- right: 右对齐文本 (12pt 仿宋)
- signature: 签署页 (需 party_a 和 party_b 参数)
- blank: 空行
"""

from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, sys

# ── 格式底座（page_setup / Normal 样式委托给 create_document）──
_sys_path = os.path.join(os.path.dirname(__file__), '..', '..', '非争议解决公共标准', 'references')
if _sys_path not in sys.path:
    sys.path.insert(0, _sys_path)
from format_docx import create_document
from format_base import *

# _docx_base 保留 set_font / make_signature_block 等自定义内容辅助
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '../_shared'))
from _docx_base import set_font, make_signature_block


def generate_clean_contract(sections, output_path=None):
    """
    生成修正版干净合同 Word 文档（无修订痕迹）。

    Args:
        sections: List of dicts，每个包含 'type' 和对应内容键
        output_path: 输出 .docx 文件路径（可选；None 时自动生成）

    Returns:
        str: output_path
    """
    # ── 委托 format_docx 处理 page_setup + Normal 样式 ──
    doc = create_document([], output_path=None)

    for sec in sections:
        sec_type = sec.get('type', 'body')

        if sec_type == 'title':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(sec['text'])
            set_font(run, TITLE_FONT, 22, bold=True)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = BODY_SIZE
            p.paragraph_format.space_after = Pt(6)

        elif sec_type == 'subtitle':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(sec['text'])
            set_font(run, TITLE_FONT, 16, bold=True)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)

        elif sec_type == 'chapter':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(sec['text'])
            set_font(run, TITLE_FONT, 14, bold=True)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = BODY_SIZE
            p.paragraph_format.space_after = Pt(6)

        elif sec_type == 'article':
            p = doc.add_paragraph()
            run = p.add_run(sec['text'])
            set_font(run, TITLE_FONT, 12, bold=True)
            p.paragraph_format.first_line_indent = Cm(0.85)
            p.paragraph_format.space_before = Pt(6)

        elif sec_type == 'body':
            p = doc.add_paragraph()
            run = p.add_run(sec['text'])
            set_font(run, FONT_VARIANT, 12)
            p.paragraph_format.first_line_indent = Cm(0.85)

        elif sec_type == 'body_no_indent':
            p = doc.add_paragraph()
            run = p.add_run(sec['text'])
            set_font(run, FONT_VARIANT, 12)
            p.paragraph_format.first_line_indent = Cm(0)

        elif sec_type == 'center':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(sec['text'])
            set_font(run, FONT_VARIANT, 12)
            p.paragraph_format.first_line_indent = Cm(0)

        elif sec_type == 'right':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(sec['text'])
            set_font(run, FONT_VARIANT, 12)
            p.paragraph_format.first_line_indent = Cm(0)

        elif sec_type == 'signature':
            party_a = sec.get('party_a', '【甲方】')
            party_b = sec.get('party_b', '【乙方】')
            parties = [
                {'label': '甲方', 'name': party_a},
                {'label': '乙方', 'name': party_b},
            ]
            make_signature_block(doc, parties)

        elif sec_type == 'blank':
            doc.add_paragraph()

    if output_path is None:
        output_path = build_filename('clean_contract', 'party', 'matter')
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    print(f"修正版合同已生成: {output_path}")
    return output_path


if __name__ == '__main__':
    # Demo
    sections = [
        {"type": "title", "text": "退休返聘用工协议"},
        {"type": "center", "text": "（超龄劳动者劳务用工协议）"},
        {"type": "article", "text": "甲方（用人单位）"},
        {"type": "body", "text": "单位名称【___________】"},
        {"type": "signature", "party_a": "甲方（盖章）", "party_b": "乙方（签字）"},
    ]
    generate_clean_contract(sections, '/sandbox/workspace/outputs/demo_clean_contract.docx')
