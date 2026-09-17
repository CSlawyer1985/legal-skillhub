# Maintained by Lu Lingyan, Deheng (Wuxi) Law Firm.
#!/usr/bin/env python3
"""
文件清单及盖章指引生成脚本
读取 references/文件清单及盖章指引模板.md 中的表格，生成 Word (.docx) 版文件清单。
格式与诉讼文书一致：宋体标题、仿宋正文、A4 页面。

用法：
  python3 generate_file_list.py <模板.md> <输出.docx> [--case "案件：XX诉XX XX纠纷案"] [--defendants N]
  --defendants N：被告人数（默认 1）。对模板中备注列标「随被告数」的行，份数自动 = N + 2（原告1+被告N+法院1）；默认 3 份，多一被告多一份。
"""
import re, sys, argparse
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


MARK_DEFENDANTS = "随被告数"  # 备注列标记：该行份数随被告人数浮动（仅内部识别，输出时剥离，不外显）


def parse_table(md_path, defendants=1, skip=None):
    """解析 markdown 表格为二维数组；备注含「随被告数」的行，份数按 被告人数+2 调整，
    并将「随被告数」标记从备注中剥离（内部规则不写入交付文件）。skip 为需剔除的序号列表。"""
    rows = []
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()
    in_table = False
    for line in lines:
        s = line.strip()
        if s.startswith("|"):
            if not in_table:
                in_table = True
            cells = [c.strip() for c in s.strip("|").split("|")]
            # 跳过分隔行 |---|
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            if skip and len(cells) >= 1 and cells[0] in skip:
                continue
            if len(cells) >= 5 and MARK_DEFENDANTS in cells[4]:
                cells[3] = str(defendants + 2)  # 原告1 + 被告N + 法院1
                cells[4] = re.sub(r"[；;]?\s*" + MARK_DEFENDANTS + r"\s*[；;]?", "", cells[4]).strip(" ；;")
            rows.append(cells)
        elif in_table:
            break
    return rows


def set_font(run, name, size, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def set_solid_borders(table):
    """设置表格为实线边框（single），替代 python-docx 默认的虚线/无边框"""
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr
    # 移除已有 borders 定义
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")      # 1pt
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def render(md_path, out_path, case_line="", defendants=1, skip=None, add_rows=None):
    # 表格
    rows = parse_table(md_path, defendants=defendants, skip=skip)
    if add_rows:
        for extra in add_rows:
            cells = [c.strip() for c in extra.split("|")]
            while len(cells) < 5:
                cells.append("")
            rows.append(cells[:5])

    # 排序规则（陆律师定，2026-08-27）：原件在前、复印件垫底——营业执照、法代身份证等复印件
    # 恒排最后两位；稳定性排序保证同类型内保持模板/追加行原有相对顺序
    if len(rows) > 1:
        head, body = rows[0], rows[1:]
        def copy_key(r):
            return 1 if (len(r) > 2 and r[2] == "复印件") else 0
        body = sorted(body, key=copy_key)
        rows = [head] + body

    # 先生成到临时路径 → self_check 自检通过后才落盘正式路径（不通过不存盘）
    tmp_path = out_path + ".tmp"
    _build(tmp_path, case_line, rows)
    ok, msg = self_check(tmp_path, rows)
    if not ok:
        import os
        os.remove(tmp_path)
        raise SystemExit(f"❌ self_check 未通过，已放弃生成：{msg}")
    print(f"✅ self_check 通过：{len(rows)} 行逐格一致")

    import os
    os.replace(tmp_path, out_path)
    print(f"✅ {out_path}")
    return out_path


def _build(out_path, case_line, rows):
    doc = Document()
    section = doc.sections[0]
    # 一页纸排版：收窄页边距
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.5)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)

    def _compact(p):
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        return p

    # 标题
    p = _compact(doc.add_paragraph())
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("文件清单"), "宋体", 16, bold=True)

    # 案件信息行（如有）
    if case_line:
        p = _compact(doc.add_paragraph())
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(case_line), "仿宋", 12)

    # 说明
    p = _compact(doc.add_paragraph())
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run("打印时按本清单份数准备，备注列标「公司盖章」的位置需加盖公章。"), "仿宋", 12)

    if rows:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_solid_borders(table)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = table.cell(ri, ci)
                cell.text = ""
                para = _compact(cell.paragraphs[0])
                is_header = (ri == 0)
                size = 11 if is_header else 10.5
                set_font(para.add_run(val), "宋体" if is_header else "仿宋", size, bold=is_header)

    doc.save(out_path)


def self_check(doc_path, expected_rows):
    """交付前自检：输出 docx 的表格与预期行逐格比对，不一致即失败"""
    tb = Document(doc_path).tables[0]
    actual = [[c.text.strip() for c in row.cells] for row in tb.rows]
    expect = [[str(c).strip() for c in row] for row in expected_rows]
    if len(actual) != len(expect):
        return False, f"行数不符：实际{len(actual)} vs 预期{len(expect)}"
    if any(len(a) != len(e) for a, e in zip(actual, expect)):
        return False, "存在列数不一致的行"
    for i, (a, e) in enumerate(zip(actual, expect)):
        if a != e:
            return False, f"第{i}行不符\n  实际: {a}\n  预期: {e}"
    # 内部规则不得外显
    joined = "\n".join(" | ".join(a) for a in actual)
    if MARK_DEFENDANTS in joined:
        return False, "内部标记「随被告数」泄露到交付文件"
    return True, f"{len(actual)} 行逐格一致"


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("template", help="模板 .md 路径")
    ap.add_argument("output", help="输出 .docx 路径")
    ap.add_argument("--case", default="", help="案件信息行（可选）")
    ap.add_argument("--defendants", type=int, default=1, help="被告人数（默认 1），标「随被告数」的行份数 = 被告人数 + 2")
    ap.add_argument("--skip", default="", help="需剔除的序号，逗号分隔（如 01-3）")
    ap.add_argument("--add", action="append", default=[], help='追加行，竖线分隔5列：序号|文件名称|原件/复印件|份数|备注；可多次传入')
    args = ap.parse_args()
    skip = [s.strip() for s in args.skip.split(",") if s.strip()]
    render(args.template, args.output, args.case, args.defendants, skip=skip, add_rows=args.add)
