#!/usr/bin/env python3
"""build_docx.py — 由 report.json 直出《案件检索报告》DOCX（含页脚页码）。

Skill作者：浙江金道律师事务所 龚家勇律师（微信：13967182079）

不走 Markdown 中间格式，避免转换环节静默丢弃表格与引用块。

用法：
    python build_docx.py --report _work/report.json --out ~/Desktop/案件检索报告_XXX_20260922.docx

report.json 结构：
{
  "title": "案件检索报告",
  "subtitle": "（可选副标题，一般写争议焦点简称）",
  "meta": {"检索主题": "...", "检索地域范围": "...", "检索时间范围": "...",
           "检索日期": "...", "数据库": "华宇元典法律数据"},
  "overview": ["段落一", "段落二"],
  "overview_table": {"columns": ["项目", "内容"], "rows": [["检索式", "..."]]},
  "table": {"columns": ["序号", "案号", ...], "rows": [["1", "（2023）...", ...]]},
  "cases": [
    {
      "heading": "案例一　（2023）浙01民终1234号",
      "flags": ["案号部分脱敏，援引前须另行核实"],
      "sections": [{"label": "基本案情", "text": "……"},
                   {"label": "裁判要旨", "text": "……"}],
      "quotes": [{"title": "判词原文摘录（节录）", "text": "……"}],
      "analysis": "……"
    }
  ],
  "conclusion": ["段落一", "段落二"],
  "risk": ["段落一"],                 // 缺省时自动补入两处固定声明
  "auto_declarations": true           // 默认 true：强制保证两处固定声明在位
}
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_HEAD = "黑体"
CN_BODY = "宋体"
CN_QUOTE = "仿宋"
EN_FONT = "Times New Roman"

DECLARATION_1A = (
    "声明一-A（未指定结果模式）：本报告按请求权基础、法律关系、案件事实三要素检索高匹配类案，"
    "未将裁判结果作为匹配要素，故所列案例中可能包含裁判结果与本报告委托人主张不一致的案例"
    "（已在「裁判结果方向」列逐例标注），援引时须自行甄别；本报告未设置针对不利先例的定向检索路径，"
    "所列案例与结论不得被理解为对相关争议全部类案裁判倾向的完整描述，亦不构成对案件结果的承诺。"
)
DECLARATION_1B = (
    "声明一-B（指定结果模式）：本报告仅检索与用户指定结果一致的案例，已剔除与指定结果相反者"
    "（不符案号见第一部分）；本报告未设置针对不利先例的定向检索路径，所列案例与结论不得被理解为"
    "对相关争议全部类案裁判倾向的完整描述，亦不构成对案件结果的承诺。"
)
DECLARATION_2 = (
    "声明二（案号与判词局限）：本报告所载案号均照录数据库返回原文，其中标注"
    "「案号部分脱敏」者，正式援引前须另行核实完整案号；判词摘录为节录文本，援引时以官方"
    "数据库全文为准。本报告结论受限于所设检索条件与数据库收录范围。"
)


def set_font(run, *, cn=CN_BODY, size=12.0, bold=False, color=None):
    run.font.name = EN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), EN_FONT)
    rFonts.set(qn("w:hAnsi"), EN_FONT)


def add_paragraph(doc, text, *, cn=CN_BODY, size=12.0, bold=False, align=None,
                  first_line_indent=None, space_after=6.0, color=None):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.5
    if first_line_indent:
        para.paragraph_format.first_line_indent = Pt(first_line_indent)
    run = para.add_run(text)
    set_font(run, cn=cn, size=size, bold=bold, color=color)
    return para


def add_quote(doc, text, *, title=None):
    if title:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(title)
        set_font(run, cn=CN_HEAD, size=11.0, bold=True, color=(0x44, 0x44, 0x44))
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent = Pt(24)
    pf.right_indent = Pt(12)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    run = para.add_run(text)
    set_font(run, cn=CN_QUOTE, size=11.5)
    return para


def add_heading(doc, text, *, level=1):
    size = {1: 15.0, 2: 13.0, 3: 12.0}.get(level, 12.0)
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    run = para.add_run(text)
    set_font(run, cn=CN_HEAD, size=size, bold=True)
    return para


def shade_cell(cell, fill="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(doc, columns, rows, *, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    for idx, col in enumerate(columns):
        header[idx].text = ""
        para = header[idx].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(str(col))
        set_font(run, cn=CN_HEAD, size=font_size, bold=True)
        shade_cell(header[idx])
    short_rows = 0
    for row in rows:
        cells = table.add_row().cells
        if len(row) != len(columns):
            short_rows += 1
        for idx in range(len(columns)):
            value = row[idx] if idx < len(row) else ""
            text = "" if value is None else str(value)
            cells[idx].text = ""
            para = cells[idx].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.25
            run = para.add_run(text)
            set_font(run, cn=CN_BODY, size=font_size)
    if short_rows:
        print(f"[警告] 有 {short_rows} 行数据的列数与表头不一致（表头 {len(columns)} 列），"
              f"缺项已补为空字符串；请核对 report.json，避免出现空单元格。", file=sys.stderr)
    return table


def add_page_footer(section, font_size=10.5):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)

    def field(instr: str):
        run = para.add_run()
        set_font(run, cn=CN_BODY, size=font_size)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = instr
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr_el)
        run._r.append(end)

    # 依次写入：第 [PAGE] 页 / 共 [NUMPAGES] 页
    for text, instr in (("第 ", None), (None, "PAGE"), (" 页 / 共 ", None),
                        (None, "NUMPAGES"), (" 页", None)):
        if instr:
            field(instr)
        else:
            run = para.add_run(text)
            set_font(run, cn=CN_BODY, size=font_size)
    return footer


def build(report: dict, out_path: Path) -> int:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)

    # 标题
    add_paragraph(doc, report.get("title", "案件检索报告"), cn=CN_HEAD, size=18.0, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    subtitle = report.get("subtitle")
    if subtitle:
        add_paragraph(doc, subtitle, cn=CN_HEAD, size=12.0,
                      align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # 头部五项
    meta = report.get("meta") or {}
    items = list(meta.items()) if isinstance(meta, dict) else list(meta)
    if items:
        add_table(doc, ["项目", "内容"], [[str(k), str(v if v not in (None, "") else "—")]
                                          for k, v in items])

    # 一、检索概况
    add_heading(doc, "一、检索概况", level=1)
    for para in report.get("overview", []) or ["（本次检索概况说明缺失）"]:
        add_paragraph(doc, para, first_line_indent=24.0)
    overview_table = report.get("overview_table")
    if overview_table:
        add_paragraph(doc, "", space_after=2)
        add_table(doc, overview_table.get("columns", []), overview_table.get("rows", []))

    # 二、案例匹配度一览表
    add_heading(doc, "二、案例匹配度一览表", level=1)
    table = report.get("table") or {}
    if table.get("rows"):
        add_table(doc, table.get("columns", []), table["rows"])
        mode = str(report.get("result_mode", "unspecified")).strip().lower()
        if mode == "specified":
            score_note = ("注：「要素匹配度」为请求权基础、法律关系、案件事实、结果相似度四要素加权得分"
                          "（满分 10 分，本地计算，已指定检索结果）；")
        else:
            score_note = ("注：「要素匹配度」为请求权基础、法律关系、案件事实三要素加权得分"
                          "（满分 8 分，本地计算，未指定检索结果，结果相似度不计入）；")
        add_paragraph(doc, score_note +
                      "「语义相关度」为数据库返回的语义相似度评分，区分度极低，不得作为排序依据。"
                      "两项指标口径不同，不得相互换算。",
                      size=10.5, color=(0x5A, 0x5A, 0x5A))
    else:
        add_paragraph(doc, "本次未检索到符合入报门槛的类案，一览表空缺。",
                      first_line_indent=24.0)

    # 三、逐案详析
    add_heading(doc, "三、逐案详析", level=1)
    cases = report.get("cases") or []
    if not cases:
        add_paragraph(doc, "本次无符合入报门槛的案例，本节空缺。", first_line_indent=24.0)
    for case in cases:
        add_heading(doc, case.get("heading", "（案例标题缺失）"), level=2)
        for flag in case.get("flags", []) or []:
            add_paragraph(doc, "⚠ " + str(flag), size=10.5, color=(0xA3, 0x2D, 0x2D))
        for sec in case.get("sections", []) or []:
            label = sec.get("label")
            text = sec.get("text", "")
            if label:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(2)
                run = para.add_run(f"{label}：")
                set_font(run, cn=CN_HEAD, size=12.0, bold=True)
                run2 = para.add_run(str(text))
                set_font(run2, cn=CN_BODY, size=12.0)
                para.paragraph_format.line_spacing = 1.5
                para.paragraph_format.first_line_indent = Pt(24)
            else:
                add_paragraph(doc, str(text), first_line_indent=24.0)
        for quote in case.get("quotes", []) or []:
            add_quote(doc, str(quote.get("text", "")), title=quote.get("title"))
        analysis = case.get("analysis")
        if analysis:
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(8)
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.first_line_indent = Pt(24)
            run = para.add_run("【分析意见】")
            set_font(run, cn=CN_HEAD, size=12.0, bold=True, color=(0x2E, 0x5F, 0xB0))
            run2 = para.add_run(str(analysis))
            set_font(run2, cn=CN_BODY, size=12.0)

    # 四、结论与建议
    add_heading(doc, "四、结论与建议", level=1)
    for para in report.get("conclusion", []) or ["（结论部分缺失）"]:
        add_paragraph(doc, para, first_line_indent=24.0)

    # 五、风险提示与检索局限
    add_heading(doc, "五、风险提示与检索局限", level=1)
    risk = list(report.get("risk") or [])
    if report.get("auto_declarations", True):
        # 声明一按匹配模式择一：未指定结果（三要素）／指定结果（含结果相似度）
        decl_1 = (DECLARATION_1B if str(report.get("result_mode", "unspecified")).strip().lower()
                  == "specified" else DECLARATION_1A)
        for decl in (decl_1, DECLARATION_2):
            if decl not in risk:
                risk.append(decl)
    for para in risk:
        add_paragraph(doc, para, first_line_indent=24.0)

    add_page_footer(section)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="由 report.json 生成《案件检索报告》DOCX")
    parser.add_argument("--report", required=True, help="report.json 路径")
    parser.add_argument("--out", required=True, help="输出 DOCX 路径（默认存桌面）")
    args = parser.parse_args()

    report_path = Path(args.report).expanduser()
    out_path = Path(args.out).expanduser()
    if not report_path.exists():
        print(f"[错误] report.json 不存在：{report_path}", file=sys.stderr)
        return 2

    with report_path.open("r", encoding="utf-8") as fh:
        report = json.load(fh)

    rc = build(report, out_path)
    if rc == 0:
        print(f"DOCX 已生成：{out_path}")
        print(f"入报案例 {len(report.get('cases') or [])} 例；"
              f"一览表 {len((report.get('table') or {}).get('rows', []) or [])} 行")
    return rc


if __name__ == "__main__":
    sys.exit(main())
