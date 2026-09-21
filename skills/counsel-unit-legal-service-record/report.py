#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
报告生成模块：根据筛选后的记录生成 DOCX（含页脚页码）与 Excel 汇总报告。
支持两种模式：
  - 单表模式（默认）：一张总表（或单单位分表）。
  - 分表模式（split_by_unit=True）：按单位拆分为多张独立分表
    （DOCX 为每个单位生成独立小节；Excel 为每个单位生成独立 worksheet）。
"""
import io
from datetime import datetime
from collections import OrderedDict

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 主题色（与前端一致，法务简洁风）
C_TITLE = RGBColor(0x1F, 0x3A, 0x5F)      # 深蓝
C_HEADER_BG = "1F3A5F"
C_HEADER_FONT = "FFFFFF"
C_BAND = "EEF2F7"                          # 斑马纹浅灰蓝

EAST_ASIA_FONT = "宋体"

# 列定义（全量，含"顾问单位""经办律师"列）
DOCX_HEADERS_FULL = ["序号", "日期", "顾问单位", "律师工作内容", "工作时间(小时)", "后续工作", "备注", "经办律师"]
# 总宽度控制在 A4 可用宽度(约17.4cm)内；单单位/分表导出时单位列被移除，宽度更宽松
DOCX_WIDTHS_FULL = [1.0, 1.8, 2.2, 3.8, 1.6, 2.2, 2.2, 1.6]
EXCEL_HEADERS_FULL = ["序号", "日期", "顾问单位", "律师工作内容", "工作时间(小时)", "后续工作", "备注", "经办律师"]
EXCEL_WIDTHS_FULL = [6, 13, 18, 40, 13, 28, 28, 12]


# --------------------------------------------------------------------------- #
# 通用工具
# --------------------------------------------------------------------------- #
def _set_cjk(run, font=EAST_ASIA_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)


def _add_page_field(paragraph, field):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " %s " % field
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def _add_page_number_footer(doc):
    section = doc.sections[0]
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = "第 "
    _add_page_field(p, "PAGE")
    p.add_run(" 页 / 共 ")
    _add_page_field(p, "NUMPAGES")
    p.add_run(" 页")
    for r in p.runs:
        _set_cjk(r)
        r.font.size = Pt(9)


def _period_label(date_from, date_to):
    if date_from or date_to:
        return "%s 至 %s" % (date_from or "最早", date_to or "最新")
    return "全部期间"


def _group_by_unit(records, unit_order=None):
    """按单位分组，返回 (有序单位列表, {单位: 记录列表})。未分类('')置于最后。"""
    groups = OrderedDict()
    for r in records:
        u = (r.get("unit") or "").strip()
        groups.setdefault(u, []).append(r)
    ordered = []
    if unit_order:
        for u in unit_order:
            if u in groups:
                ordered.append(u)
    for u in groups:
        if u not in ordered:
            ordered.append(u)
    if "" in ordered:
        ordered.remove("")
        ordered.append("")  # 未分类单位放最后
    return ordered, groups


def _label_for_unit(u):
    return u if u else "未分类单位"


# --------------------------------------------------------------------------- #
# DOCX —— 单表（保持原有精确行为）
# --------------------------------------------------------------------------- #
def _fill_docx_table(doc, records, show_unit_col):
    headers = list(DOCX_HEADERS_FULL)
    widths = list(DOCX_WIDTHS_FULL)
    if not show_unit_col:
        idx = headers.index("顾问单位")
        headers.pop(idx)
        widths.pop(idx)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        _set_cjk(run)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), C_HEADER_BG)
        hdr[i]._tc.get_or_add_tcPr().append(shd)

    for idx, r in enumerate(records, 1):
        cells = table.add_row().cells
        vals = [
            str(idx),
            r.get("date", ""),
            r.get("unit", ""),
            r.get("content", ""),
            str(r.get("hours", "")),
            r.get("followup", ""),
            r.get("remark", ""),
            r.get("lawyer", ""),
        ]
        if not show_unit_col:
            vals.pop(2)  # 去掉单位列（经办律师列保持显示）
        for i, v in enumerate(vals):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(v)
            _set_cjk(run)
            run.font.size = Pt(9)
            if idx % 2 == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), C_BAND)
                cells[i]._tc.get_or_add_tcPr().append(shd)

    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def build_docx(records, unit_label, date_from, date_to, split_by_unit=False, unit_order=None, show_unit_col=True, overview=False, unit_meta=None):
    unit_meta = unit_meta or {}
    if split_by_unit:
        return _build_docx_split(records, date_from, date_to, unit_order, unit_meta)
    if overview:
        # 总表概览：仅列示服务对象名称与服务期限，不含任何工作明细
        return _build_docx_overview(unit_order or [], unit_meta)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("顾问单位法律服务工作记录汇总报告")
    _set_cjk(trun)
    trun.font.size = Pt(16)
    trun.font.bold = True
    trun.font.color.rgb = C_TITLE

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub.add_run("顾问单位：%s　　汇总期间：%s" % (unit_label, _period_label(date_from, date_to)))
    _set_cjk(srun)
    srun.font.size = Pt(10.5)
    srun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    gen = doc.add_paragraph()
    gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    grun = gen.add_run("报告生成时间：%s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    _set_cjk(grun)
    grun.font.size = Pt(9)
    grun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    total_hours = round(sum(float(r.get("hours") or 0) for r in records), 2)
    stat = doc.add_paragraph()
    strun = stat.add_run("本期间共记录工作事项 %d 项，累计法律服务工作时间 %s 小时。" % (len(records), total_hours))
    _set_cjk(strun)
    strun.font.size = Pt(10.5)
    strun.font.bold = True

    _fill_docx_table(doc, records, show_unit_col=show_unit_col)
    _add_page_number_footer(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_overview(units, unit_meta):
    """总表概览：仅列示各顾问单位的服务对象名称与服务期限（依据律师服务合同）。"""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("顾问单位一览（总表）")
    _set_cjk(trun)
    trun.font.size = Pt(16)
    trun.font.bold = True
    trun.font.color.rgb = C_TITLE

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub.add_run("本表仅列示各顾问单位的服务对象名称与服务期限（依据律师服务合同约定）；"
                       "具体法律服务工作明细见各单位的独立分表。")
    _set_cjk(srun)
    srun.font.size = Pt(10.5)
    srun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    gen = doc.add_paragraph()
    gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    grun = gen.add_run("报告生成时间：%s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    _set_cjk(grun)
    grun.font.size = Pt(9)
    grun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    headers = ["序号", "服务对象名称", "服务开始日期（合同约定）", "服务结束日期（合同约定）"]
    widths = [1.2, 4.5, 4.0, 4.0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        _set_cjk(run)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), C_HEADER_BG)
        hdr[i]._tc.get_or_add_tcPr().append(shd)

    real_units = [u for u in units if u]  # 排除未分类单位
    if not real_units:
        cells = table.add_row().cells
        cells[0].merge(cells[-1])
        pc = cells[0].paragraphs[0]
        r = pc.add_run("（暂无顾问单位。请先在网页「单位管理」中新建单位并设置服务期限。）")
        _set_cjk(r)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    else:
        for idx, u in enumerate(real_units, 1):
            meta = unit_meta.get(u, {}) or {}
            sd = (meta.get("start_date") or "").strip() or "—"
            ed = (meta.get("end_date") or "").strip() or "—"
            cells = table.add_row().cells
            vals = [str(idx), u, sd, ed]
            for i, v in enumerate(vals):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                run = p.add_run(v)
                _set_cjk(run)
                run.font.size = Pt(10)
                if idx % 2 == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), C_BAND)
                    cells[i]._tc.get_or_add_tcPr().append(shd)

    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)

    _add_page_number_footer(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_split(records, date_from, date_to, unit_order, unit_meta=None):
    unit_meta = unit_meta or {}
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("顾问单位法律服务工作记录·分单位汇总报告")
    _set_cjk(trun)
    trun.font.size = Pt(16)
    trun.font.bold = True
    trun.font.color.rgb = C_TITLE

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub.add_run("汇总期间：%s" % _period_label(date_from, date_to))
    _set_cjk(srun)
    srun.font.size = Pt(10.5)
    srun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    gen = doc.add_paragraph()
    gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    grun = gen.add_run("报告生成时间：%s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    _set_cjk(grun)
    grun.font.size = Pt(9)
    grun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    ordered, groups = _group_by_unit(records, unit_order)
    total_hours = round(sum(float(r.get("hours") or 0) for r in records), 2)
    real_units = [u for u in ordered if u]  # 排除未分类
    unclassified = groups.get("", [])
    ov = doc.add_paragraph()
    ovrun = ov.add_run("本报告涵盖 %d 家顾问单位；另含未分类记录 %d 项，共 %d 项工作事项，累计 %s 小时。" % (
        len(real_units), len(unclassified), len(records), total_hours))
    _set_cjk(ovrun)
    ovrun.font.size = Pt(10.5)
    ovrun.font.bold = True

    # 第一部分：总表概览（仅服务对象名称 + 服务期限，不含工作明细）
    h1 = doc.add_paragraph()
    h1run = h1.add_run("一、顾问单位一览（总表）")
    _set_cjk(h1run)
    h1run.font.size = Pt(13)
    h1run.font.bold = True
    h1run.font.color.rgb = C_TITLE

    ov_headers = ["序号", "服务对象名称", "服务开始日期（合同约定）", "服务结束日期（合同约定）"]
    ov_widths = [1.2, 4.5, 4.0, 4.0]
    ov_table = doc.add_table(rows=1, cols=len(ov_headers))
    ov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ov_table.style = "Table Grid"
    ov_hdr = ov_table.rows[0].cells
    for i, h in enumerate(ov_headers):
        ov_hdr[i].text = ""
        p = ov_hdr[i].paragraphs[0]
        run = p.add_run(h)
        _set_cjk(run)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), C_HEADER_BG)
        ov_hdr[i]._tc.get_or_add_tcPr().append(shd)
    if not real_units:
        oc = ov_table.add_row().cells
        oc[0].merge(oc[-1])
        orun = oc[0].paragraphs[0].add_run("（暂无顾问单位）")
        _set_cjk(orun)
        orun.font.size = Pt(10)
        orun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    else:
        for idx, u in enumerate(real_units, 1):
            meta = unit_meta.get(u, {}) or {}
            sd = (meta.get("start_date") or "").strip() or "—"
            ed = (meta.get("end_date") or "").strip() or "—"
            oc = ov_table.add_row().cells
            for i, v in enumerate([str(idx), u, sd, ed]):
                oc[i].text = ""
                run = oc[i].paragraphs[0].add_run(v)
                _set_cjk(run)
                run.font.size = Pt(10)
                if idx % 2 == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), C_BAND)
                    oc[i]._tc.get_or_add_tcPr().append(shd)
    for row in ov_table.rows:
        for i, w in enumerate(ov_widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()

    # 第二部分：各单位工作明细分表
    h2 = doc.add_paragraph()
    h2run = h2.add_run("二、各单位法律服务工作明细分表")
    _set_cjk(h2run)
    h2run.font.size = Pt(13)
    h2run.font.bold = True
    h2run.font.color.rgb = C_TITLE
    doc.add_paragraph()

    for i, u in enumerate(ordered, 1):
        recs = groups[u]
        h = round(sum(float(r.get("hours") or 0) for r in recs), 2)
        hpar = doc.add_paragraph()
        hrun = hpar.add_run("%d、%s（%d 项，%s 小时）" % (i, _label_for_unit(u), len(recs), h))
        _set_cjk(hrun)
        hrun.font.size = Pt(12.5)
        hrun.font.bold = True
        hrun.font.color.rgb = C_TITLE
        _fill_docx_table(doc, recs, show_unit_col=False)
        doc.add_paragraph()  # 分表间隔

    _add_page_number_footer(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Excel —— 单表 / 多分表
# --------------------------------------------------------------------------- #
def _excel_cell(ws, row, col, value, font, alignment, border, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.alignment = alignment
    cell.border = border
    if fill:
        cell.fill = fill
    return cell


def _fill_excel_sheet(ws, records, label, date_from, date_to, show_unit_col):
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=C_HEADER_BG)
    band_fill = PatternFill("solid", fgColor=C_BAND)
    header_font = Font(name="微软雅黑", bold=True, color=C_HEADER_FONT, size=10)
    cell_font = Font(name="微软雅黑", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = list(EXCEL_HEADERS_FULL)
    widths = list(EXCEL_WIDTHS_FULL)
    if not show_unit_col:
        idx = headers.index("顾问单位")
        headers.pop(idx)
        widths.pop(idx)

    # 大标题
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    t = ws.cell(row=1, column=1, value="顾问单位法律服务工作记录（%s / %s）" % (label, _period_label(date_from, date_to)))
    t.font = Font(name="微软雅黑", bold=True, size=13, color="1F3A5F")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for c, h in enumerate(headers, 1):
        _excel_cell(ws, 2, c, h, header_font, center, border, header_fill)

    for i, r in enumerate(records, 1):
        row = 2 + i
        vals = [
            i,
            r.get("date", ""),
            r.get("unit", ""),
            r.get("content", ""),
            float(r.get("hours") or 0),
            r.get("followup", ""),
            r.get("remark", ""),
            r.get("lawyer", ""),
        ]
        if not show_unit_col:
            vals.pop(2)
        # 左对齐列按表头名称动态判定，避免"单位列"被移除后列号错位
        # （如工时列曾被硬编码为左对齐）。这里仅文本类字段左对齐，序号/日期/工时居中。
        LEFT_COLS = {"顾问单位", "律师工作内容", "后续工作", "备注", "经办律师"}
        for c, v in enumerate(vals, 1):
            cell = _excel_cell(ws, row, c, v, cell_font, border=border,
                               alignment=(left if headers[c - 1] in LEFT_COLS else center))
            if i % 2 == 0:
                cell.fill = band_fill

    # 合计行
    total_row = 2 + len(records) + 1
    hours_col = headers.index("工作时间(小时)") + 1
    content_col = headers.index("律师工作内容") + 1
    _excel_cell(ws, total_row, 1, "合计", Font(name="微软雅黑", bold=True, size=10), center, border)
    _excel_cell(ws, total_row, content_col, "本期间共 %d 项" % len(records),
                Font(name="微软雅黑", bold=True, size=10), left, border)
    _excel_cell(ws, total_row, hours_col, round(sum(float(r.get("hours") or 0) for r in records), 2),
                Font(name="微软雅黑", bold=True, size=10), center, border)
    for c in range(1, len(headers) + 1):
        if c not in (1, content_col, hours_col):
            _excel_cell(ws, total_row, c, "", Font(name="微软雅黑", bold=True, size=10), center, border)

    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A3"


def _safe_sheet_name(name):
    if not name:
        return "未分类单位"
    bad = ':\\/?*[]'
    s = ''.join('_' if c in bad else c for c in name)
    return s[:31]


def _fill_excel_overview(ws, units, unit_meta):
    """总表概览 sheet：仅服务对象名称与服务期限。"""
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=C_HEADER_BG)
    band_fill = PatternFill("solid", fgColor=C_BAND)
    header_font = Font(name="微软雅黑", bold=True, color=C_HEADER_FONT, size=10)
    cell_font = Font(name="微软雅黑", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["序号", "服务对象名称", "服务开始日期（合同约定）", "服务结束日期（合同约定）"]
    widths = [6, 32, 26, 26]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    t = ws.cell(row=1, column=1, value="顾问单位一览（总表）— 仅列示服务对象名称与服务期限，工作明细见各分表")
    t.font = Font(name="微软雅黑", bold=True, size=13, color="1F3A5F")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for c, h in enumerate(headers, 1):
        _excel_cell(ws, 2, c, h, header_font, center, border, header_fill)

    real_units = [u for u in units if u]
    if not real_units:
        _excel_cell(ws, 3, 1, "（暂无顾问单位。请先在网页「单位管理」中新建单位并设置服务期限。）",
                    cell_font, left, border)
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    else:
        for i, u in enumerate(real_units, 1):
            meta = unit_meta.get(u, {}) or {}
            sd = (meta.get("start_date") or "").strip() or "—"
            ed = (meta.get("end_date") or "").strip() or "—"
            row = 2 + i
            for c, v in enumerate([i, u, sd, ed], 1):
                cell = _excel_cell(ws, row, c, v, cell_font, (left if c == 2 else center), border)
                if i % 2 == 0:
                    cell.fill = band_fill

    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A3"


def build_excel(records, unit_label, date_from, date_to, split_by_unit=False, unit_order=None, show_unit_col=True, overview=False, unit_meta=None):
    unit_meta = unit_meta or {}
    wb = openpyxl.Workbook()

    if overview and not split_by_unit:
        # 总表概览：仅单位一览（名称 + 服务期限）
        ws = wb.active
        ws.title = "总表（单位一览）"
        _fill_excel_overview(ws, unit_order or [], unit_meta)
    elif split_by_unit:
        # 第一个 sheet：总表概览；其后每家单位一张独立明细分表
        ws_all = wb.active
        ws_all.title = "总表（单位一览）"
        _fill_excel_overview(ws_all, unit_order or [], unit_meta)
        ordered, groups = _group_by_unit(records, unit_order)
        for u in ordered:
            ws = wb.create_sheet(title=_safe_sheet_name(u))
            _fill_excel_sheet(ws, groups[u], _label_for_unit(u), date_from, date_to, show_unit_col=False)
    else:
        ws = wb.active
        ws.title = "工作记录汇总"
        _fill_excel_sheet(ws, records, unit_label, date_from, date_to, show_unit_col=show_unit_col)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
