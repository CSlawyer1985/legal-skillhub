#!/usr/bin/env python3
"""合同智审竞赛 Skill 的通用演示工具。

功能：
- inspect: 检查 DOCX 的段落、表格、金额、日期和基础风险；
- demo-redline: 对竞赛示例合同执行定点红线修改，同时生成清洁版；
- verify: 核验 DOCX 可打开、红色覆盖、关键条款、金额和 PDF 页数。
- compare-text: 对两个DOCX的全部正文与表格文本进行哈希比对，用于验证内容一致性。封面等局部修改场景可先提取非目标区域后调用。

该脚本不调用任何大模型，便于评委复现文档处理与验证能力。
"""

from __future__ import annotations

import argparse
import json
import hashlib
import re
import shutil
import subprocess
import sys
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

RED = RGBColor(0xC0, 0x00, 0x00)
MONEY_RE = re.compile(r"(?:人民币\s*)?[¥￥]?\s*([0-9][0-9,]*(?:\.\d{1,2})?)\s*元")
DATE_RE = re.compile(r"(?:20\d{2}年\d{1,2}月\d{1,2}日|20\d{2}-\d{1,2}-\d{1,2})")


def iter_paragraphs(doc: Document) -> Iterable[Paragraph]:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def document_text(doc: Document) -> str:
    return "\n".join(p.text for p in iter_paragraphs(doc))


def set_run_font(run, red=False, bold=None, name="宋体", size=10.5):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if red:
        run.font.color.rgb = RED


def clear_paragraph(paragraph: Paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph(paragraph: Paragraph, text: str, red: bool):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, red=red)


def insert_after(paragraph: Paragraph, text: str, red: bool) -> Paragraph:
    new_p = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        new_p.append(deepcopy(paragraph._p.pPr))
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    run = p.add_run(text)
    set_run_font(run, red=red)
    return p


def find_exact(doc: Document, text: str) -> Paragraph:
    matches = [p for p in iter_paragraphs(doc) if p.text.strip() == text.strip()]
    if len(matches) != 1:
        raise ValueError(f"目标文本应唯一，实际命中 {len(matches)} 次：{text}")
    return matches[0]


def red_block_count(doc: Document) -> int:
    n = 0
    for p in iter_paragraphs(doc):
        if any(r.font.color and r.font.color.rgb and str(r.font.color.rgb) == "C00000" for r in p.runs):
            n += 1
    return n


def extract_amounts(text: str):
    values = []
    for raw in MONEY_RE.findall(text):
        try:
            values.append(Decimal(raw.replace(",", "")))
        except Exception:
            pass
    return values


def inspect_docx(path: Path) -> dict:
    doc = Document(path)
    text = document_text(doc)
    amounts = extract_amounts(text)
    issues = []

    if "身份证号码：________________" in text or "统一社会信用代码：________________" in text:
        issues.append({"level": "B", "issue": "主体识别信息存在空白"})
    if re.search(r"账号[:：]\s*_+", text):
        issues.append({"level": "A", "issue": "收付款账户尚未填写"})
    if "任何损失" in text or "全部损失" in text:
        issues.append({"level": "B", "issue": "损失范围可能过宽，应审查过错、因果关系和合理性"})
    if "具体范围由乙方在交付时确定" in text or "以乙方实际提供为准" in text:
        issues.append({"level": "A", "issue": "合同标的及交付范围由乙方单方确定，甲方付款后仍无法锁定取得内容"})
    if "一次性支付全部价款" in text and "收到全部价款后" in text:
        issues.append({"level": "A", "issue": "甲方先全额付款、乙方后交付，付款与交付未形成制约"})
    if "任一期逾期" in text and ("全部" in text or "提前到期" in text):
        issues.append({"level": "A", "issue": "存在一期逾期触发全部债务提前到期的风险"})
    if "每日千分之五" in text:
        issues.append({"level": "A", "issue": "日千分之五违约金按简单年化计算约为182.5%，存在显著调减与谈判风险"})
    if "乙方逾期交付的，双方另行协商" in text:
        issues.append({"level": "A", "issue": "乙方逾期交付无明确违约责任、退款或解除机制"})
    if "甲方所在地人民法院" in text or "甲方住所地人民法院" in text:
        issues.append({"level": "B", "issue": "管辖条款可能单方有利，应结合委托方立场审查"})
    if "乙方所在地人民法院" in text:
        issues.append({"level": "B", "issue": "约定乙方所在地法院管辖，可能增加甲方维权成本"})

    return {
        "file": str(path),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "characters": len(text),
        "amounts_found": [str(v) for v in amounts],
        "dates_found": DATE_RE.findall(text),
        "red_blocks": red_block_count(doc),
        "issues": issues,
    }


def demo_redline(src: Path, out_red: Path, out_clean: Path):
    old_scope = "乙方向甲方转让生产设备、技术资料及相关配套物品，具体范围由乙方在交付时确定。"
    old_payment = "甲方应于本协议签订之日起三日内一次性支付全部价款。"
    old_delivery = "乙方在收到全部价款后十日内办理交付，具体交付内容以乙方实际提供为准。"
    old_default = "甲方任何一期付款逾期，全部未付款项立即到期，并按合同总价每日千分之五支付违约金。"
    old_seller_default = "乙方逾期交付的，双方另行协商。"
    old_forum = "因本合同发生争议，由乙方所在地人民法院管辖。"

    new_scope = "乙方向甲方转让的设备、技术资料及配套物品，以双方签署并作为本合同附件的《交付清单》为准；未经甲方书面同意，乙方不得单方变更或减少转让范围。"
    new_payment = "甲方分期支付价款；每一期付款均以乙方完成对应交付义务并提交书面交付清单为条件。甲方在收到完整交付材料并确认无误后五个工作日内支付对应款项。"
    new_delivery = "乙方应按照附件《交付清单》完整交付合同标的、资料及相关凭证。交付内容、数量和状态经双方签字确认；未经甲方书面确认，不视为完成交付。"
    new_default = "甲方逾期支付已经到期且付款条件已经成就的款项，应以该期逾期未付款为基数承担违约责任；任一期逾期不导致其他未到期款项提前到期。"
    new_seller_default = "乙方逾期交付超过十日，甲方有权要求继续履行并赔偿损失；逾期超过三十日或经催告仍未完成交付的，甲方有权解除合同，乙方应在五个工作日内返还已收款项并赔偿甲方实际损失。"
    new_forum = "因本合同发生争议，双方应先协商解决；协商不成的，任何一方可向与争议有实际联系且依法具有管辖权的人民法院提起诉讼。具体管辖连接点应结合合同履行地、标的所在地及委托方维权便利在签署前确定。"
    new_warranty = "乙方保证对合同标的享有完整处分权，标的不存在质押、冻结、重复转让或其他权利负担；因违反该保证造成甲方损失的，乙方应承担赔偿责任。"

    for output, red in [(out_red, True), (out_clean, False)]:
        shutil.copy2(src, output)
        doc = Document(output)
        p_scope = find_exact(doc, old_scope)
        p_payment = find_exact(doc, old_payment)
        p_delivery = find_exact(doc, old_delivery)
        p_default = find_exact(doc, old_default)
        p_seller_default = find_exact(doc, old_seller_default)
        p_forum = find_exact(doc, old_forum)
        replace_paragraph(p_scope, new_scope, red)
        replace_paragraph(p_payment, new_payment, red)
        replace_paragraph(p_delivery, new_delivery, red)
        insert_after(p_delivery, new_warranty, red)
        replace_paragraph(p_default, new_default, red)
        replace_paragraph(p_seller_default, new_seller_default, red)
        replace_paragraph(p_forum, new_forum, red)
        doc.save(output)

    return {
        "redline": str(out_red),
        "clean": str(out_clean),
        "red_blocks": red_block_count(Document(out_red)),
        "clean_red_blocks": red_block_count(Document(out_clean)),
    }


def pdf_pages(path: Path):
    try:
        p = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, timeout=30)
        m = re.search(r"^Pages:\s+(\d+)", p.stdout, re.M)
        return int(m.group(1)) if m else None
    except Exception:
        return None


def verify(path: Path, pdf: Path | None = None, must_contain=None) -> dict:
    doc = Document(path)
    text = document_text(doc)
    required = must_contain or []
    missing = [x for x in required if x not in text]
    result = {
        "file": str(path),
        "openable": True,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "red_blocks": red_block_count(doc),
        "missing_required_text": missing,
        "amounts_found": [str(v) for v in extract_amounts(text)],
    }
    if pdf:
        result["pdf"] = str(pdf)
        result["pdf_pages"] = pdf_pages(pdf)
    result["passed"] = not missing and result["openable"]
    return result


def build_review_memo(out: Path):
    content = """# 示例合同审核意见书（甲方／受让方立场）

## 一、审查结论

现有《设备及资料转让合同》不建议由甲方直接签署。核心风险在于：转让范围由乙方单方确定，甲方须先全额付款，乙方交付义务和逾期责任不明确，甲方逾期责任畸重，且争议由乙方所在地法院管辖。

## 二、请求—抗辩—再抗辩示例

甲方付款后要求完整交付 → 乙方可能主张“实际提供即为约定范围” → 原合同没有清单，甲方举证困难 → 应以双方签署的《交付清单》锁定标的，并将付款与对应交付联动。

## 三、主要风险及条款化建议

1. **A 必须修改｜标的不确定。**以双方签署的《交付清单》锁定名称、数量和状态，未经甲方书面同意不得减少或变更。
2. **A 必须修改｜先全额付款后交付。**改为分期支付，每期付款以对应交付和甲方书面确认作为成就条件。
3. **A 必须修改｜交付及权利保证缺失。**增加签字验收及无质押、冻结、重复转让等权利负担保证。
4. **A 必须修改｜一期逾期导致全部加速且违约金畸高。**仅以已经到期且付款条件成就的当期未付款为基数，不触发其他未到期款项。
5. **A 必须修改｜乙方违约责任空缺。**增加继续履行、赔偿、解除及五个工作日退款机制。
6. **B 建议修改｜管辖偏向乙方。**调整为合同履行地有管辖权的人民法院。

## 四、签署前开放问题

双方统一社会信用代码、真实开户行和账号、《交付清单》、分期金额和节点、签章及日期尚待补齐。因此生成文件应称为“待签署清洁版”，不得表述为已经签署或生效。

> 本示例全部为虚构数据，仅用于展示Skill工作流程，不构成针对真实交易的法律意见。
"""
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(content, encoding="utf-8")
    return {"file": str(out), "characters": len(content), "passed": True}


def apply_plan(src: Path, plan_path: Path, out_red: Path, out_clean: Path):
    """按结构化JSON计划对任意DOCX执行精确替换/插入。

    计划格式：
    {"operations":[
      {"action":"replace", "target":"唯一原文", "text":"新条款"},
      {"action":"insert_after", "target":"唯一锚点", "text":"新增条款"}
    ]}
    """
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    operations = plan.get("operations", [])
    if not operations:
        raise ValueError("修改计划缺少operations")
    for output, red in [(out_red, True), (out_clean, False)]:
        shutil.copy2(src, output)
        doc = Document(output)
        for idx, op in enumerate(operations, 1):
            action = op.get("action")
            target = op.get("target", "")
            new_text = op.get("text", "")
            if not target or not new_text:
                raise ValueError(f"第{idx}项操作缺少target或text")
            p = find_exact(doc, target)
            if action == "replace":
                replace_paragraph(p, new_text, red)
            elif action == "insert_after":
                insert_after(p, new_text, red)
            else:
                raise ValueError(f"不支持的action：{action}")
        doc.save(output)
    return {
        "source": str(src), "plan": str(plan_path), "operations": len(operations),
        "redline": str(out_red), "clean": str(out_clean),
        "red_blocks": red_block_count(Document(out_red)),
        "clean_red_blocks": red_block_count(Document(out_clean)),
        "passed": red_block_count(Document(out_red)) > 0 and red_block_count(Document(out_clean)) == 0,
    }


def verify_pair(red_path: Path, clean_path: Path, red_pdf: Path | None = None,
                clean_pdf: Path | None = None, must_contain=None, must_absent=None) -> dict:
    red_doc, clean_doc = Document(red_path), Document(clean_path)
    red_text, clean_text = document_text(red_doc), document_text(clean_doc)
    required = must_contain or []
    forbidden = must_absent or []
    missing = [x for x in required if x not in red_text or x not in clean_text]
    old_remaining = [x for x in forbidden if x in red_text or x in clean_text]
    result = {
        "redline": str(red_path), "clean": str(clean_path),
        "openable": True,
        "red_blocks": red_block_count(red_doc),
        "clean_red_blocks": red_block_count(clean_doc),
        "texts_identical": red_text == clean_text,
        "missing_required_text": missing,
        "forbidden_text_remaining": old_remaining,
        "amounts_identical": extract_amounts(red_text) == extract_amounts(clean_text),
    }
    if red_pdf:
        result["red_pdf_pages"] = pdf_pages(red_pdf)
    if clean_pdf:
        result["clean_pdf_pages"] = pdf_pages(clean_pdf)
    pdf_ok = (not red_pdf or result.get("red_pdf_pages")) and (not clean_pdf or result.get("clean_pdf_pages"))
    result["passed"] = bool(
        result["red_blocks"] > 0 and result["clean_red_blocks"] == 0 and
        result["texts_identical"] and result["amounts_identical"] and
        not missing and not old_remaining and pdf_ok
    )
    return result


def docx_text_hash(path: Path, skip_paragraphs: int = 0) -> dict:
    doc = Document(path)
    body = list(iter_paragraphs(doc))
    text = "\n".join(p.text for p in body[skip_paragraphs:])
    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    return {
        "file": str(path),
        "skip_paragraphs": skip_paragraphs,
        "characters": len(normalized),
        "sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "text": normalized,
    }


def compare_text(before: Path, after: Path, skip_paragraphs: int = 0) -> dict:
    a = docx_text_hash(before, skip_paragraphs)
    b = docx_text_hash(after, skip_paragraphs)
    return {
        "before": {k: v for k, v in a.items() if k != "text"},
        "after": {k: v for k, v in b.items() if k != "text"},
        "identical": a["text"] == b["text"],
        "passed": a["text"] == b["text"],
    }


def main():
    ap = argparse.ArgumentParser(description="合同智审演示与验证工具")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p1 = sub.add_parser("inspect")
    p1.add_argument("docx", type=Path)
    p1.add_argument("--out", type=Path)

    p2 = sub.add_parser("demo-redline")
    p2.add_argument("src", type=Path)
    p2.add_argument("--red", type=Path, required=True)
    p2.add_argument("--clean", type=Path, required=True)

    p3 = sub.add_parser("verify")
    p3.add_argument("docx", type=Path)
    p3.add_argument("--pdf", type=Path)
    p3.add_argument("--contains", action="append", default=[])
    p3.add_argument("--out", type=Path)

    p4 = sub.add_parser("compare-text")
    p4.add_argument("before", type=Path)
    p4.add_argument("after", type=Path)
    p4.add_argument("--skip-paragraphs", type=int, default=0,
                    help="忽略文档遍历顺序开头的N个段落，用于只改封面等局部修改验证")
    p4.add_argument("--out", type=Path)

    p5 = sub.add_parser("review-memo")
    p5.add_argument("--out", type=Path, required=True)

    p6 = sub.add_parser("apply-plan")
    p6.add_argument("src", type=Path)
    p6.add_argument("--plan", type=Path, required=True)
    p6.add_argument("--red", type=Path, required=True)
    p6.add_argument("--clean", type=Path, required=True)
    p6.add_argument("--out", type=Path)

    p7 = sub.add_parser("verify-pair")
    p7.add_argument("red", type=Path)
    p7.add_argument("clean", type=Path)
    p7.add_argument("--red-pdf", type=Path)
    p7.add_argument("--clean-pdf", type=Path)
    p7.add_argument("--contains", action="append", default=[])
    p7.add_argument("--absent", action="append", default=[])
    p7.add_argument("--out", type=Path)

    args = ap.parse_args()
    if args.cmd == "inspect":
        data = inspect_docx(args.docx)
    elif args.cmd == "demo-redline":
        args.red.parent.mkdir(parents=True, exist_ok=True)
        args.clean.parent.mkdir(parents=True, exist_ok=True)
        data = demo_redline(args.src, args.red, args.clean)
    elif args.cmd == "verify":
        data = verify(args.docx, args.pdf, args.contains)
    elif args.cmd == "compare-text":
        data = compare_text(args.before, args.after, args.skip_paragraphs)
    elif args.cmd == "review-memo":
        data = build_review_memo(args.out)
    elif args.cmd == "apply-plan":
        args.red.parent.mkdir(parents=True, exist_ok=True)
        args.clean.parent.mkdir(parents=True, exist_ok=True)
        data = apply_plan(args.src, args.plan, args.red, args.clean)
    else:
        data = verify_pair(args.red, args.clean, args.red_pdf, args.clean_pdf,
                           args.contains, args.absent)

    payload = json.dumps(data, ensure_ascii=False, indent=2)
    if getattr(args, "out", None) and args.cmd != "review-memo":
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(payload, encoding="utf-8")
    print(payload)
    return 0 if data.get("passed", True) else 2


if __name__ == "__main__":
    sys.exit(main())
